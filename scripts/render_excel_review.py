#!/usr/bin/env python3
"""Render workbook sheets and wide-sheet panels to PNG for visual review."""

from __future__ import annotations

import argparse
import io
import textwrap
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont


REGULAR_FONT = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
BOLD_FONT = Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
ITALIC_FONT = Path("/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf")
DEFAULT_FILL = "#FFFFFF"
DEFAULT_TEXT = "#292C2F"


def load_font(size: float, *, bold: bool = False, italic: bool = False) -> ImageFont.FreeTypeFont:
    if bold:
        path = BOLD_FONT
    elif italic:
        path = ITALIC_FONT
    else:
        path = REGULAR_FONT
    pixel_size = max(9, round((size or 10) * 1.18))
    return ImageFont.truetype(str(path), pixel_size)


def rgb(value, fallback: str) -> str:
    if value is None:
        return fallback
    color_type = getattr(value, "type", None)
    raw = getattr(value, "rgb", None)
    if color_type == "rgb" and raw:
        return "#" + raw[-6:]
    return fallback


def column_pixels(sheet, column: int) -> int:
    letter = get_column_letter(column)
    width = sheet.column_dimensions[letter].width
    if width is None:
        width = 10
    return max(28, round(width * 7.1 + 5))


def row_pixels(sheet, row: int) -> int:
    height = sheet.row_dimensions[row].height
    if height is None:
        height = 18
    return max(20, round(height * 1.33))


def wrap_line(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    if text == "":
        return [""]
    words = text.split(" ")
    lines: list[str] = []
    current = ""
    for word in words:
        candidate = word if not current else f"{current} {word}"
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
            continue
        if current:
            lines.append(current)
        if draw.textbbox((0, 0), word, font=font)[2] <= max_width:
            current = word
        else:
            estimated = max(5, int(len(word) * max_width / max(1, draw.textbbox((0, 0), word, font=font)[2])))
            chunks = textwrap.wrap(word, width=estimated, break_long_words=True, break_on_hyphens=False)
            lines.extend(chunks[:-1])
            current = chunks[-1] if chunks else word
    if current:
        lines.append(current)
    return lines or [""]


def wrapped_text(
    draw: ImageDraw.ImageDraw,
    value: object,
    font,
    max_width: int,
) -> str:
    raw = "" if value is None else str(value)
    lines: list[str] = []
    for paragraph in raw.splitlines() or [""]:
        lines.extend(wrap_line(draw, paragraph, font, max_width))
    return "\n".join(lines)


def merged_range_for(sheet, row: int, column: int):
    for merged_range in sheet.merged_cells.ranges:
        if (
            merged_range.min_row <= row <= merged_range.max_row
            and merged_range.min_col <= column <= merged_range.max_col
        ):
            return merged_range
    return None


def paste_logo(sheet, canvas: Image.Image, x_positions: dict[int, int], y_positions: dict[int, int]) -> None:
    if not sheet._images or 1 not in x_positions or 1 not in y_positions:
        return
    workbook_image = sheet._images[0]
    try:
        logo = Image.open(io.BytesIO(workbook_image._data())).convert("RGBA")
    except Exception:
        return
    logo.thumbnail((115, 58), Image.Resampling.LANCZOS)
    canvas.alpha_composite(logo, (x_positions[1] + 6, y_positions[1] + 3))


def render_panel(
    sheet,
    columns: list[int],
    rows: list[int],
    output_path: Path,
) -> None:
    widths = {column: column_pixels(sheet, column) for column in columns}
    heights = {row: row_pixels(sheet, row) for row in rows}
    x_positions: dict[int, int] = {}
    y_positions: dict[int, int] = {}

    x = 0
    for column in columns:
        x_positions[column] = x
        x += widths[column]
    y = 0
    for row in rows:
        y_positions[row] = y
        y += heights[row]

    canvas = Image.new("RGBA", (x, y), "#F6F2ED")
    draw = ImageDraw.Draw(canvas)
    selected_columns = set(columns)
    selected_rows = set(rows)

    for row in rows:
        for column in columns:
            cell = sheet.cell(row=row, column=column)
            merged = merged_range_for(sheet, row, column)
            if isinstance(cell, MergedCell):
                continue
            if merged and (cell.row != merged.min_row or cell.column != merged.min_col):
                continue

            span_columns = [column]
            span_rows = [row]
            if merged:
                span_columns = [
                    candidate
                    for candidate in columns
                    if merged.min_col <= candidate <= merged.max_col
                ]
                span_rows = [
                    candidate
                    for candidate in rows
                    if merged.min_row <= candidate <= merged.max_row
                ]
            if not span_columns or not span_rows:
                continue

            left = x_positions[min(span_columns)]
            top = y_positions[min(span_rows)]
            right = x_positions[max(span_columns)] + widths[max(span_columns)]
            bottom = y_positions[max(span_rows)] + heights[max(span_rows)]

            fill_color = rgb(cell.fill.fgColor, DEFAULT_FILL)
            draw.rectangle((left, top, right, bottom), fill=fill_color)

            bottom_side = cell.border.bottom
            top_side = cell.border.top
            border_color = rgb(getattr(bottom_side, "color", None), "#D8D0C8")
            if getattr(bottom_side, "style", None):
                line_width = 2 if bottom_side.style in {"medium", "thick"} else 1
                draw.line((left, bottom - 1, right, bottom - 1), fill=border_color, width=line_width)
            if getattr(top_side, "style", None):
                top_color = rgb(getattr(top_side, "color", None), "#6B1F2B")
                line_width = 2 if top_side.style in {"medium", "thick"} else 1
                draw.line((left, top, right, top), fill=top_color, width=line_width)

            if cell.value is None:
                continue

            font = load_font(
                cell.font.sz or 10,
                bold=bool(cell.font.bold),
                italic=bool(cell.font.italic),
            )
            font_color = rgb(cell.font.color, DEFAULT_TEXT)
            padding_x = 6
            padding_y = 5
            text = wrapped_text(draw, cell.value, font, max(20, right - left - 2 * padding_x))
            text_box = draw.multiline_textbbox((0, 0), text, font=font, spacing=3)
            text_height = text_box[3] - text_box[1]

            horizontal = cell.alignment.horizontal or "left"
            vertical = cell.alignment.vertical or "top"
            if horizontal == "center":
                anchor_x = left + (right - left) / 2
                anchor = "ma"
            elif horizontal == "right":
                anchor_x = right - padding_x
                anchor = "ra"
            else:
                anchor_x = left + padding_x
                anchor = "la"

            if vertical == "center":
                anchor_y = top + max(padding_y, (bottom - top - text_height) / 2)
            elif vertical == "bottom":
                anchor_y = bottom - text_height - padding_y
            else:
                anchor_y = top + padding_y

            cell_layer = Image.new("RGBA", (max(1, right - left), max(1, bottom - top)), (0, 0, 0, 0))
            cell_draw = ImageDraw.Draw(cell_layer)
            local_x = anchor_x - left
            local_y = anchor_y - top
            cell_draw.multiline_text(
                (local_x, local_y),
                text,
                font=font,
                fill=font_color,
                spacing=3,
                anchor=anchor,
            )
            canvas.alpha_composite(cell_layer, (left, top))

    paste_logo(sheet, canvas, x_positions, y_positions)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.convert("RGB").save(output_path, quality=95)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--workbook",
        type=Path,
        default=Path(
            "outputs/client-preview/"
            "NeuraFlow_Product_Opportunity_Evidence_Illustrative_2026-08-15.xlsx"
        ),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("media/client-preview/excel-review"),
    )
    args = parser.parse_args()

    workbook = load_workbook(args.workbook, data_only=False, keep_links=False)
    panels = [
        ("Shortlist", list(range(1, 16)), list(range(1, 19)), "01-shortlist-full.png"),
        ("Evidence", list(range(1, 20)), list(range(1, 19)), "02-evidence-identity-meta.png"),
        (
            "Evidence",
            [1, 2, 3, *range(20, 38)],
            list(range(1, 19)),
            "03-evidence-market-search.png",
        ),
        (
            "Evidence",
            [1, 2, 3, *range(38, 50)],
            list(range(1, 19)),
            "04-evidence-context-review.png",
        ),
        ("Supplier Leads", list(range(1, 16)), list(range(1, 43)), "05-supplier-leads-full.png"),
    ]

    for sheet_name, columns, rows, filename in panels:
        render_panel(
            workbook[sheet_name],
            columns,
            rows,
            args.output_dir / filename,
        )
        print(args.output_dir / filename)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
