#!/usr/bin/env python3
"""Generate the NeuraFlow client-preview evidence workbook from canonical JSON."""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.worksheet import Worksheet


CREAM = "F9F3EC"
CREAM_DARK = "EEE5DB"
MAROON = "6B1F2B"
NAVY = "24384B"
MUTED_NAVY = "647487"
MUSTARD = "C49A37"
SAGE = "93A28E"
ROSE = "C9949A"
CHARCOAL = "292C2F"
MID_GREY = "696A6C"
LIGHT_GREY = "F3F0EC"
GRID = "D8D0C8"
WHITE = "FFFFFF"
PALE_NAVY = "E5EAF0"
PALE_MAROON = "EEDDE0"
PALE_MUSTARD = "F1E5C5"
PALE_SAGE = "E1E8DE"
PALE_ROSE = "F0E1E3"

THIN_GRID = Side(style="thin", color=GRID)
MEDIUM_MAROON = Side(style="medium", color=MAROON)

LABEL_FILLS = {
    "Strong": PALE_SAGE,
    "Promising": PALE_MAROON,
    "Worth Watching": PALE_MUSTARD,
}
STATUS_FILLS = {
    "New": PALE_NAVY,
    "Returning": LIGHT_GREY,
    "Strengthened": PALE_SAGE,
    "Weakened": PALE_ROSE,
    "Coverage Changed": PALE_MUSTARD,
}
CLARITY_FILLS = {
    "Clear": PALE_SAGE,
    "Partial": PALE_MUSTARD,
    "Uncertain": PALE_ROSE,
}


def safe_value(value: Any) -> Any:
    """Neutralise spreadsheet formula prefixes in untrusted strings."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def join_values(values: Iterable[Any], separator: str = " · ") -> str:
    return separator.join(str(value) for value in values if value not in (None, ""))


def bool_label(value: bool) -> str:
    return "Available" if value else "Not available"


def format_timestamp(value: str | None) -> str:
    if not value:
        return "Not available"
    parsed = datetime.fromisoformat(value)
    return parsed.strftime("%d %b %Y, %H:%M")


def quick_presence(product: dict[str, Any]) -> int:
    return sum(
        state == "Present"
        for city_values in product["quick_commerce"].values()
        for state in city_values.values()
    )


def quick_summary(product: dict[str, Any], platform: str) -> str:
    parts = []
    for city in ("Mumbai", "Delhi", "Bengaluru"):
        state = product["quick_commerce"][city][platform]
        parts.append(f"{city}: {state}")
    return "; ".join(parts)


def set_cell(
    sheet: Worksheet,
    row: int,
    column: int,
    value: Any,
    *,
    font: Font | None = None,
    fill: PatternFill | None = None,
    alignment: Alignment | None = None,
    border: Border | None = None,
    number_format: str | None = None,
) -> Cell:
    cell = sheet.cell(row=row, column=column, value=safe_value(value))
    if font is not None:
        cell.font = font
    if fill is not None:
        cell.fill = fill
    if alignment is not None:
        cell.alignment = alignment
    if border is not None:
        cell.border = border
    if number_format is not None:
        cell.number_format = number_format
    return cell


def add_hyperlink(cell: Cell, target: str, label: str) -> None:
    cell.value = label
    cell.hyperlink = target
    cell.font = Font(name="Aptos", size=10, color=MAROON, underline="single")
    cell.alignment = Alignment(vertical="top", wrap_text=True)


def solid(color: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=color)


def apply_all_borders(sheet: Worksheet, min_row: int, max_row: int, min_col: int, max_col: int) -> None:
    border = Border(bottom=THIN_GRID)
    for row in sheet.iter_rows(
        min_row=min_row,
        max_row=max_row,
        min_col=min_col,
        max_col=max_col,
    ):
        for cell in row:
            cell.border = border


def configure_page(sheet: Worksheet, max_column: int, max_row: int, fit_width: int) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = fit_width
    sheet.page_setup.fitToHeight = 0
    sheet.print_area = f"A1:{get_column_letter(max_column)}{max_row}"
    sheet.print_title_rows = "1:6"
    sheet.page_margins = PageMargins(
        left=0.25,
        right=0.25,
        top=0.35,
        bottom=0.35,
        header=0.15,
        footer=0.15,
    )
    sheet.oddFooter.center.text = "NeuraFlow · Product Opportunity Radar"
    sheet.oddFooter.center.size = 8
    sheet.oddFooter.center.color = MID_GREY
    sheet.oddFooter.right.text = "Page &P of &N"
    sheet.oddFooter.right.size = 8
    sheet.oddFooter.right.color = MID_GREY


def add_brand_header(
    sheet: Worksheet,
    *,
    title: str,
    period: str,
    purpose: str,
    last_column: int,
    logo_path: Path,
    disclosure: str | None = None,
    reading_note: str | None = None,
) -> None:
    last_letter = get_column_letter(last_column)
    sheet.merge_cells(start_row=1, start_column=3, end_row=1, end_column=last_column)
    sheet.merge_cells(start_row=2, start_column=3, end_row=2, end_column=last_column)
    sheet.merge_cells(start_row=3, start_column=3, end_row=3, end_column=last_column)
    sheet.merge_cells(start_row=4, start_column=3, end_row=4, end_column=last_column)

    set_cell(
        sheet,
        1,
        3,
        title,
        font=Font(name="Aptos Display", size=22, bold=True, color=NAVY),
        fill=solid(CREAM),
        alignment=Alignment(vertical="center"),
    )
    set_cell(
        sheet,
        2,
        3,
        period,
        font=Font(name="Aptos", size=10.5, bold=True, color=MAROON),
        fill=solid(CREAM),
        alignment=Alignment(vertical="center"),
    )
    third_line = disclosure if disclosure is not None else purpose
    set_cell(
        sheet,
        3,
        3,
        third_line,
        font=Font(
            name="Aptos",
            size=9,
            italic=disclosure is not None,
            color=MID_GREY if disclosure is not None else CHARCOAL,
        ),
        fill=solid(CREAM),
        alignment=Alignment(vertical="center", wrap_text=True),
    )
    fourth_line = purpose if disclosure is not None else (reading_note or "")
    set_cell(
        sheet,
        4,
        3,
        fourth_line,
        font=Font(name="Aptos", size=9.5, color=CHARCOAL),
        fill=solid(CREAM),
        alignment=Alignment(vertical="center", wrap_text=True),
    )

    for row in range(1, 5):
        for column in range(1, last_column + 1):
            sheet.cell(row=row, column=column).fill = solid(CREAM)

    if logo_path is not None:
        # New frozen packets can supply verified owned image bytes instead of
        # rereading the mutable global logo. Existing path calls stay exact.
        logo = Image(logo_path if hasattr(logo_path, "read") else str(logo_path))
        logo.width = 115
        logo.height = 58
        logo.anchor = "A1"
        sheet.add_image(logo)

    sheet.row_dimensions[1].height = 31
    sheet.row_dimensions[2].height = 18
    sheet.row_dimensions[3].height = 18
    sheet.row_dimensions[4].height = 24
    for column in ("A", "B"):
        sheet.column_dimensions[column].width = max(sheet.column_dimensions[column].width or 0, 8)

    if last_letter:
        sheet.sheet_properties.tabColor = MAROON


def style_header_row(sheet: Worksheet, row: int, max_column: int) -> None:
    for column in range(1, max_column + 1):
        cell = sheet.cell(row=row, column=column)
        cell.fill = solid(NAVY)
        cell.font = Font(name="Aptos", size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = Border(bottom=MEDIUM_MAROON)
    sheet.row_dimensions[row].height = 42


def build_shortlist(
    workbook: Workbook,
    data: dict[str, Any],
    logo_path: Path,
) -> Worksheet:
    sheet = workbook.active
    sheet.title = "Shortlist"
    headers = [
        "Rank",
        "Product ID",
        "Product",
        "Category",
        "Market Signal",
        "Weekly Status",
        "Positive Core Signals",
        "Available Core Signals",
        "Why It Surfaced",
        "Retail Listing Band",
        "Quick-Commerce Context",
        "Compliance Check",
        "Main Warning",
        "Representative Image",
        "Evidence Row",
    ]
    report = data["report"]
    period = f"Week of 9–15 August 2026 · {report['report_type']}"
    add_brand_header(
        sheet,
        title=report["workbook_title"],
        period=period,
        purpose="Ranked weekly decision list. Use Evidence for source detail and Supplier Leads for investigation starting points.",
        disclosure=report["disclosure"],
        last_column=len(headers),
        logo_path=logo_path,
    )

    for column in range(1, len(headers) + 1):
        sheet.cell(row=5, column=column).fill = solid(MAROON)
    sheet.row_dimensions[5].height = 3

    for column, header in enumerate(headers, start=1):
        set_cell(sheet, 6, column, header)
    style_header_row(sheet, 6, len(headers))

    for offset, product in enumerate(data["products"], start=7):
        present = quick_presence(product)
        values = [
            product["rank"],
            product["product_id"],
            product["product"],
            product["category"],
            product["market_signal"],
            product["weekly_status"],
            product["positive_core_signals"],
            product["available_core_signals"],
            product["why_it_surfaced"],
            product["retail_listing_band"],
            f"Present in {present} of 9 sampled platform-city checks; context only.",
            product["compliance_check"],
            product["main_warning"],
            None,
            None,
        ]
        row_fill = WHITE if product["rank"] % 2 else LIGHT_GREY
        for column, value in enumerate(values, start=1):
            cell = set_cell(
                sheet,
                offset,
                column,
                value,
                font=Font(name="Aptos", size=9.5, color=CHARCOAL),
                fill=solid(row_fill),
                alignment=Alignment(
                    vertical="top",
                    wrap_text=True,
                    horizontal="center" if column in {1, 7, 8} else "left",
                ),
            )
            if column in {1, 7, 8}:
                cell.font = Font(name="Aptos", size=10, bold=True, color=NAVY)

        label_cell = sheet.cell(offset, 5)
        label_cell.fill = solid(LABEL_FILLS[product["market_signal"]])
        label_cell.font = Font(name="Aptos", size=9.5, bold=True, color=CHARCOAL)

        status_cell = sheet.cell(offset, 6)
        status_cell.fill = solid(STATUS_FILLS[product["weekly_status"]])
        status_cell.font = Font(name="Aptos", size=9.5, bold=True, color=CHARCOAL)

        compliance_cell = sheet.cell(offset, 12)
        if product["compliance_check"] == "Client Compliance Check Required":
            compliance_cell.fill = solid(PALE_ROSE)
            compliance_cell.font = Font(name="Aptos", size=9, bold=True, color=MAROON)

        add_hyperlink(
            sheet.cell(offset, 14),
            product["representative_image"],
            "Open image",
        )
        add_hyperlink(
            sheet.cell(offset, 15),
            f"#'Evidence'!A{product['rank'] + 6}",
            "Open evidence",
        )
        sheet.row_dimensions[offset].height = 76

    widths = {
        "A": 7,
        "B": 13,
        "C": 31,
        "D": 15,
        "E": 18,
        "F": 18,
        "G": 12,
        "H": 12,
        "I": 48,
        "J": 17,
        "K": 38,
        "L": 27,
        "M": 44,
        "N": 17,
        "O": 17,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    apply_all_borders(sheet, 7, 18, 1, len(headers))
    sheet.freeze_panes = "E7"
    sheet.auto_filter.ref = "A6:O18"
    sheet.sheet_view.zoomScale = 82
    sheet.sheet_properties.tabColor = MAROON
    sheet.print_title_cols = "A:D"
    configure_page(sheet, len(headers), 18, fit_width=1)
    return sheet


def build_evidence(
    workbook: Workbook,
    data: dict[str, Any],
    logo_path: Path,
) -> Worksheet:
    sheet = workbook.create_sheet("Evidence")
    headers = [
        "Rank",
        "Product ID",
        "Product",
        "Aliases",
        "Category",
        "Market Signal",
        "Weekly Status",
        "Positive Core Signals",
        "Available Core Signals",
        "Prior Label",
        "Prior Coverage",
        "Meta Available",
        "Meta Query",
        "Active Relevant Ads",
        "Distinct Advertisers",
        "Oldest Active Ad (days)",
        "Meta Result",
        "Meta Title Samples",
        "Meta Example Links",
        "Marketplace Available",
        "Amazon Evidence Type",
        "Amazon Title",
        "Amazon Position",
        "Amazon Link",
        "Flipkart Match",
        "Flipkart Title",
        "Flipkart Position",
        "Flipkart Link",
        "Marketplace Result",
        "Search Available",
        "Search Query",
        "Recent 4-Week Average",
        "Prior 4-Week Average",
        "Ratio",
        "Absolute Increase",
        "Related Rising Query",
        "Search Result",
        "Blinkit Sample Summary",
        "Zepto Sample Summary",
        "Swiggy Instamart Summary",
        "Presence Checks",
        "Sampled Checks",
        "Coverage Note",
        "Evidence Warning",
        "Compliance Check",
        "Main Warning",
        "Meta Collected At",
        "Marketplace Collected At",
        "Search Collected At",
    ]
    report = data["report"]
    add_brand_header(
        sheet,
        title="Evidence Ledger",
        period="Week of 9–15 August 2026 · 12 selected products",
        purpose="Source-level detail behind each label. Missing values remain missing; quick-commerce is context only.",
        reading_note="Use the outline controls to collapse or reveal secondary source detail and collection times.",
        last_column=len(headers),
        logo_path=logo_path,
    )

    groups = [
        ("Identity", 1, 5, NAVY, WHITE),
        ("Final Result", 6, 11, MAROON, WHITE),
        ("Meta Advertising", 12, 19, MUTED_NAVY, WHITE),
        ("Marketplace", 20, 29, NAVY, WHITE),
        ("Google Search", 30, 37, MUSTARD, CHARCOAL),
        ("Quick Commerce · Context Only", 38, 42, SAGE, CHARCOAL),
        ("Review", 43, 46, MAROON, WHITE),
        ("Provenance", 47, 49, MUTED_NAVY, WHITE),
    ]
    for title, start, end, color, font_color in groups:
        if start != end:
            sheet.merge_cells(start_row=5, start_column=start, end_row=5, end_column=end)
        set_cell(
            sheet,
            5,
            start,
            title,
            font=Font(name="Aptos", size=9, bold=True, color=font_color),
            fill=solid(color),
            alignment=Alignment(horizontal="left", vertical="center"),
        )
        for column in range(start, end + 1):
            sheet.cell(row=5, column=column).fill = solid(color)
            sheet.cell(row=5, column=column).border = Border(bottom=MEDIUM_MAROON)
    sheet.row_dimensions[5].height = 21

    for column, header in enumerate(headers, start=1):
        set_cell(sheet, 6, column, header)
    style_header_row(sheet, 6, len(headers))

    for row, product in enumerate(data["products"], start=7):
        meta = product["meta"]
        marketplace = product["marketplace"]
        search = product["search"]
        prior_coverage = (
            "—"
            if product["prior_available_core_signals"] is None
            else f"{product['prior_positive_core_signals']} positive / {product['prior_available_core_signals']} available"
        )
        values = [
            product["rank"],
            product["product_id"],
            product["product"],
            "; ".join(product["aliases"]),
            product["category"],
            product["market_signal"],
            product["weekly_status"],
            product["positive_core_signals"],
            product["available_core_signals"],
            product["prior_market_signal"],
            prior_coverage,
            bool_label(meta["available"]),
            meta["query"],
            meta["active_relevant_ads"],
            meta["distinct_advertisers"],
            meta["oldest_active_ad_days"],
            meta["result"],
            "\n".join(meta["source_title_samples"]),
            None,
            bool_label(marketplace["available"]),
            marketplace["amazon_evidence_type"],
            marketplace["amazon_title"],
            marketplace["amazon_position"],
            None,
            marketplace["flipkart_match"],
            marketplace["flipkart_title"] or "Not available",
            marketplace["flipkart_position"],
            None,
            marketplace["result"],
            bool_label(search["available"]),
            search["query"],
            search["recent_four_week_average"],
            search["prior_four_week_average"],
            search["ratio"],
            search["absolute_increase"],
            search["related_rising_query"] or "Not available",
            search["result"],
            quick_summary(product, "Blinkit"),
            quick_summary(product, "Zepto"),
            quick_summary(product, "Swiggy Instamart"),
            quick_presence(product),
            9,
            product["coverage_note"],
            product["evidence_warning"],
            product["compliance_check"],
            product["main_warning"],
            format_timestamp(meta["collected_at"]),
            format_timestamp(marketplace["collected_at"]),
            format_timestamp(search["collected_at"]),
        ]
        base_fill = WHITE if product["rank"] % 2 else LIGHT_GREY
        for column, value in enumerate(values, start=1):
            numeric_center = column in {
                1,
                8,
                9,
                14,
                15,
                16,
                23,
                27,
                32,
                33,
                34,
                35,
                41,
                42,
            }
            set_cell(
                sheet,
                row,
                column,
                value,
                font=Font(name="Aptos", size=9, color=CHARCOAL),
                fill=solid(base_fill),
                alignment=Alignment(
                    vertical="top",
                    wrap_text=True,
                    horizontal="center" if numeric_center else "left",
                ),
                number_format="0.00" if column == 34 else None,
            )

        sheet.cell(row, 6).fill = solid(LABEL_FILLS[product["market_signal"]])
        sheet.cell(row, 6).font = Font(name="Aptos", size=9, bold=True, color=CHARCOAL)
        sheet.cell(row, 7).fill = solid(STATUS_FILLS[product["weekly_status"]])
        sheet.cell(row, 7).font = Font(name="Aptos", size=9, bold=True, color=CHARCOAL)
        if product["compliance_check"] == "Client Compliance Check Required":
            sheet.cell(row, 45).fill = solid(PALE_ROSE)
            sheet.cell(row, 45).font = Font(name="Aptos", size=9, bold=True, color=MAROON)

        add_hyperlink(
            sheet.cell(row, 19),
            meta["example_links"][0],
            f"Open {len(meta['example_links'])} example" + ("s" if len(meta["example_links"]) > 1 else ""),
        )
        add_hyperlink(
            sheet.cell(row, 24),
            marketplace["amazon_link"],
            "Open Amazon",
        )
        add_hyperlink(
            sheet.cell(row, 28),
            marketplace["flipkart_link"],
            "Open Flipkart",
        )
        sheet.row_dimensions[row].height = 86

    widths = {
        "A": 7,
        "B": 13,
        "C": 30,
        "D": 34,
        "E": 14,
        "F": 18,
        "G": 18,
        "H": 11,
        "I": 11,
        "J": 17,
        "K": 21,
        "L": 14,
        "M": 28,
        "N": 11,
        "O": 11,
        "P": 13,
        "Q": 14,
        "R": 42,
        "S": 17,
        "T": 14,
        "U": 19,
        "V": 40,
        "W": 11,
        "X": 16,
        "Y": 17,
        "Z": 38,
        "AA": 11,
        "AB": 16,
        "AC": 16,
        "AD": 14,
        "AE": 28,
        "AF": 13,
        "AG": 13,
        "AH": 10,
        "AI": 12,
        "AJ": 28,
        "AK": 14,
        "AL": 39,
        "AM": 39,
        "AN": 43,
        "AO": 11,
        "AP": 10,
        "AQ": 44,
        "AR": 48,
        "AS": 28,
        "AT": 48,
        "AU": 21,
        "AV": 21,
        "AW": 21,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    for start, end in [("R", "S"), ("V", "X"), ("Z", "AB"), ("AJ", "AJ")]:
        sheet.column_dimensions.group(start, end, hidden=False, outline_level=1)
    sheet.column_dimensions.group("AU", "AW", hidden=True, outline_level=1)
    sheet.sheet_properties.outlinePr.summaryRight = True

    apply_all_borders(sheet, 7, 18, 1, len(headers))
    sheet.freeze_panes = "F7"
    sheet.auto_filter.ref = "A6:AW18"
    sheet.sheet_view.zoomScale = 58
    sheet.sheet_properties.tabColor = NAVY
    sheet.print_title_cols = "A:E"
    configure_page(sheet, len(headers), 18, fit_width=3)
    return sheet


def build_supplier_leads(
    workbook: Workbook,
    data: dict[str, Any],
    logo_path: Path,
) -> Worksheet:
    sheet = workbook.create_sheet("Supplier Leads")
    headers = [
        "Product Rank",
        "Product ID",
        "Product",
        "Lead ID",
        "Supplier",
        "City & State",
        "Supplier Type",
        "Visible Price",
        "Visible MOQ",
        "Evidence Clarity",
        "Evidence Note",
        "Supplier Page",
        "Source Domain",
        "URL Status",
        "Lead Limitation",
    ]
    add_brand_header(
        sheet,
        title="Supplier Lead Ledger",
        period="Week of 9–15 August 2026 · 36 investigation starting points",
        purpose="Possible leads only. Visible price is not landed cost; contact, stock, terms and suitability require client investigation.",
        reading_note="Missing commercial fields are left visibly missing rather than inferred.",
        last_column=len(headers),
        logo_path=logo_path,
    )

    for column in range(1, len(headers) + 1):
        sheet.cell(row=5, column=column).fill = solid(MAROON)
    sheet.row_dimensions[5].height = 3

    for column, header in enumerate(headers, start=1):
        set_cell(sheet, 6, column, header)
    style_header_row(sheet, 6, len(headers))

    products = {product["product_id"]: product for product in data["products"]}
    sorted_leads = sorted(
        data["supplier_leads"],
        key=lambda lead: (lead["product_rank"], lead["lead_id"]),
    )
    previous_rank: int | None = None
    for row, lead in enumerate(sorted_leads, start=7):
        product = products[lead["product_id"]]
        values = [
            lead["product_rank"],
            lead["product_id"],
            product["product"],
            lead["lead_id"],
            lead["supplier"],
            lead["city_state"],
            lead["supplier_type"],
            lead["visible_price"] if lead["visible_price"] is not None else "Not shown",
            lead["visible_moq"] if lead["visible_moq"] is not None else "Not shown",
            lead["evidence_clarity"],
            lead["evidence_note"],
            None,
            lead["source_domain"],
            lead["url_status"],
            lead["lead_limitation"],
        ]
        rank_fill = WHITE if lead["product_rank"] % 2 else LIGHT_GREY
        for column, value in enumerate(values, start=1):
            set_cell(
                sheet,
                row,
                column,
                value,
                font=Font(name="Aptos", size=9.25, color=CHARCOAL),
                fill=solid(rank_fill),
                alignment=Alignment(
                    vertical="top",
                    wrap_text=True,
                    horizontal="center" if column in {1, 4} else "left",
                ),
            )

        clarity_cell = sheet.cell(row, 10)
        clarity_cell.fill = solid(CLARITY_FILLS[lead["evidence_clarity"]])
        clarity_cell.font = Font(name="Aptos", size=9.25, bold=True, color=CHARCOAL)

        status_cell = sheet.cell(row, 14)
        if lead["url_status"] == "Reachable":
            status_cell.fill = solid(PALE_SAGE)
        elif lead["url_status"] in {"Redirected", "Unchecked"}:
            status_cell.fill = solid(PALE_MUSTARD)
        else:
            status_cell.fill = solid(PALE_ROSE)
        status_cell.font = Font(
            name="Aptos",
            size=9.25,
            bold=True,
            color=MAROON if lead["url_status"] in {"Blocked", "Unreachable"} else CHARCOAL,
        )

        add_hyperlink(
            sheet.cell(row, 12),
            lead["supplier_page"],
            "Open supplier page",
        )
        sheet.row_dimensions[row].height = 64

        if previous_rank is not None and previous_rank != lead["product_rank"]:
            for column in range(1, len(headers) + 1):
                sheet.cell(row, column).border = Border(top=MEDIUM_MAROON)
        previous_rank = lead["product_rank"]

    widths = {
        "A": 10,
        "B": 13,
        "C": 31,
        "D": 11,
        "E": 28,
        "F": 24,
        "G": 17,
        "H": 18,
        "I": 18,
        "J": 16,
        "K": 46,
        "L": 20,
        "M": 17,
        "N": 15,
        "O": 50,
    }
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width

    apply_all_borders(sheet, 7, 42, 1, len(headers))
    previous_rank = None
    for row in range(7, 43):
        rank = sheet.cell(row, 1).value
        if previous_rank is not None and rank != previous_rank:
            for column in range(1, len(headers) + 1):
                existing = sheet.cell(row, column).border
                sheet.cell(row, column).border = Border(
                    left=existing.left,
                    right=existing.right,
                    top=MEDIUM_MAROON,
                    bottom=existing.bottom,
                )
        previous_rank = rank

    sheet.freeze_panes = "E7"
    sheet.auto_filter.ref = "A6:O42"
    sheet.sheet_view.zoomScale = 78
    sheet.sheet_properties.tabColor = MUSTARD
    sheet.print_title_cols = "A:D"
    configure_page(sheet, len(headers), 42, fit_width=2)
    return sheet


def generate_workbook(
    data_path: Path,
    output_path: Path,
    logo_path: Path,
) -> None:
    with data_path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)

    workbook = Workbook()
    workbook.properties.creator = "NeuraFlow"
    workbook.properties.lastModifiedBy = "NeuraFlow"
    workbook.properties.title = data["report"]["workbook_title"]
    workbook.properties.subject = "Illustrative weekly product-opportunity evidence workbook"
    workbook.properties.description = (
        "Client-preview workbook generated from one canonical manufactured dataset."
    )
    workbook.properties.keywords = (
        "product research, opportunity radar, evidence, supplier leads, NeuraFlow"
    )
    workbook.properties.category = "Client Preview"
    workbook.iso_dates = True

    build_shortlist(workbook, data, logo_path)
    build_evidence(workbook, data, logo_path)
    build_supplier_leads(workbook, data, logo_path)

    workbook.active = 0
    for sheet in workbook.worksheets:
        sheet.sheet_view.selection[0].activeCell = "A1"
        sheet.sheet_view.selection[0].sqref = "A1"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--data",
        type=Path,
        default=Path("preview/sample-data/client_preview_dataset.json"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path(
            "outputs/client-preview/"
            "NeuraFlow_Product_Opportunity_Evidence_Illustrative_2026-08-15.xlsx"
        ),
    )
    parser.add_argument(
        "--logo",
        type=Path,
        default=Path("media/NeuraFlow/long_maroon_path_transparent_1024.png"),
    )
    args = parser.parse_args()

    generate_workbook(args.data, args.output, args.logo)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
