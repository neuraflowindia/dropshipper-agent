#!/usr/bin/env python3
"""Validate the manufactured client-preview dataset and optional workbook."""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse


LABEL_STRENGTH = {"Weak": 0, "Worth Watching": 1, "Promising": 2, "Strong": 3}
ALLOWED_STATUSES = {
    "New",
    "Returning",
    "Strengthened",
    "Weakened",
    "Coverage Changed",
}
ALLOWED_URL_STATUSES = {
    "Reachable",
    "Redirected",
    "Unchecked",
    "Blocked",
    "Unreachable",
}
ALLOWED_SUPPLIER_TYPES = {
    "Manufacturer",
    "Wholesaler",
    "Importer",
    "Trader",
    "Unclear",
}
ALLOWED_EVIDENCE_CLARITY = {"Clear", "Partial", "Uncertain"}
EXPECTED_DISCLOSURE = (
    "Illustrative preview using sample data; not a live market-research report."
)
EXPECTED_SUPPLIER_STATUS_COUNTS = Counter(
    {
        "Reachable": 21,
        "Redirected": 4,
        "Unchecked": 5,
        "Blocked": 3,
        "Unreachable": 3,
    }
)
EXPECTED_SUPPLIER_CLARITY_COUNTS = Counter(
    {"Clear": 23, "Partial": 9, "Uncertain": 4}
)
EXPECTED_SUPPLIER_TYPE_COUNTS = Counter(
    {
        "Manufacturer": 11,
        "Wholesaler": 13,
        "Trader": 7,
        "Importer": 3,
        "Unclear": 2,
    }
)
EXPECTED_QUICK_COUNTS = {
    1: 7,
    2: 1,
    3: 3,
    4: 6,
    5: 6,
    6: 2,
    7: 4,
    8: 5,
    9: 4,
    10: 5,
    11: 3,
    12: 4,
}
FORBIDDEN_SUPPLIER_CLAIMS = {"verified", "approved", "trusted", "recommended"}


class Validation:
    def __init__(self) -> None:
        self.errors: list[str] = []
        self.checks = 0

    def require(self, condition: bool, message: str) -> None:
        self.checks += 1
        if not condition:
            self.errors.append(message)


def is_positive_meta(meta: dict) -> bool:
    if not meta["available"]:
        return False
    activity = (
        meta["distinct_advertisers"] >= 3
        or meta["active_relevant_ads"] >= 5
    )
    return activity and meta["oldest_active_ad_days"] >= 14


def is_positive_marketplace(marketplace: dict) -> bool:
    if not marketplace["available"]:
        return False
    if marketplace["amazon_evidence_type"] == "Bestseller":
        return True
    both_relevant = (
        marketplace["amazon_position"] is not None
        and marketplace["flipkart_match"] == "Relevant"
        and marketplace["flipkart_position"] is not None
    )
    top_ten = (
        marketplace["amazon_position"] is not None
        and marketplace["amazon_position"] <= 10
    ) or (
        marketplace["flipkart_position"] is not None
        and marketplace["flipkart_position"] <= 10
    )
    return both_relevant and top_ten


def is_positive_search(search: dict) -> bool:
    if not search["available"]:
        return False
    numeric_growth = (
        search["ratio"] is not None
        and search["absolute_increase"] is not None
        and search["ratio"] >= 1.15
        and search["absolute_increase"] >= 5
    )
    return numeric_growth


def expected_weekly_status(product: dict) -> str:
    prior_label = product["prior_market_signal"]
    if prior_label == "Not shortlisted":
        return "New"
    if product["prior_available_core_signals"] != product["available_core_signals"]:
        return "Coverage Changed"
    current_strength = LABEL_STRENGTH[product["market_signal"]]
    prior_strength = LABEL_STRENGTH[prior_label]
    if current_strength > prior_strength:
        return "Strengthened"
    if current_strength < prior_strength:
        return "Weakened"
    return "Returning"


def bool_label(value: bool) -> str:
    return "Available" if value else "Not available"


def format_timestamp(value: str | None) -> str:
    if not value:
        return "Not available"
    return datetime.fromisoformat(value).strftime("%d %b %Y, %H:%M")


def quick_presence(product: dict) -> int:
    return sum(
        state == "Present"
        for city_values in product["quick_commerce"].values()
        for state in city_values.values()
    )


def quick_summary(product: dict, platform: str) -> str:
    return "; ".join(
        f"{city}: {product['quick_commerce'][city][platform]}"
        for city in ("Mumbai", "Delhi", "Bengaluru")
    )


def validate_dataset(path: Path) -> tuple[Validation, dict]:
    validation = Validation()
    with path.open("r", encoding="utf-8") as handle:
        data = json.load(handle)

    report = data["report"]
    products = data["products"]
    suppliers = data["supplier_leads"]

    validation.require(
        report["disclosure"] == EXPECTED_DISCLOSURE,
        "The manufactured-preview disclosure differs from the approved line.",
    )
    validation.require(report["disclosure_sheet"] == "Shortlist", "Disclosure must be assigned to Shortlist only.")
    validation.require(len(products) == 12, f"Expected 12 products, found {len(products)}.")
    validation.require(len(suppliers) == 36, f"Expected 36 supplier leads, found {len(suppliers)}.")

    ranks = [product["rank"] for product in products]
    product_ids = [product["product_id"] for product in products]
    validation.require(ranks == list(range(1, 13)), f"Ranks are not 1–12 in order: {ranks}.")
    validation.require(len(set(product_ids)) == 12, "Product IDs must be unique.")

    category_counts = Counter(product["category"] for product in products)
    validation.require(
        category_counts == Counter({"Kitchen": 3, "Home Decor": 3, "Pet": 3, "Baby": 3}),
        f"Unexpected category balance: {dict(category_counts)}.",
    )

    label_counts = Counter(product["market_signal"] for product in products)
    validation.require(
        label_counts == Counter({"Strong": 4, "Promising": 6, "Worth Watching": 2}),
        f"Unexpected label distribution: {dict(label_counts)}.",
    )

    status_counts = Counter(product["weekly_status"] for product in products)
    validation.require(
        status_counts
        == Counter(
            {
                "New": 5,
                "Returning": 3,
                "Strengthened": 2,
                "Weakened": 1,
                "Coverage Changed": 1,
            }
        ),
        f"Unexpected weekly-status distribution: {dict(status_counts)}.",
    )

    validation.require(
        set(status_counts).issubset(ALLOWED_STATUSES),
        f"Unknown weekly status: {set(status_counts) - ALLOWED_STATUSES}.",
    )

    total_positive = 0
    total_available = 0
    total_quick_present = 0
    prior_strength = 99
    products_by_id: dict[str, dict] = {}

    for product in products:
        products_by_id[product["product_id"]] = product
        rank = product["rank"]
        signals = [
            ("Meta", product["meta"], is_positive_meta(product["meta"])),
            (
                "Marketplace",
                product["marketplace"],
                is_positive_marketplace(product["marketplace"]),
            ),
            ("Search", product["search"], is_positive_search(product["search"])),
        ]
        positive = sum(expected for _, _, expected in signals)
        available = sum(bool(payload["available"]) for _, payload, _ in signals)
        total_positive += positive
        total_available += available

        validation.require(
            positive == product["positive_core_signals"],
            f"Rank {rank}: calculated {positive} positive signals, stored {product['positive_core_signals']}.",
        )
        validation.require(
            available == product["available_core_signals"],
            f"Rank {rank}: calculated {available} available signals, stored {product['available_core_signals']}.",
        )
        validation.require(
            LABEL_STRENGTH[product["market_signal"]] == positive,
            f"Rank {rank}: label {product['market_signal']} does not match {positive} positives.",
        )

        for signal_name, payload, expected_positive in signals:
            expected_result = "Missing" if not payload["available"] else ("Positive" if expected_positive else "Not positive")
            validation.require(
                payload["result"] == expected_result,
                f"Rank {rank} {signal_name}: stored result {payload['result']!r}, expected {expected_result!r}.",
            )

        validation.require(
            product["weekly_status"] == expected_weekly_status(product),
            f"Rank {rank}: weekly status {product['weekly_status']} does not reconcile with prior coverage and label.",
        )

        current_strength = LABEL_STRENGTH[product["market_signal"]]
        validation.require(
            current_strength <= prior_strength,
            f"Rank {rank}: a stronger label appears below a weaker label.",
        )
        prior_strength = current_strength

        quick_states = [
            state
            for city_values in product["quick_commerce"].values()
            for state in city_values.values()
        ]
        validation.require(
            len(quick_states) == 9,
            f"Rank {rank}: expected 9 quick-commerce checks, found {len(quick_states)}.",
        )
        validation.require(
            set(quick_states).issubset({"Present", "No match"}),
            f"Rank {rank}: unexpected quick-commerce state.",
        )
        quick_present = quick_states.count("Present")
        total_quick_present += quick_present
        validation.require(
            quick_present == EXPECTED_QUICK_COUNTS[rank],
            f"Rank {rank}: expected {EXPECTED_QUICK_COUNTS[rank]} quick-commerce presences, found {quick_present}.",
        )

        for url in [
            product["representative_image"],
            *product["meta"]["example_links"],
            product["marketplace"]["amazon_link"],
            product["marketplace"]["flipkart_link"],
        ]:
            validation.require(
                urlparse(url).hostname == "example.com",
                f"Rank {rank}: non-reserved preview URL {url}.",
            )

    validation.require(total_positive == 26, f"Expected 26 positive signal outcomes, found {total_positive}.")
    validation.require(total_available == 35, f"Expected 35 available signal assessments, found {total_available}.")
    validation.require(total_quick_present == 50, f"Expected 50 quick-commerce presences, found {total_quick_present}.")

    leads_by_product: defaultdict[str, list[dict]] = defaultdict(list)
    lead_ids: set[str] = set()
    supplier_pages: set[str] = set()
    for lead in suppliers:
        lead_id = lead["lead_id"]
        validation.require(lead_id not in lead_ids, f"Duplicate supplier lead ID {lead_id}.")
        lead_ids.add(lead_id)
        leads_by_product[lead["product_id"]].append(lead)
        product = products_by_id.get(lead["product_id"])
        validation.require(product is not None, f"Lead {lead_id}: unknown product ID.")
        if product is not None:
            validation.require(
                lead["product_rank"] == product["rank"],
                f"Lead {lead_id}: product rank does not match {lead['product_id']}.",
            )
        validation.require(
            lead["url_status"] in ALLOWED_URL_STATUSES,
            f"Lead {lead_id}: invalid URL status {lead['url_status']}.",
        )
        validation.require(
            lead["supplier_type"] in ALLOWED_SUPPLIER_TYPES,
            f"Lead {lead_id}: invalid supplier type {lead['supplier_type']}.",
        )
        validation.require(
            lead["evidence_clarity"] in ALLOWED_EVIDENCE_CLARITY,
            f"Lead {lead_id}: invalid evidence clarity {lead['evidence_clarity']}.",
        )
        validation.require(
            bool(lead["supplier"].strip()),
            f"Lead {lead_id}: supplier name is empty.",
        )
        validation.require(
            not lead["supplier"].startswith("Sample Lead"),
            f"Lead {lead_id}: supplier name still uses a placeholder pattern.",
        )
        validation.require(
            urlparse(lead["supplier_page"]).hostname == "example.com",
            f"Lead {lead_id}: supplier URL is not under example.com.",
        )
        validation.require(
            lead["source_domain"] == "example.com",
            f"Lead {lead_id}: source domain differs from the reserved preview domain.",
        )
        validation.require(
            lead["supplier_page"] not in supplier_pages,
            f"Lead {lead_id}: supplier URL is duplicated.",
        )
        supplier_pages.add(lead["supplier_page"])
        combined_text = " ".join(
            str(lead[field] or "")
            for field in ["supplier", "evidence_note", "lead_limitation"]
        ).lower()
        validation.require(
            not any(claim in combined_text for claim in FORBIDDEN_SUPPLIER_CLAIMS),
            f"Lead {lead_id}: contains forbidden supplier-assurance wording.",
        )

    for product_id in product_ids:
        validation.require(
            len(leads_by_product[product_id]) == 3,
            f"{product_id}: expected 3 supplier leads, found {len(leads_by_product[product_id])}.",
        )

    validation.require(
        Counter(lead["url_status"] for lead in suppliers)
        == EXPECTED_SUPPLIER_STATUS_COUNTS,
        "Supplier URL-status distribution differs from the approved manufactured sample.",
    )
    validation.require(
        Counter(lead["evidence_clarity"] for lead in suppliers)
        == EXPECTED_SUPPLIER_CLARITY_COUNTS,
        "Supplier evidence-clarity distribution differs from the approved manufactured sample.",
    )
    validation.require(
        Counter(lead["supplier_type"] for lead in suppliers)
        == EXPECTED_SUPPLIER_TYPE_COUNTS,
        "Supplier-type distribution differs from the approved manufactured sample.",
    )
    validation.require(
        sum(lead["visible_price"] is None for lead in suppliers) == 6,
        "The manufactured sample must retain exactly six visibly missing supplier prices.",
    )
    validation.require(
        sum(lead["visible_moq"] is None for lead in suppliers) == 4,
        "The manufactured sample must retain exactly four visibly missing supplier MOQs.",
    )

    return validation, data


def validate_workbook(path: Path, data: dict, validation: Validation) -> None:
    try:
        from openpyxl import load_workbook
    except ImportError:
        validation.errors.append("openpyxl is required to validate the workbook.")
        return

    workbook = load_workbook(path, data_only=False, keep_links=False)
    validation.require(
        workbook.sheetnames == ["Shortlist", "Evidence", "Supplier Leads"],
        f"Unexpected sheet order: {workbook.sheetnames}.",
    )

    expected_headers = {
        "Shortlist": [
            "Rank",
            "Product ID",
            "Product",
            "Category",
            "Market Signal",
            "Weekly Status",
            "Positive Core Signals",
            "Available Core Signals",
            "Why It Surfaced",
            "Retail Listing Band",
            "Quick-Commerce Context",
            "Compliance Check",
            "Main Warning",
            "Representative Image",
            "Evidence Row",
        ],
        "Evidence": [
            "Rank",
            "Product ID",
            "Product",
            "Aliases",
            "Category",
            "Market Signal",
            "Weekly Status",
            "Positive Core Signals",
            "Available Core Signals",
            "Prior Label",
            "Prior Coverage",
            "Meta Available",
            "Meta Query",
            "Active Relevant Ads",
            "Distinct Advertisers",
            "Oldest Active Ad (days)",
            "Meta Result",
            "Meta Title Samples",
            "Meta Example Links",
            "Marketplace Available",
            "Amazon Evidence Type",
            "Amazon Title",
            "Amazon Position",
            "Amazon Link",
            "Flipkart Match",
            "Flipkart Title",
            "Flipkart Position",
            "Flipkart Link",
            "Marketplace Result",
            "Search Available",
            "Search Query",
            "Recent 4-Week Average",
            "Prior 4-Week Average",
            "Ratio",
            "Absolute Increase",
            "Related Rising Query",
            "Search Result",
            "Blinkit Sample Summary",
            "Zepto Sample Summary",
            "Swiggy Instamart Summary",
            "Presence Checks",
            "Sampled Checks",
            "Coverage Note",
            "Evidence Warning",
            "Compliance Check",
            "Main Warning",
            "Meta Collected At",
            "Marketplace Collected At",
            "Search Collected At",
        ],
        "Supplier Leads": [
            "Product Rank",
            "Product ID",
            "Product",
            "Lead ID",
            "Supplier",
            "City & State",
            "Supplier Type",
            "Visible Price",
            "Visible MOQ",
            "Evidence Clarity",
            "Evidence Note",
            "Supplier Page",
            "Source Domain",
            "URL Status",
            "Lead Limitation",
        ],
    }
    expected_specs = {
        "Shortlist": {
            "records": 12,
            "max_row": 18,
            "max_column": 15,
            "freeze": "E7",
            "filter": "A6:O18",
        },
        "Evidence": {
            "records": 12,
            "max_row": 18,
            "max_column": 49,
            "freeze": "F7",
            "filter": "A6:AW18",
        },
        "Supplier Leads": {
            "records": 36,
            "max_row": 42,
            "max_column": 15,
            "freeze": "E7",
            "filter": "A6:O42",
        },
    }
    for sheet_name, spec in expected_specs.items():
        sheet = workbook[sheet_name]
        validation.require(
            sheet.max_row == spec["max_row"],
            f"{sheet_name}: expected max row {spec['max_row']}, found {sheet.max_row}.",
        )
        validation.require(
            sheet.max_column == spec["max_column"],
            f"{sheet_name}: expected {spec['max_column']} columns, found {sheet.max_column}.",
        )
        validation.require(
            str(sheet.freeze_panes) == spec["freeze"],
            f"{sheet_name}: expected freeze pane {spec['freeze']}, found {sheet.freeze_panes}.",
        )
        validation.require(
            sheet.auto_filter.ref == spec["filter"],
            f"{sheet_name}: expected filter {spec['filter']}, found {sheet.auto_filter.ref}.",
        )
        validation.require(
            len(sheet._images) == 1,
            f"{sheet_name}: expected exactly one NeuraFlow logo, found {len(sheet._images)}.",
        )
        actual_headers = [
            sheet.cell(row=6, column=column).value
            for column in range(1, spec["max_column"] + 1)
        ]
        validation.require(
            actual_headers == expected_headers[sheet_name],
            f"{sheet_name}: header contract differs.",
        )

    products = sorted(data["products"], key=lambda product: product["rank"])
    shortlist = workbook["Shortlist"]
    evidence = workbook["Evidence"]
    for product in products:
        row = product["rank"] + 6
        expected_shortlist = [
            product["rank"],
            product["product_id"],
            product["product"],
            product["category"],
            product["market_signal"],
            product["weekly_status"],
            product["positive_core_signals"],
            product["available_core_signals"],
            product["why_it_surfaced"],
            product["retail_listing_band"],
            f"Present in {quick_presence(product)} of 9 sampled platform-city checks; context only.",
            product["compliance_check"],
            product["main_warning"],
            "Open image",
            "Open evidence",
        ]
        validation.require(
            [shortlist.cell(row, column).value for column in range(1, 16)]
            == expected_shortlist,
            f"Shortlist row {row}: one or more displayed fields differ from source.",
        )
        validation.require(
            shortlist.cell(row, 14).hyperlink.target == product["representative_image"],
            f"Shortlist row {row}: representative-image link differs from source.",
        )
        validation.require(
            shortlist.cell(row, 15).hyperlink.target == f"#'Evidence'!A{row}",
            f"Shortlist row {row}: evidence-row link is incorrect.",
        )

        meta = product["meta"]
        marketplace = product["marketplace"]
        search = product["search"]
        prior_coverage = (
            "—"
            if product["prior_available_core_signals"] is None
            else f"{product['prior_positive_core_signals']} positive / {product['prior_available_core_signals']} available"
        )
        expected_evidence = [
            product["rank"],
            product["product_id"],
            product["product"],
            "; ".join(product["aliases"]),
            product["category"],
            product["market_signal"],
            product["weekly_status"],
            product["positive_core_signals"],
            product["available_core_signals"],
            product["prior_market_signal"],
            prior_coverage,
            bool_label(meta["available"]),
            meta["query"],
            meta["active_relevant_ads"],
            meta["distinct_advertisers"],
            meta["oldest_active_ad_days"],
            meta["result"],
            "\n".join(meta["source_title_samples"]),
            f"Open {len(meta['example_links'])} example" + ("s" if len(meta["example_links"]) > 1 else ""),
            bool_label(marketplace["available"]),
            marketplace["amazon_evidence_type"],
            marketplace["amazon_title"],
            marketplace["amazon_position"],
            "Open Amazon",
            marketplace["flipkart_match"],
            marketplace["flipkart_title"] or "Not available",
            marketplace["flipkart_position"],
            "Open Flipkart",
            marketplace["result"],
            bool_label(search["available"]),
            search["query"],
            search["recent_four_week_average"],
            search["prior_four_week_average"],
            search["ratio"],
            search["absolute_increase"],
            search["related_rising_query"] or "Not available",
            search["result"],
            quick_summary(product, "Blinkit"),
            quick_summary(product, "Zepto"),
            quick_summary(product, "Swiggy Instamart"),
            quick_presence(product),
            9,
            product["coverage_note"],
            product["evidence_warning"],
            product["compliance_check"],
            product["main_warning"],
            format_timestamp(meta["collected_at"]),
            format_timestamp(marketplace["collected_at"]),
            format_timestamp(search["collected_at"]),
        ]
        validation.require(
            [evidence.cell(row, column).value for column in range(1, 50)]
            == expected_evidence,
            f"Evidence row {row}: one or more displayed fields differ from source.",
        )
        for column, target in [
            (19, meta["example_links"][0]),
            (24, marketplace["amazon_link"]),
            (28, marketplace["flipkart_link"]),
        ]:
            validation.require(
                evidence.cell(row, column).hyperlink.target == target,
                f"Evidence row {row}: hyperlink in column {column} differs from source.",
            )

    products_by_id = {product["product_id"]: product for product in products}
    expected_leads = sorted(
        data["supplier_leads"],
        key=lambda lead: (lead["product_rank"], lead["lead_id"]),
    )
    supplier_sheet = workbook["Supplier Leads"]
    for row, lead in enumerate(expected_leads, start=7):
        product = products_by_id[lead["product_id"]]
        expected_supplier_row = [
            lead["product_rank"],
            lead["product_id"],
            product["product"],
            lead["lead_id"],
            lead["supplier"],
            lead["city_state"],
            lead["supplier_type"],
            lead["visible_price"] if lead["visible_price"] is not None else "Not shown",
            lead["visible_moq"] if lead["visible_moq"] is not None else "Not shown",
            lead["evidence_clarity"],
            lead["evidence_note"],
            "Open supplier page",
            lead["source_domain"],
            lead["url_status"],
            lead["lead_limitation"],
        ]
        validation.require(
            [supplier_sheet.cell(row, column).value for column in range(1, 16)]
            == expected_supplier_row,
            f"Supplier Leads row {row}: one or more displayed fields differ from source.",
        )
        validation.require(
            supplier_sheet.cell(row, 12).hyperlink.target == lead["supplier_page"],
            f"Supplier Leads row {row}: supplier-page link differs from source.",
        )

    validation.require(
        evidence.column_dimensions["AU"].hidden
        and evidence.column_dimensions["AU"].outlineLevel == 1,
        "Evidence: provenance columns must be grouped and collapsed.",
    )
    validation.require(
        all(evidence.column_dimensions[column].outlineLevel == 1 for column in ["R", "V", "Z", "AJ"]),
        "Evidence: secondary source-detail groups are missing.",
    )

    disclosure = data["report"]["disclosure"]
    disclosure_hits: list[str] = []
    formula_cells: list[str] = []
    unsafe_prefix_cells: list[str] = []
    external_hyperlinks: list[str] = []
    internal_hyperlinks: list[str] = []
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value == disclosure:
                    disclosure_hits.append(f"{sheet.title}!{cell.coordinate}")
                if cell.data_type == "f" or (
                    isinstance(cell.value, str) and cell.value.startswith("=")
                ):
                    formula_cells.append(f"{sheet.title}!{cell.coordinate}")
                if (
                    isinstance(cell.value, str)
                    and cell.value.startswith(("+", "-", "@"))
                    and not cell.value.startswith("'")
                ):
                    unsafe_prefix_cells.append(f"{sheet.title}!{cell.coordinate}")
                if cell.hyperlink:
                    target = cell.hyperlink.target
                    if target.startswith("#"):
                        internal_hyperlinks.append(target)
                    else:
                        external_hyperlinks.append(target)

    validation.require(
        disclosure_hits == ["Shortlist!C3"],
        f"Disclosure must appear only in Shortlist!C3, found {disclosure_hits}.",
    )
    validation.require(not formula_cells, f"Workbook contains formulas: {formula_cells[:5]}.")
    validation.require(
        not unsafe_prefix_cells,
        f"Workbook contains unsafe formula-prefix strings: {unsafe_prefix_cells[:5]}.",
    )
    validation.require(
        len(internal_hyperlinks) == 12,
        f"Expected 12 internal evidence hyperlinks, found {len(internal_hyperlinks)}.",
    )
    validation.require(
        len(external_hyperlinks) == 84,
        f"Expected 84 external preview hyperlinks, found {len(external_hyperlinks)}.",
    )
    validation.require(
        all(urlparse(target).hostname == "example.com" for target in external_hyperlinks),
        "Workbook contains an external preview hyperlink outside example.com.",
    )
    validation.require(workbook.vba_archive is None, "Workbook unexpectedly contains VBA.")
    validation.require(
        not getattr(workbook, "_external_links", []),
        "Workbook unexpectedly contains external workbook links.",
    )

    with zipfile.ZipFile(path) as archive:
        members = archive.namelist()
        validation.require("[Content_Types].xml" in members, "XLSX package is missing [Content_Types].xml.")
        validation.require("xl/workbook.xml" in members, "XLSX package is missing xl/workbook.xml.")
        validation.require(
            not any(member.endswith("vbaProject.bin") for member in members),
            "XLSX package contains a VBA project.",
        )
        validation.require(
            not any(member.startswith("xl/externalLinks/") for member in members),
            "XLSX package contains external workbook links.",
        )
    validation.require(path.stat().st_size < 5 * 1024 * 1024, "Workbook exceeds the 5 MB preview target.")


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--data",
        type=Path,
        default=Path("preview/sample-data/client_preview_dataset.json"),
    )
    parser.add_argument("--workbook", type=Path)
    args = parser.parse_args()

    validation, data = validate_dataset(args.data)
    if args.workbook:
        validate_workbook(args.workbook, data, validation)

    if validation.errors:
        print(f"FAILED: {len(validation.errors)} issue(s) across {validation.checks} checks")
        for error in validation.errors:
            print(f"- {error}")
        return 1

    subject = "dataset and workbook" if args.workbook else "dataset"
    print(f"PASS: {subject} passed {validation.checks} checks")
    return 0


if __name__ == "__main__":
    sys.exit(main())
