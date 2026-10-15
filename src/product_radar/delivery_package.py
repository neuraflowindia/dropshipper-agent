"""Bind checked local outputs into one immutable, unapproved delivery packet.

The caller supplies draft manifests returned by ``write_issue_package`` and
receipts returned by the owned PDF renderer, keyed by language. No approval,
upload, activation timestamp or publication history is created here. Recheck
the projected ordinal against the actual queue before owner approval.
"""

from copy import deepcopy
from io import BytesIO
import os
from pathlib import Path
import posixpath
import re
import tempfile
from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook

from .live_outputs import (LEGACY_PRICE_DISPLAY_VERSION, PRICE_DISPLAY_VERSION, digest_bytes,
                           price_display, validate_dataset, validate_stage_provenance, supplier_search_display, _url,
                           evidence_workbook_rows, supplier_workbook_rows, workbook_contract)
from .report_source_facts import CONTRACTS as SOURCE_ROUTE_CONTRACTS
from .report_examples import RECOMMENDATION_CONTRACT, WORKBOOK_HEADERS as EXAMPLE_WORKBOOK_HEADERS, example_workbook_rows
from .models import SafetyError, canonical, digest, identifier
from .opportunity_score import workbook_headers as opportunity_workbook_headers, workbook_values as opportunity_workbook_values
from .runtime_storage import runtime_path
from .storage import load_json, safe_path

MAX_FILE = 32 * 1024 * 1024  # Same bound as the private portal.
LANGUAGES = ("en", "hi-Latn")
HASH = re.compile(r"[a-f0-9]{64}\Z")


def _hash(value):
    if not isinstance(value, str) or not HASH.fullmatch(value):
        raise SafetyError("DELIVERY_HASH_INVALID")
    return value


def _read(path, limit=MAX_FILE):
    path = Path(path)
    safe_path(path.parent)
    if path.is_symlink():
        raise SafetyError("DELIVERY_PATH_INVALID")
    try:
        with path.open("rb") as stream:
            raw = stream.read(limit + 1)
    except OSError:
        raise SafetyError("DELIVERY_FILE_MISSING") from None
    if not raw or len(raw) > limit:
        raise SafetyError("DELIVERY_FILE_SIZE_INVALID")
    return raw


def _artifact(record, expected, limit=MAX_FILE):
    if not isinstance(record, dict) or record.get("path") != str(expected):
        raise SafetyError("DELIVERY_ARTIFACT_PATH_MISMATCH")
    raw = _read(expected, limit)
    if digest_bytes(raw) != _hash(record.get("sha256")):
        raise SafetyError("DELIVERY_ARTIFACT_CHANGED")
    return {"path": str(expected), "sha256": record["sha256"], "size": len(raw)}, raw


def _source_hyperlinks(zipped):
    """Retain the real cell targets hidden by read-only worksheet loading."""
    namespace = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    relationships = {row.get("Id"): row for row in ElementTree.fromstring(zipped.read("xl/_rels/workbook.xml.rels"))}
    output = {}
    for sheet in ElementTree.fromstring(zipped.read("xl/workbook.xml")).iter():
        name = sheet.get("name")
        if name not in ("Evidence", "Supplier leads", "Example products"):
            continue
        relationship = relationships.get(sheet.get(namespace + "id"))
        if relationship is None or not relationship.get("Type", "").endswith("/worksheet"):
            raise SafetyError("DELIVERY_WORKBOOK_LINK_INVALID")
        target = relationship.get("Target", "")
        path = posixpath.normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
        if relationship.get("TargetMode") == "External" or not path.startswith("xl/worksheets/"):
            raise SafetyError("DELIVERY_WORKBOOK_LINK_INVALID")
        rel_path = posixpath.join(posixpath.dirname(path), "_rels", posixpath.basename(path) + ".rels")
        links = {row.get("Id"): row.get("Target") for row in ElementTree.fromstring(zipped.read(rel_path))
                 if row.get("Type", "").endswith("/hyperlink")} if rel_path in zipped.namelist() else {}
        output[name] = {row.get("ref"): links.get(row.get(namespace + "id"))
                        for row in ElementTree.fromstring(zipped.read(path)).iter()
                        if row.tag.endswith("}hyperlink")}
    return output


def _workbook(raw, dataset, *, price_display_version=PRICE_DISPLAY_VERSION):
    try:
        functional = dataset.get("schema_version") == 3
        marked = any(product.get("source_route_contract") in SOURCE_ROUTE_CONTRACTS for product in dataset["products"])
        with ZipFile(BytesIO(raw)) as zipped:
            if len(zipped.infolist()) > 500 or sum(entry.file_size for entry in zipped.infolist()) > 128 * 1024 * 1024:
                raise SafetyError("DELIVERY_WORKBOOK_SIZE_INVALID")
            # Read-only worksheet cells omit hyperlink destinations. Check the
            # actual relationships, not the harmless displayed "Open source".
            # The same public-link rule preserves valid historical packets.
            for entry in zipped.infolist():
                if not entry.filename.endswith(".rels"):
                    continue
                for relationship in ElementTree.fromstring(zipped.read(entry)):
                    if not relationship.get("Type", "").endswith("/hyperlink"):
                        continue
                    target = relationship.get("Target")
                    if relationship.get("TargetMode") != "External" or not isinstance(target, str) or not target:
                        raise SafetyError("DELIVERY_WORKBOOK_LINK_INVALID")
                    try:
                        _url(target)
                    except SafetyError:
                        raise SafetyError("DELIVERY_WORKBOOK_LINK_INVALID") from None
            source_hyperlinks = _source_hyperlinks(zipped) if marked or functional else {}
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        try:
            new_fit = dataset.get("fit_policy_id") == "fit-10-v1"
            if book.sheetnames != ["Shortlist", "Evidence", "Supplier leads"] + (["Checklist"] if new_fit else []) + (["Example products"] if functional else []):
                raise SafetyError("DELIVERY_WORKBOOK_MISMATCH")
            expected_headers = ("Rank", "Product", "Category", "Status", "Checklist", "Market", "Retail prices", "Supplier quotes", "Concept ID", "Supplier search")
            if "opportunity_policy_id" in dataset:
                expected_headers += opportunity_workbook_headers()
            if tuple(next(book["Shortlist"].iter_rows(min_row=6, max_row=6, values_only=True))) != expected_headers:
                raise SafetyError("DELIVERY_WORKBOOK_MISMATCH")
            rows = list(book["Shortlist"].iter_rows(min_row=7, values_only=True))
            if len(rows) != len(dataset["products"]):
                raise SafetyError("DELIVERY_WORKBOOK_MISMATCH")
            for row, product in zip(rows, dataset["products"]):
                expected_row = [product["rank"], product["title"], product["category"], product["status"], f'{product["kpi_matches"]}/{dataset.get("kpi_total", 13)}',
                                product["market_label"], "\n".join(price_display(p, render_version=price_display_version) for p in product["prices"] if p["kind"] == "retail") or "Not stated",
                                "\n".join(price_display(p, render_version=price_display_version) for p in product["prices"] if p["kind"] == "supplier") or "Not stated", product["concept_id"],
                                supplier_search_display(product)]
                if "opportunity_policy_id" in dataset:
                    expected_row.extend(opportunity_workbook_values(product))
                expected_row = ["'" + value if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")) else value for value in expected_row]
                if tuple(row) != tuple(expected_row):
                    raise SafetyError("DELIVERY_WORKBOOK_MISMATCH")
            wanted = {p["concept_id"] for p in dataset["products"]}
            if new_fit:
                from .research_ai import CRITERIA_BY_ID
                headers = ("Rank", "Product", "Criterion ID", "Criterion", "State", "Basis", "Reason", "References", "Fit", "Concept ID")
                if tuple(next(book["Checklist"].iter_rows(min_row=6, max_row=6, values_only=True))) != headers:
                    raise SafetyError("DELIVERY_WORKBOOK_MISMATCH")
                expected = []
                for product in dataset["products"]:
                    for criterion in product["criteria"]:
                        row = [product["rank"], product["title"], criterion["id"], CRITERIA_BY_ID[criterion["id"]], criterion["state"],
                               # XLSX stores an empty string as a blank cell.
                               # Unknown judgments may legitimately have no
                               # references; do not mistake that for damage.
                               criterion["basis"], criterion["reason"], ", ".join(criterion["references"]) or None,
                               "Good fit" if product["fit_tier"] == "good" else "Decent fit", product["concept_id"]]
                        expected.append(tuple("'" + value if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")) else value for value in row))
                if list(book["Checklist"].iter_rows(min_row=7, values_only=True)) != expected:
                    raise SafetyError("DELIVERY_WORKBOOK_MISMATCH")
            for name, column in (("Evidence", 7), ("Supplier leads", 9)):
                for row in book[name].iter_rows(min_row=7, values_only=True):
                    if any(v is not None for v in row) and row[column] not in wanted:
                        raise SafetyError("DELIVERY_WORKBOOK_MISMATCH")
            if marked:
                # Old packets keep their historical validation contract. New
                # saved-route packets must really contain all frozen outcomes,
                # numbers and supplier leads, not just correct product IDs.
                source_sheets = [("Evidence", evidence_workbook_rows, 6),
                                 ("Supplier leads", supplier_workbook_rows, 8)]
                if functional:
                    if tuple(next(book["Example products"].iter_rows(min_row=6, max_row=6, values_only=True))) != EXAMPLE_WORKBOOK_HEADERS:
                        raise SafetyError("DELIVERY_WORKBOOK_MISMATCH")
                    source_sheets.append(("Example products", example_workbook_rows, 8))
                for name, make_rows, link in source_sheets:
                    expected = []
                    for product in dataset["products"]:
                        for original in make_rows(product):
                            row = list(original)
                            url = row[link] if isinstance(row[link], str) and row[link].startswith(("http://", "https://")) else None
                            cell = chr(65 + link) + str(len(expected) + 7)
                            if source_hyperlinks.get(name, {}).get(cell) != url:
                                raise SafetyError("DELIVERY_WORKBOOK_LINK_MISMATCH")
                            if url is not None:
                                row[link] = "Open source"
                            expected.append(tuple("'" + value if isinstance(value, str)
                                and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")) else value for value in row))
                    actual = [row for row in book[name].iter_rows(min_row=7, values_only=True)
                              if any(value is not None for value in row)]
                    if actual != expected:
                        raise SafetyError("DELIVERY_WORKBOOK_MISMATCH")
        finally:
            book.close()
    except SafetyError:
        raise
    except Exception:
        raise SafetyError("DELIVERY_WORKBOOK_INVALID") from None


def _pdf(raw, receipt):
    import fitz
    if not raw.startswith(b"%PDF-"):
        raise SafetyError("DELIVERY_PDF_INVALID")
    try:
        with fitz.open(stream=raw, filetype="pdf") as document:
            if not 1 <= len(document) <= 100:
                raise SafetyError("DELIVERY_PDF_INVALID")
            for page in document:
                if page.rect.width >= page.rect.height or not page.get_text().strip():
                    raise SafetyError("DELIVERY_PDF_LAYOUT_INVALID")
            if (type(receipt.get("pages")) is not int or receipt["pages"] != len(document)
                    or receipt.get("portrait") is not True or type(receipt.get("bytes")) is not int
                    or receipt["bytes"] != len(raw)):
                raise SafetyError("DELIVERY_RENDER_RECEIPT_MISMATCH")
    except SafetyError:
        raise
    except Exception:
        raise SafetyError("DELIVERY_PDF_INVALID") from None


def build_delivery_manifest(issue, language_manifests, pdf_receipts, *, root,
                            publication_ordinal=None, selected_language=None):
    """Validate frozen facts plus actual files, returning a draft manifest.

    ``language_manifests`` and ``pdf_receipts`` are dictionaries keyed by en or
    hi-Latn. Receipts contain path/html_sha256/pdf_sha256/pages/portrait/bytes.
    With no language choice after issue two, return a checked language_pending
    draft with no download entries. Missing required trial PDFs are an error,
    not a complete packet. The ordinal can be explicitly reprojected by the
    caller; a changed final packet must not replace accepted revision bytes.
    """
    root = safe_path(root).resolve()
    issue_id, client_id = identifier(issue.get("issue_id")), identifier(issue.get("client_id"))
    revision = issue.get("revision")
    if type(revision) is not int or revision < 1:
        raise SafetyError("DELIVERY_REVISION_INVALID")
    ordinal = issue.get("projected_ordinal") if publication_ordinal is None else publication_ordinal
    if type(ordinal) is not int or ordinal < 1:
        raise SafetyError("DELIVERY_ORDINAL_INVALID")
    if selected_language not in (None, *LANGUAGES):
        raise SafetyError("DELIVERY_LANGUAGE_INVALID")
    if (not isinstance(language_manifests, dict) or not language_manifests
            or set(language_manifests) - set(LANGUAGES) or not isinstance(pdf_receipts, dict)
            or set(pdf_receipts) - set(language_manifests)):
        raise SafetyError("DELIVERY_INPUT_INVALID")
    data = safe_path(runtime_path(root, "data", "outputs"), issue_id, "revision-" + str(revision))
    media = safe_path(runtime_path(root, "media", "live", "reports"), issue_id, "revision-" + str(revision))
    dataset = load_json(data / "final_report_data.json", 4 * 1024 * 1024)
    validate_dataset(issue, dataset, root=root if dataset.get("schema_version") == 3 else None)
    choice = selected_language if selected_language is not None else dataset["selected_language"]
    if selected_language is not None and dataset["selected_language"] not in (None, selected_language):
        raise SafetyError("DELIVERY_LANGUAGE_CHOICE_MISMATCH")
    required = list(LANGUAGES) if ordinal <= 2 else ([] if choice is None else [choice])
    if set(required) - set(language_manifests) or set(required) - set(pdf_receipts):
        raise SafetyError("DELIVERY_REQUIRED_LANGUAGE_MISSING")
    artifacts = {}
    workbook_version = None
    for language in sorted(language_manifests):
        draft = language_manifests[language]
        saved = load_json(data / ("manifest-" + language + ".json"))
        if not isinstance(draft, dict) or canonical(draft) != canonical(saved):
            raise SafetyError("DELIVERY_DRAFT_MANIFEST_CHANGED")
        if "price_display_version" in draft:
            if type(draft["price_display_version"]) is not str or draft["price_display_version"] != PRICE_DISPLAY_VERSION:
                raise SafetyError("DELIVERY_PRICE_DISPLAY_VERSION_INVALID")
            render_version = PRICE_DISPLAY_VERSION
        else:
            # Preserve immutable pre-versioning manifests exactly. Missing
            # means the one historical deterministic formatter, not any label.
            render_version = LEGACY_PRICE_DISPLAY_VERSION
        if workbook_version is not None and workbook_version != render_version:
            raise SafetyError("DELIVERY_SHARED_RENDER_VERSION_MISMATCH")
        workbook_version = render_version
        expected = {"schema_version": 2 if dataset["schema_version"] == 3 else 1, "issue_id": issue_id, "revision": revision, "language": language,
                    "facts_hash": digest(dataset), "selection_hash": issue["selection_hash"],
                    "evidence_type": dataset["evidence_type"], "counts": issue["counts"],
                    "concept_ids": [p["concept_id"] for p in dataset["products"]],
                    "projected_ordinal": issue["projected_ordinal"], "language_state": issue["language_state"],
                    "approved": False, "release_ready": False}
        if dataset.get("fit_policy_id") == "fit-10-v1":
            expected.update(fit_policy_id="fit-10-v1", criterion_ids=dataset["criterion_ids"], kpi_total=10, workbook_contract=workbook_contract(dataset))
        elif any(field in draft for field in ("fit_policy_id", "criterion_ids", "kpi_total", "workbook_contract")):
            raise SafetyError("DELIVERY_DRAFT_BINDING_MISMATCH")
        if "opportunity_policy_id" in dataset:
            expected["opportunity_policy_id"] = dataset["opportunity_policy_id"]
        elif "opportunity_policy_id" in draft:
            raise SafetyError("DELIVERY_DRAFT_BINDING_MISMATCH")
        if dataset["schema_version"] == 3:
            expected.update(recommendation_contract=RECOMMENDATION_CONTRACT,
                            asset_receipt_hash=digest(dataset["report_asset_receipt"]))
        elif any(field in draft for field in ("recommendation_contract", "asset_receipt_hash")):
            raise SafetyError("DELIVERY_DRAFT_BINDING_MISMATCH")
        if any(draft.get(field) != value or (type(value) in (bool, int) and type(draft.get(field)) is not type(value)) for field, value in expected.items()):
            raise SafetyError("DELIVERY_DRAFT_BINDING_MISMATCH")
        if "internal_stage" in draft:
            validate_stage_provenance(dataset, language, draft["internal_stage"])
        for field, path, limit in (("facts", data / "final_report_data.json", 4 * 1024 * 1024),
                                   ("workbook", data / "shortlist.xlsx", MAX_FILE),
                                   ("html", media / ("report-" + language + ".html"), MAX_FILE)):
            record, raw = _artifact(draft.get("artifacts", {}).get(field), path, limit)
            key = field if field != "html" else "html-" + language
            if key in artifacts and artifacts[key] != record:
                raise SafetyError("DELIVERY_SHARED_FACTS_MISMATCH")
            artifacts[key] = record
            if field == "workbook":
                _workbook(raw, dataset, price_display_version=render_version)
        if language in pdf_receipts:
            receipt = pdf_receipts[language]
            if not isinstance(receipt, dict) or receipt.get("html_sha256") != artifacts["html-" + language]["sha256"]:
                raise SafetyError("DELIVERY_RENDER_HTML_MISMATCH")
            record, raw = _artifact({"path": receipt.get("path"), "sha256": receipt.get("pdf_sha256")}, media / ("report-" + language + ".pdf"))
            _pdf(raw, receipt)
            artifacts["pdf-" + language] = record
    files = []
    if required:
        prefix = f"clients/{client_id}/issues/{issue_id}/r{revision}/"
        for name, record, media_type in [("shortlist.xlsx", artifacts["workbook"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
                *(("report-" + language + ".pdf", artifacts["pdf-" + language], "application/pdf") for language in required)]:
            files.append({"name": name, "key": prefix + name, "size": record["size"], "sha256": record["sha256"], "media_type": media_type})
    manifest = {"schema_version": 2 if dataset["schema_version"] == 3 else 1, "issue_id": issue_id, "client_id": client_id, "revision": revision,
                "release_date": issue["date"], "selection_hash": issue["selection_hash"], "evidence_fingerprint": issue["evidence_fingerprint"],
                "facts_hash": digest(dataset), "evidence_type": dataset["evidence_type"], "concept_ids": [p["concept_id"] for p in dataset["products"]],
                "counts": deepcopy(issue["counts"]), "product_count": issue["counts"]["products"], "new_count": issue["counts"]["new"],
                "publication_ordinal": ordinal, "selected_language": choice, "languages": required,
                "complete": bool(required), "state": "draft_ready_for_review" if required else "language_pending",
                "approved": False, "release_ready": False, "files": files, "artifacts": artifacts,
                "language_manifests": deepcopy(language_manifests), "pdf_receipts": deepcopy(pdf_receipts)}
    if dataset.get("fit_policy_id") == "fit-10-v1":
        manifest.update(fit_policy_id="fit-10-v1", criterion_ids=deepcopy(dataset["criterion_ids"]), kpi_total=10, workbook_contract=workbook_contract(dataset))
    if dataset["schema_version"] == 3:
        manifest.update(recommendation_contract=RECOMMENDATION_CONTRACT,
                        asset_receipt_hash=digest(dataset["report_asset_receipt"]))
    if "opportunity_policy_id" in dataset:
        manifest["opportunity_policy_id"] = dataset["opportunity_policy_id"]
    manifest["package_hash"] = digest(manifest)
    return manifest


def validate_delivery_manifest(manifest, issue, *, root, publication_ordinal=None, selected_language=None):
    """Rebuild and compare, detecting modified files, queue or language choice."""
    if not isinstance(manifest, dict):
        raise SafetyError("DELIVERY_MANIFEST_INVALID")
    rebuilt = build_delivery_manifest(issue, manifest.get("language_manifests"), manifest.get("pdf_receipts"), root=root,
                                      publication_ordinal=publication_ordinal,
                                      selected_language=manifest.get("selected_language") if selected_language is None else selected_language)
    if canonical(manifest) != canonical(rebuilt):
        raise SafetyError("DELIVERY_MANIFEST_CHANGED")
    return manifest


def write_delivery_manifest(issue, language_manifests, pdf_receipts, *, root,
                            publication_ordinal=None, selected_language=None):
    """Commit last, create-if-absent; replay never overwrites accepted bytes."""
    manifest = build_delivery_manifest(issue, language_manifests, pdf_receipts, root=root,
                                      publication_ordinal=publication_ordinal, selected_language=selected_language)
    path = safe_path(runtime_path(Path(root).resolve(), "data", "outputs"), issue["issue_id"], "revision-" + str(issue["revision"])) / "delivery-manifest.json"
    safe_path(path.parent)
    if path.is_symlink():
        raise SafetyError("DELIVERY_PATH_INVALID")
    encoded = canonical(manifest) + b"\n"
    if path.exists():
        if _read(path, 512 * 1024) != encoded:
            raise SafetyError("DELIVERY_IMMUTABLE_MANIFEST_CHANGED")
        return manifest
    if len(encoded) > 512 * 1024:
        raise SafetyError("DELIVERY_MANIFEST_SIZE_INVALID")
    descriptor, temporary = tempfile.mkstemp(prefix=".delivery-", dir=path.parent)
    try:
        with os.fdopen(descriptor, "wb") as handle:
            handle.write(encoded)
            handle.flush()
            os.fsync(handle.fileno())
        try:
            os.link(temporary, path)
        except FileExistsError:
            if _read(path, 512 * 1024) != encoded:
                raise SafetyError("DELIVERY_IMMUTABLE_MANIFEST_CHANGED") from None
        directory = os.open(path.parent, os.O_RDONLY)
        try:
            os.fsync(directory)
        finally:
            os.close(directory)
    finally:
        os.unlink(temporary)
    return manifest
