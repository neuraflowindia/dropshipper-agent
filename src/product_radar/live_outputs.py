"""Immutable factual packets and internal live-layout adapters, not a publisher.

The dataset contract is explicit: frozen issue IDs/category/status/checklist;
product title, market_checks, prices, suppliers and source_refs supplied by
research; narratives supplied independently per language. Missing facts remain
unknown. There is no model call, translation, collection or remote upload here.
"""

from base64 import b64encode
from copy import deepcopy
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from html import escape
import importlib.util
from io import BytesIO
from pathlib import Path
import re
import textwrap
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

from .models import SafetyError, canonical, digest, identifier, timestamp
from .public_pages import _public_url
from .service import validate_issue, policy_for_fit, product_fit, issue_fit_policy
from .fit_policy import FIT_POLICY_ID, LEGACY_POLICY_ID, evaluate_fit, policy_for
from .opportunity_score import (OPPORTUNITY_FIELDS, POLICY_ID as SCORE_POLICY_ID, crowding_text, score_text,
                                workbook_headers as opportunity_workbook_headers, workbook_values as opportunity_workbook_values)
from .supplier_effort import frozen_supplier_effort_valid
from .report_source_facts import (CONTRACT as SOURCE_ROUTE_CONTRACT, CONTRACTS as SOURCE_ROUTE_CONTRACTS,
                                 FUNCTIONAL_CONTRACT as FUNCTIONAL_SOURCE_ROUTE_CONTRACT, FIELDS as SOURCE_ROUTE_FIELDS,
                                  validate_route_facts)
from .runtime_storage import logical_path, runtime_path
from .storage import Journal, atomic_bytes, atomic_json, load_json, safe_path
from .report_examples import (RECOMMENDATION_CONTRACT, FUNCTIONAL_REPORT_FIELDS, OPTIONAL_FUNCTIONAL_REPORT_FIELDS,
    functional_report_fields, WORKBOOK_HEADERS as EXAMPLE_WORKBOOK_HEADERS,
                              validate_functional_examples, example_workbook_rows)

ROOT = Path(__file__).resolve().parents[2]
LOGO = ROOT / "media" / "NeuraFlow" / "long_maroon_path_transparent_1024.png"
LANGUAGES = {"en", "hi-Latn"}
PRICE_DISPLAY_VERSION = "price-display-v2"
LEGACY_PRICE_DISPLAY_VERSION = "price-display-v1"
INTERNAL_STAGE_VERSION = "checked-report-stages-v1"
WORKBOOK_GENERATION_VERSION = "deterministic-workbook-v1-style-v2"
LEGACY_WORKBOOK_GENERATION_VERSION = "legacy-verified-workbook-v1"
LEGACY_HTML_RENDER_VERSION = "functional-report-stages-v1"
HTML_RENDER_VERSION = "functional-report-stages-v1-print-v3"
SOURCE_WORKBOOK_GENERATION_VERSION = "deterministic-workbook-v2-style-v2"
LEGACY_SOURCE_HTML_RENDER_VERSION = "functional-report-stages-v2"
SOURCE_HTML_RENDER_VERSION = "functional-report-stages-v2-print-v3"
TYPE_WORKBOOK_GENERATION_VERSION = "deterministic-workbook-v3-style-v2"
LEGACY_TYPE_HTML_RENDER_VERSION = "functional-report-stages-v3"
TYPE_HTML_RENDER_VERSION = "functional-report-stages-v3-print-v3"
TYPE_WORKBOOK_CONTRACT = "fit-checklist-examples-v2"
LEGACY_SUPPLIER_TERMS_DISPLAY_VERSION = "supplier-terms-v1"
SUPPLIER_TERMS_DISPLAY_VERSION = "supplier-terms-v2"
MAX_STAGE_FILE = 32 * 1024 * 1024
MAX_CORE_FILE = 4 * 1024 * 1024
from .report_copy import (DISPATCH_LABELS, SUBSCRIPTION_LABELS, HI_DISPATCH_LABELS,  # noqa: E402  (shared wording)
                          HI_SUBSCRIPTION_LABELS)


@lru_cache(maxsize=1)
def _style_helpers():
    modules = []
    for name in ("generate_sample_excel", "generate_sample_pdf"):
        spec = importlib.util.spec_from_file_location("radar_shared_" + name, ROOT / "scripts" / (name + ".py"))
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        modules.append(module)
    return modules


def _text(value, *, optional=False):
    if value is None and optional:
        return
    if not isinstance(value, str) or (not optional and not value.strip()):
        raise SafetyError("LIVE_TEXT_INVALID")


def _url(value):
    if value is None:
        return
    try:
        # Validate only. Keep the original public link, including its fragment,
        # so an immutable packet does not acquire normalized replacement facts.
        _public_url(value)
    except SafetyError:
        raise SafetyError("LIVE_SOURCE_URL_INVALID") from None


def _price(price):
    if not isinstance(price, dict) or price.get("kind") not in ("retail", "supplier") or price.get("currency") != "INR":
        raise SafetyError("LIVE_PRICE_INVALID")
    amount = price.get("amount")
    if amount is not None:
        if not isinstance(amount, str):
            raise SafetyError("LIVE_PRICE_INVALID")
        try:
            number = Decimal(amount)
            if not number.is_finite() or number < 0 or number > Decimal("1000000000") or number.as_tuple().exponent < -8:
                raise ValueError
        except (InvalidOperation, ValueError):
            raise SafetyError("LIVE_PRICE_INVALID") from None
    for field in ("quantity", "moq"):
        if price.get(field) is not None and (type(price[field]) is not int or price[field] < 1):
            raise SafetyError("LIVE_PRICE_INVALID")
    if price.get("unit") not in ("per_item", "per_pack", "bulk_quote", "not_stated"):
        raise SafetyError("LIVE_PRICE_INVALID")
    _text(price.get("conditions", ""), optional=True)
    _url(price.get("source_url"))
    if price.get("observed_at"):
        timestamp(price["observed_at"])


def price_display(price, language="en", *, render_version=PRICE_DISPLAY_VERSION):
    _price(price)
    if language not in LANGUAGES:
        raise SafetyError("LIVE_LANGUAGE_INVALID")
    if type(render_version) is not str or render_version not in (LEGACY_PRICE_DISPLAY_VERSION, PRICE_DISPLAY_VERSION):
        raise SafetyError("LIVE_PRICE_DISPLAY_VERSION_INVALID")
    amount = price.get("amount")
    raw = price.get("raw_displayed")
    raw_retail_quote = (render_version == PRICE_DISPLAY_VERSION and amount is None
                        and price["kind"] == "retail" and isinstance(raw, str) and bool(raw.strip()))
    value = ("Price not stated" if language == "en" else "keemat nahi di gayi") if amount is None else "₹" + format(Decimal(amount), "f")
    if raw_retail_quote:
        value = ("displayed quote " if language == "en" else "dikhaya gaya quote ") + raw
    kind = ("Retail listing" if price["kind"] == "retail" else "Supplier quote") if language == "en" else ("Retail listing ki keemat" if price["kind"] == "retail" else "Supplier ka quote")
    units = {"per_item": "per item", "per_pack": "per pack", "bulk_quote": "bulk quote", "not_stated": "unit not stated"} if language == "en" else {
        "per_item": "har item ke liye", "per_pack": "har pack ke liye", "bulk_quote": "bulk order ka quote", "not_stated": "unit nahi batayi gayi"}
    parts = [kind + ": " + value, units[price["unit"]]]
    if raw_retail_quote:
        parts.append("price basis not verified" if language == "en" else "keemat ka basis abhi verify nahi hua hai")
    if price.get("quantity") is not None:
        parts.append(("pack quantity " if language == "en" else "pack ki quantity: ") + str(price["quantity"]))
    if price.get("moq") is not None:
        parts.append(("minimum order " if language == "en" else "minimum order: ") + str(price["moq"]))
    if price.get("conditions"):
        parts.append(("" if language == "en" else "Source ki shartein: ") + price["conditions"])
    return "; ".join(parts)


def supplier_display(supplier, language="en", *, terms_version=LEGACY_SUPPLIER_TERMS_DISPLAY_VERSION):
    from .supplier_scoped_display import is_scoped_supplier, scoped_supplier_lines
    if is_scoped_supplier(supplier):
        return "\n".join(scoped_supplier_lines(supplier, language))
    terms_text = supplier_terms_display(supplier, language, terms_version=terms_version)
    dispatch = (DISPATCH_LABELS if language == "en" else HI_DISPATCH_LABELS)[supplier["single_piece_dispatch"]]
    fees = _supplier_subscription_display(supplier, language, terms_version=terms_version)
    return "; ".join([dispatch, fees, terms_text] +
                     ([("minimum order " if language == "en" else "minimum order: ") + str(supplier["min_order"])] if supplier.get("min_order") is not None else []))


def _supplier_subscription_display(supplier, language="en", *, terms_version=LEGACY_SUPPLIER_TERMS_DISPLAY_VERSION):
    fees = (SUBSCRIPTION_LABELS if language == "en" else HI_SUBSCRIPTION_LABELS)[supplier["subscription"]]
    if terms_version == SUPPLIER_TERMS_DISPLAY_VERSION and supplier["subscription"] == "not_stated":
        rows = [row for row in supplier.get("terms", []) if isinstance(row, dict)
                and row.get("field") == "subscription_requirement"] if isinstance(supplier.get("terms"), list) else []
        if len(rows) == 1 and rows[0].get("basis") == "observed" and rows[0].get("value") == "Optional":
            fees = ("Optional subscription stated by the supplier" if language == "en" else
                    "Supplier ne optional subscription ka zikr kiya hai")
        elif len(rows) == 1 and rows[0].get("basis") == "conflicting" and rows[0].get("value") == "Conflicting":
            fees = ("Supplier subscription statements conflict" if language == "en" else
                    "Supplier ki subscription ki shartein aapas mein nahi milti hain")
    return fees


def supplier_terms_display(supplier, language="en", *, terms_version=LEGACY_SUPPLIER_TERMS_DISPLAY_VERSION):
    """Keep observed purchase and delivery distinct from the combined flag."""
    if (type(terms_version) is not str or terms_version not in
            (LEGACY_SUPPLIER_TERMS_DISPLAY_VERSION, SUPPLIER_TERMS_DISPLAY_VERSION)):
        raise SafetyError("LIVE_SUPPLIER_TERMS_DISPLAY_VERSION_INVALID")
    if type(language) is not str or language not in LANGUAGES:
        raise SafetyError("LIVE_LANGUAGE_INVALID")
    parts = []
    fields = (("single_piece_purchase", "Single-piece purchase", "Single-piece kharid"),
              ("direct_customer_dispatch", "Direct customer dispatch", "Seedha customer ko dispatch"))
    for field, english, hinglish in fields:
        terms = supplier.get("terms")
        rows = [row for row in terms if isinstance(row, dict) and row.get("field") == field] if isinstance(terms, list) else []
        # Duplicated/conflicting or inferred statements are not observed Yes/No.
        value = rows[0].get("value") if len(rows) == 1 and rows[0].get("basis") == "observed" else None
        if value not in ("Yes", "No"):
            value = "Unknown"
        text = value if language == "en" else {"Yes": "Haan", "No": "Nahi", "Unknown": "Abhi clear nahi hai"}[value]
        parts.append((english if language == "en" else hinglish) + ": " + text)
    if terms_version == LEGACY_SUPPLIER_TERMS_DISPLAY_VERSION:
        # Original completed and partial packets keep their exact two-field
        # projection. Only a version marker in frozen facts selects new text.
        return "; ".join(parts)
    extra_fields = (
        ("purchase_moq", "Minimum purchase", "Minimum kharid"),
        ("subscription_requirement", "Subscription requirement", "Subscription ki zaroorat"),
        ("recurring_fee_frequency", "Recurring fee frequency", "Baar-baar lagne wali fee kab deni hai"),
        ("monthly_fee", "Monthly fee", "Har mahine ki fee"),
        ("bulk_stock_prerequisite", "Bulk stock must be bought first", "Pehle bulk stock kharidna zaroori hai"),
        ("cod", "Cash on delivery", "Delivery par payment"),
        ("india_coverage", "Delivery coverage", "Delivery kahan tak hai"),
        ("white_label", "Product offered under your brand", "Aapke brand ke naam se product"),
        ("white_label_shipping", "Shipping under your brand", "Aapke brand ke naam se shipping"),
        ("neutral_packaging", "Plain packaging without supplier branding", "Supplier ke brand ke bina saada packing"),
        ("ecommerce_support", "Support for online selling", "Online bechne mein madad"),
        ("order_integration", "Order-system connection", "Order system se connection"),
        ("setup_fee", "Setup fee", "Shuru karne ki fee"),
        ("deposit", "Deposit", "Jama karne ki rakam"),
        ("per_order_fee", "Fee for each order", "Har order ki fee"),
    )
    yes_no = {"bulk_stock_prerequisite", "cod", "white_label",
              "white_label_shipping", "neutral_packaging", "ecommerce_support", "order_integration"}
    terms = supplier.get("terms")
    rows = terms if isinstance(terms, list) else []
    added = False
    for field, english, hinglish in extra_fields:
        matches = [row for row in rows if isinstance(row, dict) and row.get("field") == field]
        if not matches or all(row.get("basis") == "not_stated" for row in matches):
            # Do not pad every supplier with fifteen repeated unknown fields.
            # The existing purchase/delivery unknowns remain explicit above.
            continue
        row = matches[0] if len(matches) == 1 else {}
        value = row.get("value")
        if row.get("basis") == "conflicting" and value == "Conflicting":
            value = "Conflicting statements" if language == "en" else "Shartein aapas mein nahi milti hain"
        elif (row.get("basis") != "observed" or not isinstance(value, str) or not value.strip()
                or value.strip().casefold() in {"not stated", "unknown", "unclear", "conflicting"}
                or field in yes_no and value not in {"Yes", "No"}
                or field == "subscription_requirement" and value not in {"Required", "Optional", "Not required"}):
            value = "Unknown" if language == "en" else "Abhi clear nahi hai"
        elif language == "hi-Latn":
            # Preserve source fee amounts, currencies, minima and frequencies
            # verbatim; only the fixed Yes/No/subscription vocabulary changes.
            value = {"Yes": "Haan", "No": "Nahi", "Required": "Zaroori",
                     "Optional": "Aapki marzi par", "Not required": "Zaroori nahi"}.get(value, value)
        parts.append((english if language == "en" else hinglish) + ": " + value)
        added = True
    if added and supplier.get("relationship") in {"relevant_company", "unclear"}:
        parts.append("Company terms only; supply of this product is not confirmed." if language == "en" else
                     "Yeh company ki shartein hain; is product ki supply abhi confirm nahi hai.")
    return "; ".join(parts)


def _supplier_terms_version(product):
    version = product.get("supplier_terms_display_version", LEGACY_SUPPLIER_TERMS_DISPLAY_VERSION)
    if (type(version) is not str or version not in
            (LEGACY_SUPPLIER_TERMS_DISPLAY_VERSION, SUPPLIER_TERMS_DISPLAY_VERSION)):
        raise SafetyError("LIVE_SUPPLIER_TERMS_DISPLAY_VERSION_INVALID")
    return version


def _supplier_minimum_order(supplier):
    """Preserve an observed compound minimum; do not invent a numeric MOQ."""
    if supplier.get("min_order") is not None:
        return supplier["min_order"]
    terms = supplier.get("terms")
    rows = [row for row in terms if isinstance(row, dict) and row.get("field") == "purchase_moq"] if isinstance(terms, list) else []
    if len(rows) == 1 and rows[0].get("basis") == "observed":
        value = rows[0].get("value")
        if isinstance(value, str) and value.strip() and value.strip().casefold() not in ("not stated", "conflicting", "unknown", "unclear"):
            return value
    return "Not stated"


def supplier_search_display(product):
    if (product.get("fit_policy_id") == FIT_POLICY_ID and product.get("supplier_effort_satisfied") is True
            and not product["suppliers"]):
        return "Supplier search did not produce any usable results."
    return {"not_started": "Not started", "attempted": "Attempted", "complete": "Complete", "blocked": "Could not complete"}[product["supplier_search_status"]]


def _bulk_order_notice(price, language):
    """Display named minimum quantities only; never turn a pack/price into MOQ."""
    if price["kind"] != "supplier":
        return None
    pattern = (r"\b(?:MOQ|minimum\s+order(?:\s+(?:quantity|qty))?|min\.?\s+order(?:\s+(?:quantity|qty))?)"
               r"\s*(?:(?:is|of)\s+|[:=\-]\s*)?(?P<number>[0-9]+(?:,[0-9]{3})*)"
               r"(?![0-9,.])(?P<unit>\s*(?:pieces?|pcs|units?|sets?|rolls?))?\b"
               r"(?P<compound>\s+of\s+[0-9]+)?")
    conditions = price.get("conditions") or ""
    for match in re.finditer(pattern, conditions, re.IGNORECASE):
        if match["number"].replace(",", "").lstrip("0") in ("", "1"):
            continue
        # A volume/weight/monetary minimum is not a count of dispatch units.
        if not match["unit"] and re.match(r"\s*(?:(?:ml|litres?|liters?|kg|grams?|gm|g|cm|mm|INR|rupees?|Rs|USD)\b|₹|\$)", conditions[match.end():], re.IGNORECASE):
            continue
        return ("Some quoted offers have bulk minimum orders. Check each offer's conditions before assuming single-piece sourcing." if language == "en" else
                "Kuch quotes mein bulk minimum order hai. Single-piece kharid maan lene se pehle har offer ki shartein dekhein.")
    return None


def build_issue_dataset(issue, product_facts, *, narratives=None, evidence_type="captured"):
    if evidence_type not in ("captured", "synthetic") or not isinstance(product_facts, list):
        raise SafetyError("LIVE_DATASET_INVALID")
    supplied = {}
    for value in product_facts:
        if not isinstance(value, dict):
            raise SafetyError("LIVE_PRODUCT_INVALID")
        concept_id = identifier(value.get("concept_id"))
        if concept_id in supplied:
            raise SafetyError("LIVE_DUPLICATE_CONCEPT")
        supplied[concept_id] = value
    wanted = [product["concept_id"] for product in issue["products"]]
    if set(supplied) != set(wanted):
        raise SafetyError("LIVE_FACTS_CONCEPT_MISMATCH")
    products = []
    declared = issue_fit_policy(issue)
    for rank, frozen in enumerate(issue["products"], 1):
        value = deepcopy(supplied[frozen["concept_id"]])
        value.update(concept_id=frozen["concept_id"], rank=rank, category=frozen["category"], status=frozen["status"],
                     kpi_states=deepcopy(frozen["kpi_states"]), kpi_matches=frozen["kpi_states"].count("Meets"))
        if declared == FIT_POLICY_ID:
            for field in ("fit_policy_id", "criterion_ids", "kpi_total", "fit_tier", "source_revisions", "brand_binding"):
                value[field] = deepcopy(frozen[field])
        value.pop("fit_description_contract", None)
        if "fit_description_contract" in frozen:
            value["fit_description_contract"] = frozen["fit_description_contract"]
        for field in (*FUNCTIONAL_REPORT_FIELDS, *OPTIONAL_FUNCTIONAL_REPORT_FIELDS, "display_asset"):
            value.pop(field, None)
            if field in frozen:
                value[field] = deepcopy(frozen[field])
        for field in OPPORTUNITY_FIELDS:
            value.pop(field, None)
            if field in frozen:
                value[field] = deepcopy(frozen[field])
        if "image_provenance" in frozen:
            value["image_provenance"] = deepcopy(frozen["image_provenance"])
            value["image_path"] = frozen["image_provenance"]["path"]
        elif frozen.get("opportunity_policy_id") == SCORE_POLICY_ID and (
                value.get("image_path") is not None or value.get("image_provenance") is not None):
            raise SafetyError("LIVE_PHOTO_BINDING_INVALID")
        for field in SOURCE_ROUTE_FIELDS:
            value.pop(field, None)
            if field in frozen:
                value[field] = deepcopy(frozen[field])
        value.pop("supplier_terms_display_version", None)
        if "supplier_terms_display_version" in frozen:
            value["supplier_terms_display_version"] = frozen["supplier_terms_display_version"]
        # Update admission belongs to the frozen selection, never supplied prose.
        for field in ("material_update_reviewed", "update_reason"):
            value.pop(field, None)
            if field in frozen:
                value[field] = deepcopy(frozen[field])
        for field in ("market_checks", "prices", "suppliers", "source_refs"):
            value.setdefault(field, [])
        value.setdefault("image_path", None)
        value.setdefault("supplier_search_status", "not_started")
        products.append(value)
    dataset = {"schema_version": 1, "evidence_type": evidence_type, "issue_id": issue["issue_id"], "revision": issue["revision"],
               "release_date": issue["date"], "selection_hash": issue["selection_hash"], "products": products,
               "narratives": deepcopy(narratives or {}), "selected_language": None}
    if declared == FIT_POLICY_ID:
        dataset.update(schema_version=2, fit_policy_id=declared, criterion_ids=list(policy_for(declared)["active_ids"]), kpi_total=10)
    if "opportunity_policy_id" in issue:
        dataset["opportunity_policy_id"] = issue["opportunity_policy_id"]
    if "recommendation_contract" in issue:
        if issue["recommendation_contract"] != RECOMMENDATION_CONTRACT or declared != FIT_POLICY_ID:
            raise SafetyError("LIVE_RECOMMENDATION_CONTRACT_INVALID")
        dataset.update(schema_version=3, recommendation_contract=RECOMMENDATION_CONTRACT,
                       report_asset_receipt=deepcopy(issue.get("report_asset_receipt")))
    validate_dataset(issue, dataset)
    return dataset


def validate_dataset(issue, dataset, *, root=None):
    if (not isinstance(dataset, dict) or type(dataset.get("schema_version")) is not int or dataset.get("schema_version") not in (1, 2, 3) or dataset.get("evidence_type") not in ("captured", "synthetic")
            or type(dataset.get("revision")) is not int or type(issue.get("revision")) is not int or issue["revision"] < 1
            or any(dataset.get(field) != issue.get(field) for field in ("issue_id", "revision", "selection_hash"))
            or dataset.get("release_date") != issue["date"]):
        raise SafetyError("LIVE_DATASET_ISSUE_MISMATCH")
    functional = dataset["schema_version"] == 3
    if functional:
        if (dataset.get("recommendation_contract") != RECOMMENDATION_CONTRACT
                or issue.get("recommendation_contract") != RECOMMENDATION_CONTRACT
                or dataset.get("report_asset_receipt") != issue.get("report_asset_receipt")):
            raise SafetyError("LIVE_RECOMMENDATION_CONTRACT_INVALID")
    elif any(field in dataset or field in issue for field in ("recommendation_contract", "report_asset_receipt")):
        raise SafetyError("LIVE_RECOMMENDATION_CONTRACT_INVALID")
    declared = FIT_POLICY_ID if dataset["schema_version"] in (2, 3) else LEGACY_POLICY_ID
    try:
        if issue_fit_policy(issue) != declared:
            raise SafetyError("LIVE_FIT_POLICY_MISMATCH")
    except SafetyError:
        raise SafetyError("LIVE_FIT_POLICY_MISMATCH") from None
    if declared == FIT_POLICY_ID:
        expected_ids = list(policy_for(declared)["active_ids"])
        if (dataset.get("fit_policy_id") != declared or issue.get("fit_policy_id") != declared
                or dataset.get("criterion_ids") != expected_ids or issue.get("criterion_ids") != expected_ids
                or any(type(number) is not int for number in dataset["criterion_ids"])
                or type(dataset.get("kpi_total")) is not int or dataset["kpi_total"] != 10 or issue.get("kpi_total") != 10):
            raise SafetyError("LIVE_FIT_POLICY_MISMATCH")
    elif any(field in dataset or field in issue for field in ("fit_policy_id", "criterion_ids", "kpi_total")):
        raise SafetyError("LIVE_FIT_POLICY_MISMATCH")
    if dataset.get("opportunity_policy_id") != issue.get("opportunity_policy_id"):
        raise SafetyError("LIVE_OPPORTUNITY_POLICY_MISMATCH")
    # The allocator determines history flags; the renderer does not reinterpret them.
    updated = [product["concept_id"] for product in issue["products"] if product["status"] == "Updated"]
    counts = validate_issue(issue["products"], updated, policy=policy_for_fit(declared))
    if issue.get("counts") != counts or issue.get("ready_for_review") is not True:
        raise SafetyError("LIVE_INCOMPLETE_ISSUE")
    from .issue_allocation import selection_fingerprint
    frozen_hash = selection_fingerprint(issue)
    if frozen_hash != issue["selection_hash"]:
        raise SafetyError("LIVE_ISSUE_FACTS_CHANGED")
    products = dataset.get("products")
    if not isinstance(products, list) or any(not isinstance(p, dict) for p in products) or [p.get("concept_id") for p in products] != [p["concept_id"] for p in issue["products"]]:
        raise SafetyError("LIVE_PRODUCT_ORDER_MISMATCH")
    for rank, (value, frozen) in enumerate(zip(products, issue["products"]), 1):
        if type(value.get("rank")) is not int or type(value.get("kpi_matches")) is not int:
            raise SafetyError("LIVE_PROTECTED_FACT_CHANGED")
        if any(value.get(field) != frozen.get(field) for field in ("concept_id", "category", "status", "kpi_states", "material_update_reviewed", "update_reason")) or value.get("rank") != rank or value.get("kpi_matches") != frozen["kpi_states"].count("Meets"):
            raise SafetyError("LIVE_PROTECTED_FACT_CHANGED")
        if any(type(value.get(field)) is not type(frozen.get(field)) for field in ("material_update_reviewed", "update_reason")):
            raise SafetyError("LIVE_PROTECTED_FACT_CHANGED")
        _text(value.get("title"))
        _text(value.get("market_label"))
        if functional:
            if (any(field not in frozen or field not in value or value[field] != frozen[field]
                    for field in (*functional_report_fields(frozen), "display_asset"))
                    or any((field in value) != (field in frozen) for field in OPTIONAL_FUNCTIONAL_REPORT_FIELDS)
                    or value["title"] != value["functional_definition"]["report_title"]):
                raise SafetyError("LIVE_FUNCTIONAL_FACTS_CHANGED")
            validate_functional_examples(value, concept_id=value["concept_id"])
            from .market_validation import validate_frozen_functional_market
            from .report_selection import validate_frozen_functional_fit
            from .report_examples import validate_frozen_functional_photo
            from .report_source_facts import validate_functional_route_proof
            validate_frozen_functional_fit(value)
            validate_frozen_functional_market(value)
            from .report_visuals import validate_quick_commerce_checks
            if validate_quick_commerce_checks(value["quick_commerce_checks"])["as_of"] != value["functional_market_proof"]["as_of"]:
                raise SafetyError("LIVE_QUICK_COMMERCE_CHECKS_INVALID")
            validate_frozen_functional_photo(value)
            validate_functional_route_proof(value)
        elif any(field in value or field in frozen for field in (*FUNCTIONAL_REPORT_FIELDS, *OPTIONAL_FUNCTIONAL_REPORT_FIELDS, "display_asset")):
            raise SafetyError("LIVE_RECOMMENDATION_CONTRACT_INVALID")
        if (value.get("fit_description_contract") != frozen.get("fit_description_contract")
                or ("fit_description_contract" in value) != ("fit_description_contract" in frozen)):
            raise SafetyError("LIVE_FIT_DESCRIPTION_BINDING_CHANGED")
        if (("supplier_terms_display_version" in value) != ("supplier_terms_display_version" in frozen)
                or value.get("supplier_terms_display_version") != frozen.get("supplier_terms_display_version")):
            raise SafetyError("LIVE_SUPPLIER_TERMS_BINDING_CHANGED")
        if "supplier_terms_display_version" in value:
            if (value["supplier_terms_display_version"] != SUPPLIER_TERMS_DISPLAY_VERSION
                    or value.get("opportunity_policy_id") != SCORE_POLICY_ID
                    or value.get("source_route_contract") not in SOURCE_ROUTE_CONTRACTS):
                raise SafetyError("LIVE_SUPPLIER_TERMS_DISPLAY_VERSION_INVALID")
        if "fit_description_contract" in value:
            from .fit_description import DESCRIPTION_INFERENCE_KIND, FUNCTIONAL_INFERENCE_KIND
            expected_contract = FUNCTIONAL_INFERENCE_KIND if functional else DESCRIPTION_INFERENCE_KIND
            if declared != FIT_POLICY_ID or value["fit_description_contract"] != expected_contract:
                raise SafetyError("LIVE_FIT_DESCRIPTION_CONTRACT_INVALID")
            if (not isinstance(value.get("criteria"), list)
                    or any(not isinstance(row, dict) or "description_inference" not in row for row in value["criteria"])):
                raise SafetyError("LIVE_FIT_DESCRIPTION_CONTRACT_INVALID")
        if frozen.get("opportunity_policy_id") == SCORE_POLICY_ID and "image_provenance" not in frozen:
            if value.get("image_path") is not None or value.get("image_provenance") is not None:
                raise SafetyError("LIVE_PHOTO_BINDING_INVALID")
        if any(field in frozen or field in value for field in SOURCE_ROUTE_FIELDS):
            if (frozen.get("opportunity_policy_id") != SCORE_POLICY_ID
                    or any(value.get(field) != frozen.get(field) for field in SOURCE_ROUTE_FIELDS)):
                raise SafetyError("LIVE_SOURCE_ROUTE_FACTS_CHANGED")
            validate_route_facts(value)
            if (value.get("source_route_contract") == FUNCTIONAL_SOURCE_ROUTE_CONTRACT) != functional:
                raise SafetyError("LIVE_SOURCE_ROUTE_CONTRACT_INVALID")
            for check in value["source_route_checks"]:
                _url(check.get("url"))
        if "image_provenance" in frozen:
            photo = value.get("image_provenance")
            if (photo != frozen["image_provenance"] or not isinstance(photo, dict)
                    or photo.get("owner_id") != frozen["concept_id"] or photo.get("source_kind") != "public_product_image"
                    or photo.get("relationship") not in ("exact_product", "relevant_family")
                    or photo.get("path") != value.get("image_path")
                    or photo.get("id") not in frozen.get("source_revisions", {}).get("source_checks", {})):
                raise SafetyError("LIVE_PHOTO_BINDING_INVALID")
            if not functional:
                from .product_images import MEDIA_ROOT, inspect_image
                from .storage import safe_path
                path = safe_path(photo["path"])
                if (not path.is_relative_to(safe_path(MEDIA_ROOT)) or not path.is_file()
                        or path.stat().st_size > 5 * 1024 * 1024
                        or inspect_image(path.read_bytes())["sha256"] != photo["sha256"]):
                    raise SafetyError("LIVE_PHOTO_BYTES_CHANGED")
            _url(photo["source_url"]); timestamp(photo["observed_at"])
        if declared == FIT_POLICY_ID:
            product_fit(value, declared)
            if any(value.get(field) != frozen.get(field) for field in OPPORTUNITY_FIELDS):
                raise SafetyError("LIVE_PROTECTED_FACT_CHANGED")
            if any(value.get(field) != frozen.get(field) for field in ("fit_policy_id", "criterion_ids", "kpi_total", "fit_tier", "source_revisions", "brand_binding")):
                raise SafetyError("LIVE_PROTECTED_FACT_CHANGED")
            binding = value.get("brand_binding")
            if (not isinstance(binding, dict) or binding.get("verdict") not in ("generic_or_small", "unclear")
                    or binding.get("source") not in ("original", "overlay", "operator", "origin")
                    or binding.get("description_hash") != digest(value.get("description"))):
                raise SafetyError("LIVE_BRAND_BINDING_INVALID")
            if binding["source"] in ("overlay", "origin") and binding.get("review_id") not in value["source_revisions"].get("brand_reviews", {}):
                raise SafetyError("LIVE_BRAND_BINDING_INVALID")
            fit_context = ({"description": value.get("fit_description") if functional else value.get("description"),
                            "description_contract": value["fit_description_contract"]}
                           if "fit_description_contract" in value else {})
            fit = evaluate_fit(value.get("criteria"), declared, **fit_context)
            if ([row["state"] for row in value["criteria"]] != value["kpi_states"] or fit["fit_tier"] != value["fit_tier"]):
                raise SafetyError("LIVE_CHECKLIST_FACT_MISMATCH")
            known = {ref.get("id") for ref in value.get("source_refs", []) if isinstance(ref, dict)}
            if any(not set(row["references"]) <= known for row in value["criteria"]):
                raise SafetyError("LIVE_CHECKLIST_REFERENCE_INVALID")
        if value.get("supplier_search_status") not in ("not_started", "attempted", "complete", "blocked"):
            raise SafetyError("LIVE_SUPPLIER_SEARCH_STATE_INVALID")
        if declared == FIT_POLICY_ID and not frozen_supplier_effort_valid(value):
            raise SafetyError("LIVE_SUPPLIER_EFFORT_REQUIRED")
        for field in ("market_checks", "prices", "suppliers", "source_refs"):
            if not isinstance(value.get(field), list) or any(not isinstance(record, dict) for record in value[field]):
                raise SafetyError("LIVE_FACT_LIST_INVALID")
        for check in value["market_checks"]:
            for field in ("source", "state", "detail"):
                _text(check.get(field))
            _url(check.get("url"))
            if check.get("observed_at"):
                timestamp(check["observed_at"])
        for price in value["prices"]:
            _price(price)
        for supplier in value["suppliers"]:
            identifier(supplier.get("lead_id"))
            _text(supplier.get("name"))
            _url(supplier.get("url"))
            if supplier.get("single_piece_dispatch") not in ("supported", "not_supported", "not_stated") or supplier.get("subscription") not in ("no_subscription_stated", "subscription_stated", "not_stated"):
                raise SafetyError("LIVE_SUPPLIER_STATE_INVALID")
            if supplier.get("min_order") is not None and (type(supplier["min_order"]) is not int or supplier["min_order"] < 1):
                raise SafetyError("LIVE_SUPPLIER_STATE_INVALID")
            for field in ("ecommerce_support", "reason"):
                _text(supplier.get(field, ""), optional=True)
        for ref in value["source_refs"]:
            identifier(ref.get("id"))
            _url(ref.get("url"))
            _text(ref.get("source"))
            timestamp(ref["observed_at"])
    if functional:
        from .report_assets import validate_packet_assets
        validate_packet_assets(issue, products, dataset["report_asset_receipt"], root=root)
    if dataset.get("selected_language") not in (None, *LANGUAGES):
        raise SafetyError("REPORT_LANGUAGE_DECISION_REQUIRED")
    if not isinstance(dataset.get("narratives"), dict):
        raise SafetyError("LIVE_NARRATIVE_INVALID")
    for language, narrative in dataset["narratives"].items():
        _narrative(narrative, products, language)
        if functional:
            from .report_narration import validate_functional_writing
            validate_functional_writing(dataset, language, narrative)
    return dataset


def _narrative(narrative, products, language):
    if language not in LANGUAGES or not isinstance(narrative, dict):
        raise SafetyError("LIVE_NARRATIVE_INVALID")
    _text(narrative.get("title"))
    rows = narrative.get("products")
    if not isinstance(rows, list) or any(not isinstance(row, dict) for row in rows) or [row.get("concept_id") for row in rows] != [p["concept_id"] for p in products]:
        raise SafetyError("LIVE_NARRATIVE_CONCEPT_MISMATCH")
    for row, product in zip(rows, products):
        if product.get("fit_policy_id") == FIT_POLICY_ID:
            phrases = row.get("checklist_phrases")
            ids = list(policy_for(FIT_POLICY_ID)["active_ids"])
            if (not isinstance(phrases, list) or len(phrases) != 10
                    or any(not isinstance(phrase, dict) or type(phrase.get("id")) is not int for phrase in phrases)
                    or [phrase.get("id") for phrase in phrases if isinstance(phrase, dict)] != ids):
                raise SafetyError("LIVE_CHECKLIST_PHRASES_INVALID")
            for phrase in phrases:
                if set(phrase) != {"id", "text"}:
                    raise SafetyError("LIVE_CHECKLIST_PHRASES_INVALID")
                from .report_narration import WRITING_LIMITS, _plain
                try:
                    _plain(phrase["text"], "checklist", WRITING_LIMITS[language]["checklist"])
                    if not phrase["text"].strip() or language == "hi-Latn" and re.search(r"[\u0900-\u097f]", phrase["text"]):
                        raise SafetyError("LIVE_CHECKLIST_PHRASES_INVALID")
                except SafetyError:
                    raise SafetyError("LIVE_CHECKLIST_PHRASES_INVALID") from None
        if product["status"] == "New" and row.get("update_sentence", ""):
            raise SafetyError("LIVE_NARRATIVE_UPDATE_NOT_ADMITTED")
        for field in ("reason", "limitation", "update_sentence"):
            _text(row.get(field, ""), optional=field != "reason")
        from .report_narration import LIMITATION_LIMIT, WRITING_LIMITS
        limits = WRITING_LIMITS[language]
        if (len(row["reason"]) > limits["reason"] or len(row.get("update_sentence") or "") > limits["update_sentence"]
                or len(row.get("limitation") or "") > LIMITATION_LIMIT
                or isinstance(row.get("points"), list) and any(isinstance(point, str) and len(point) > limits["point"]
                                                               for point in row["points"])):
            raise SafetyError("LIVE_NARRATIVE_TOO_LONG")
        if not isinstance(row.get("points"), list) or len(row["points"]) > 3:
            raise SafetyError("LIVE_NARRATIVE_INVALID")
        allowed = {"title", "rank", "kpi_matches", "date", "retail_price", "supplier_price"}
        for text in [row["reason"], *row["points"], row.get("limitation", ""), row.get("update_sentence", "")]:
            _text(text, optional=True)
            placeholders = re.findall(r"\{\{(.*?)\}\}", text)
            if any(value not in allowed for value in placeholders) or re.search(r"\d", re.sub(r"\{\{.*?\}\}", "", text)):
                raise SafetyError("LIVE_NARRATIVE_UNPROTECTED_NUMBER")


def _prose(text, product, dataset, language="en"):
    prices = product["prices"]
    values = {"title": product["title"], "rank": str(product["rank"]), "kpi_matches": str(product["kpi_matches"]),
              "date": dataset["release_date"],
              "retail_price": next((price_display(p, language) for p in prices if p["kind"] == "retail"), "Retail price not stated" if language == "en" else "Retail keemat nahi di gayi"),
              "supplier_price": next((price_display(p, language) for p in prices if p["kind"] == "supplier"), "Supplier price not stated" if language == "en" else "Supplier ki keemat nahi di gayi")}
    return re.sub(r"\{\{(.*?)\}\}", lambda match: values[match.group(1)], text)


def _image(path):
    if path is None:
        return None
    path = Path(path)
    original = path.absolute()
    safe_path(original.parent)
    if original.is_symlink():
        raise SafetyError("LIVE_IMAGE_INVALID")
    resolved = original.resolve()
    logical = logical_path(ROOT, resolved)
    if logical is None or logical.parts[:1] != ("media",):
        raise SafetyError("LIVE_IMAGE_OUTSIDE_MEDIA")
    safe_path(resolved.parent)
    if resolved.is_symlink() or resolved.suffix.casefold() not in (".png", ".jpg", ".jpeg", ".webp"):
        raise SafetyError("LIVE_IMAGE_INVALID")
    if not resolved.exists():
        return None
    with resolved.open("rb") as stream:
        raw = stream.read(8 * 1024 * 1024 + 1)
    if len(raw) > 8 * 1024 * 1024:
        raise SafetyError("LIVE_IMAGE_LIMIT")
    from PIL import Image
    try:
        with Image.open(BytesIO(raw)) as image:
            image.verify()
    except (OSError, ValueError):
        raise SafetyError("LIVE_IMAGE_INVALID") from None
    mime = "image/jpeg" if resolved.suffix.casefold() in (".jpg", ".jpeg") else "image/" + resolved.suffix[1:].casefold()
    return "data:" + mime + ";base64," + b64encode(raw).decode("ascii")


def display_market_checks(product):
    """One shared list for both PDF languages and the evidence worksheet.

    Historical packets retain their original list. New packets add the saved
    route status and numbers without replacing their measured/AI support rows.
    These factual blocks are code-owned; G need not rewrite fourteen statuses.
    """
    return product["market_checks"] + (product["source_route_checks"]
        if product.get("source_route_contract") in SOURCE_ROUTE_CONTRACTS else [])


def _owned_asset(dataset, owner_id, *, root, as_bytes=False):
    from .report_assets import resolve_report_asset
    receipt = dataset["report_asset_receipt"]
    entry = next((row for row in receipt["plan"]["entries"] if row["source"]["owner_id"] == owner_id), None)
    if entry is None or entry["state"] == "omitted":
        return None
    saved = resolve_report_asset(receipt, entry["asset_id"], root=root,
                                 issue_id=dataset["issue_id"], revision=dataset["revision"])
    return saved["bytes"] if as_bytes else saved["data_url"]


def workbook_contract(dataset):
    if dataset.get("schema_version") == 3:
        return TYPE_WORKBOOK_CONTRACT
    return "fit-checklist-v1" if dataset.get("fit_policy_id") == FIT_POLICY_ID else "legacy-checklist-v1"


def _generation_version(dataset):
    return TYPE_WORKBOOK_GENERATION_VERSION if dataset.get("schema_version") == 3 else (
        SOURCE_WORKBOOK_GENERATION_VERSION if _has_source_routes(dataset) else WORKBOOK_GENERATION_VERSION)


def _render_version(dataset):
    return TYPE_HTML_RENDER_VERSION if dataset.get("schema_version") == 3 else (
        SOURCE_HTML_RENDER_VERSION if _has_source_routes(dataset) else HTML_RENDER_VERSION)


def _supported_render_versions(dataset):
    # The original saved HTML is immutable. New editions use the repaired
    # print layout; old exact-byte stages/manifests retain their own version.
    previous = LEGACY_TYPE_HTML_RENDER_VERSION if dataset.get("schema_version") == 3 else (
        LEGACY_SOURCE_HTML_RENDER_VERSION if _has_source_routes(dataset) else LEGACY_HTML_RENDER_VERSION)
    return (_render_version(dataset), previous)


def render_html(issue, dataset, *, language="en", root=None):
    """One self-contained HTML edition (no script, no external resource).

    The packet is validated first; the layout itself lives in ``report_view``
    (pages, charts, theme and fixed wording in their own small modules).
    """
    validate_dataset(issue, dataset, root=root or ROOT)
    if language not in LANGUAGES:
        raise SafetyError("LIVE_LANGUAGE_INVALID")
    narrative = dataset["narratives"].get(language)
    if not narrative:
        raise SafetyError("LIVE_NARRATIVE_LANGUAGE_MISSING")
    functional = dataset["schema_version"] == 3
    logo = _owned_asset(dataset, "brand-logo", root=root or ROOT) if functional else _image(LOGO)
    images = [_owned_asset(dataset, product["concept_id"], root=root or ROOT) if functional else _image(product["image_path"])
              for product in dataset["products"]]
    from .report_view import render_document
    html = render_document(issue, dataset, narrative, language, logo=logo, images=images)
    # The same refusal the native PDF renderer applies, plus: embedded resources
    # are data only (links may open public pages; nothing is fetched to print).
    if (re.search(r"<(?:script|iframe|object|embed)\b", html, re.I)
            or re.search(r"\bsrc=(?![\"']?data:)", html, re.I) or re.search(r"url\((?![\"']?data:)", html, re.I)
            or re.search(r"\bhref=(?![\"']?(?:https?://|#))", html, re.I)):
        raise SafetyError("LIVE_RENDER_HTML_INVALID")
    return html


def _safe_cell(value):
    if isinstance(value, str) and value.lstrip(" \t\r\n").startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _row_height(values, widths):
    """Conservative wrapped-text estimate, within Excel's 409-point limit.

    Long content remains in the original cell even when the maximum height is
    reached; nothing is shortened to make the visual row fit.
    """
    lines = 1
    for value, width in zip(values, widths):
        capacity = max(1, int(width * .9))
        count = sum(max(1, len(textwrap.wrap(line, width=capacity, replace_whitespace=False)))
                    for line in str(value if value is not None else "").split("\n"))
        lines = max(lines, count)
    return min(409, max(42, lines * 14 + 10))


def evidence_workbook_rows(product):
    """Exact Evidence rows shared with the downstream delivery checker."""
    checks = display_market_checks(product)
    rows = [[product["rank"], product["title"], check["source"], check["state"], check["detail"],
             check.get("observed_at") or "Not stated", check.get("url") or "Not stated", product["concept_id"]]
            for check in checks]
    used = {(check.get("url"), check.get("observed_at")) for check in checks}
    rows.extend([product["rank"], product["title"], ref["source"], "Supporting source", "Saved source reference " + ref["id"],
                 ref["observed_at"], ref.get("url") or "Not stated", product["concept_id"]]
                for ref in product["source_refs"] if (ref.get("url"), ref["observed_at"]) not in used)
    return rows


def supplier_workbook_rows(product):
    """Every accepted lead, with unchanged commercial uncertainty/links."""
    from .supplier_scoped_display import is_scoped_supplier, scoped_supplier_lines
    return [[product["rank"], product["title"], supplier["name"],
             "See each offer's dispatch terms" if is_scoped_supplier(supplier) else DISPATCH_LABELS[supplier["single_piece_dispatch"]],
             "See company and offer terms" if is_scoped_supplier(supplier) else
             _supplier_subscription_display(supplier, terms_version=_supplier_terms_version(product)),
             "See each offer's minimum purchase" if is_scoped_supplier(supplier) else _supplier_minimum_order(supplier),
             supplier.get("ecommerce_support") or "Not stated",
             "\n".join(scoped_supplier_lines(supplier)) if is_scoped_supplier(supplier) else
             "\n".join([supplier.get("reason") or "Not stated",
                         supplier_terms_display(supplier, terms_version=_supplier_terms_version(product))]),
             supplier.get("url") or "Not stated", product["concept_id"]] for supplier in product["suppliers"]]


def _workbook_bytes(issue, dataset, *, deterministic=False, root=None):
    style, _ = _style_helpers()
    book = Workbook()
    book.remove(book.active)
    definitions = {
        "Shortlist": (["Rank", "Product", "Category", "Status", "Checklist", "Market", "Retail prices", "Supplier quotes", "Concept ID", "Supplier search"], []),
        "Evidence": (["Rank", "Product", "Source", "State", "Evidence", "Observed", "Link", "Concept ID"], []),
        "Supplier leads": (["Rank", "Product", "Supplier", "Single-piece sourcing and dispatch", "Subscription", "Minimum order", "E-commerce support", "Notes", "Link", "Concept ID"], []),
    }
    widths = {"Shortlist": [8, 36, 24, 12, 12, 28, 48, 70, 46, 22],
              "Evidence": [8, 36, 18, 18, 70, 24, 34, 46],
              "Supplier leads": [8, 36, 28, 42, 40, 18, 40, 80, 34, 46]}
    if dataset.get("fit_policy_id") == FIT_POLICY_ID:
        definitions["Checklist"] = (["Rank", "Product", "Criterion ID", "Criterion", "State", "Basis", "Reason", "References", "Fit", "Concept ID"], [])
        widths["Checklist"] = [8, 36, 12, 30, 20, 18, 65, 50, 16, 46]
    if dataset["schema_version"] == 3:
        definitions["Example products"] = (list(EXAMPLE_WORKBOOK_HEADERS), [])
        widths["Example products"] = [8, 36, 45, 20, 25, 25, 40, 24, 34, 46, 46, 46, 46]
    if "opportunity_policy_id" in dataset:
        definitions["Shortlist"][0].extend(opportunity_workbook_headers())
        widths["Shortlist"].extend([10, 11, 11, 11, 12, 20, 12])
    for product in dataset["products"]:
        prices = product["prices"]
        definitions["Shortlist"][1].append([product["rank"], product["title"], product["category"], product["status"],
            f'{product["kpi_matches"]}/{dataset.get("kpi_total", 13)}', product["market_label"], "\n".join(price_display(p) for p in prices if p["kind"] == "retail") or "Not stated",
            "\n".join(price_display(p) for p in prices if p["kind"] == "supplier") or "Not stated", product["concept_id"],
            supplier_search_display(product)])
        if "opportunity_policy_id" in dataset:
            definitions["Shortlist"][1][-1].extend(opportunity_workbook_values(product))
        if dataset.get("fit_policy_id") == FIT_POLICY_ID:
            from .research_ai import CRITERIA_BY_ID
            for criterion in product["criteria"]:
                definitions["Checklist"][1].append([product["rank"], product["title"], criterion["id"], CRITERIA_BY_ID[criterion["id"]],
                    criterion["state"], criterion["basis"], criterion["reason"], ", ".join(criterion["references"]),
                    "Good fit" if product["fit_tier"] == "good" else "Decent fit", product["concept_id"]])
        definitions["Evidence"][1].extend(evidence_workbook_rows(product))
        definitions["Supplier leads"][1].extend(supplier_workbook_rows(product))
        if dataset["schema_version"] == 3:
            definitions["Example products"][1].extend(example_workbook_rows(product))
    for name, (headings, rows) in definitions.items():
        sheet = book.create_sheet(name)
        logo_bytes = _owned_asset(dataset, "brand-logo", root=root or ROOT, as_bytes=True) if dataset["schema_version"] == 3 else None
        if logo_bytes is not None:
            # Excel cannot reliably display WebP. Read the owned bytes and
            # package a PNG, without touching that immutable display original.
            # Legacy path-based workbooks keep their exact existing behavior.
            from PIL import Image as PillowImage
            with PillowImage.open(BytesIO(logo_bytes)) as image:
                if image.format not in {"PNG", "JPEG", "GIF"}:
                    converted = BytesIO()
                    image.save(converted, format="PNG")
                    logo_bytes = converted.getvalue()
        logo_source = (BytesIO(logo_bytes) if logo_bytes is not None else None) if dataset["schema_version"] == 3 else LOGO
        style.add_brand_header(sheet, title=name, period=dataset["release_date"], purpose="Product research", last_column=len(headings), logo_path=logo_source)
        for column, title in enumerate(headings, 1):
            style.set_cell(sheet, 6, column, title, font=Font(name="Aptos", size=10, bold=True, color=style.CREAM), fill=style.solid(style.NAVY), alignment=Alignment(wrap_text=True, vertical="center"))
            sheet.column_dimensions[sheet.cell(6, column).column_letter].width = widths[name][column - 1]
        sheet.row_dimensions[6].height = 38
        for row_number, values in enumerate(rows, 7):
            for column, value in enumerate(values, 1):
                style.set_cell(sheet, row_number, column, _safe_cell(value), font=Font(name="Aptos", size=10, color=style.CHARCOAL),
                    fill=style.solid(style.CREAM if row_number % 2 else style.LIGHT_GREY), alignment=Alignment(wrap_text=True, vertical="top"))
                if headings[column - 1] in ("Link", "Source link") and isinstance(value, str) and value.startswith(("https://", "http://")):
                    style.add_hyperlink(sheet.cell(row_number, column), value, "Open source")
            displayed = ["Open source" if heading in ("Link", "Source link") and isinstance(value, str) and value.startswith(("http://", "https://")) else value
                         for heading, value in zip(headings, values)]
            sheet.row_dimensions[row_number].height = _row_height(displayed, widths[name])
        from .report_workbook_style import style_sheet
        style_sheet(style, sheet, name, headings, len(rows))
        sheet.freeze_panes = "C7"
        sheet.auto_filter.ref = f"A6:{sheet.cell(max(6, len(rows) + 6), len(headings)).coordinate}"
        style.configure_page(sheet, len(headings), len(rows) + 6, 1)
    stream = BytesIO()
    book.save(stream)
    raw = stream.getvalue()
    if not deterministic:
        # Preserve original unstaged packets and their historical byte contract.
        return raw
    # XLSX contains both ZIP timestamps and wall-clock core properties. Pin
    # those for new checked stages so a crash before intent can be reconciled
    # against the exact expected bytes, rather than adopting an arbitrary file.
    output = BytesIO()
    fixed_date = dataset["release_date"] + "T00:00:00Z"
    with ZipFile(BytesIO(raw)) as source, ZipFile(output, "w", compression=ZIP_DEFLATED, compresslevel=9) as target:
        for name in sorted(source.namelist()):
            body = source.read(name)
            if name == "docProps/core.xml":
                body = re.sub(rb"(<dcterms:(?:created|modified)\b[^>]*>)[^<]*(</dcterms:(?:created|modified)>)",
                              lambda match: match[1] + fixed_date.encode("ascii") + match[2], body)
            entry = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            entry.create_system = 3
            entry.external_attr = 0o600 << 16
            target.writestr(entry, body, compress_type=ZIP_DEFLATED, compresslevel=9)
    return output.getvalue()


def _stage_paths(issue, destination):
    root = safe_path(destination)
    issue_id, revision = identifier(issue["issue_id"]), identifier("revision-" + str(issue["revision"]))
    data = safe_path(runtime_path(root, "data", "outputs"), issue_id, revision)
    media = safe_path(runtime_path(root, "media", "live", "reports"), issue_id, revision)
    return data, media, safe_path(data, "internal-stages"), safe_path(media, "internal-stages")


def _stage_read(path, limit=MAX_STAGE_FILE):
    path = Path(path)
    safe_path(path)
    try:
        with path.open("rb") as stream:
            raw = stream.read(limit + 1)
    except OSError:
        raise SafetyError("LIVE_STAGE_FILE_MISSING") from None
    if not raw or len(raw) > limit:
        raise SafetyError("LIVE_STAGE_FILE_SIZE_INVALID")
    return raw


def _stage_exact(path, raw):
    safe_path(path)
    if path.exists():
        if _stage_read(path) != raw:
            raise SafetyError("LIVE_STAGE_ARTIFACT_CHANGED")
    else:
        atomic_bytes(path, raw)


def _stage_binding(issue, core, kind):
    binding = {"schema_version": 1, "stage_version": INTERNAL_STAGE_VERSION, "kind": kind,
               "issue_id": issue["issue_id"], "revision": issue["revision"], "issue_hash": digest(issue),
               "core_facts_hash": digest(core), "selection_hash": issue["selection_hash"]}
    if kind in ("workbook", "edition"):
        binding["price_display_version"] = PRICE_DISPLAY_VERSION
        if core.get("schema_version") == 3:
            entry = next((row for row in core["report_asset_receipt"]["plan"]["entries"]
                          if row["source"]["kind"] == "brand_logo"), None)
            binding["logo_sha256"] = entry["display"]["sha256"] if entry and entry["state"] == "ready" else None
            binding["asset_receipt_hash"] = digest(core["report_asset_receipt"])
            binding["recommendation_contract"] = RECOMMENDATION_CONTRACT
        else:
            binding["logo_sha256"] = digest_bytes(_stage_read(LOGO))
        binding["workbook_contract"] = workbook_contract(core)
    if kind == "workbook":
        binding.update(generation_version=_generation_version(core),
                       style_sha256=digest_bytes(_stage_read(ROOT / "scripts" / "generate_sample_excel.py")))
    elif kind == "edition":
        binding.update(render_version=_render_version(core),
                       style_sha256=digest_bytes(_stage_read(ROOT / "scripts" / "generate_sample_pdf.py")))
    return binding


def _has_source_routes(dataset):
    return any(product.get("source_route_contract") in SOURCE_ROUTE_CONTRACTS for product in dataset["products"])


def _stage_saved(stages, name, binding, specifications):
    """Verify a ready receipt, or finish its original exact-byte save intent."""
    intent_path, receipt_path = stages / (name + "-intent.json"), stages / (name + "-receipt.json")
    safe_path(intent_path); safe_path(receipt_path)
    if not intent_path.exists():
        if receipt_path.exists():
            raise SafetyError("LIVE_STAGE_INTENT_REQUIRED")
        return None
    intent = load_json(intent_path, 128 * 1024)
    if (not isinstance(intent, dict) or set(intent) != {"binding", "artifacts"}
            or canonical(intent["binding"]) != canonical(binding) or not isinstance(intent["artifacts"], dict)
            or set(intent["artifacts"]) != set(specifications)):
        raise SafetyError("LIVE_STAGE_BINDING_CHANGED")
    raw_by_field = {}
    for field, (path, payload, limit) in specifications.items():
        record = intent["artifacts"][field]
        if (not isinstance(record, dict) or set(record) != {"path", "payload_path", "sha256", "bytes"}
                or record["path"] != str(path) or record["payload_path"] != str(payload)
                or not isinstance(record["sha256"], str) or not re.fullmatch(r"[a-f0-9]{64}", record["sha256"])
                or type(record["bytes"]) is not int or not 1 <= record["bytes"] <= limit):
            raise SafetyError("LIVE_STAGE_ARTIFACT_INVALID")
        raw = _stage_read(payload, limit)
        if len(raw) != record["bytes"] or digest_bytes(raw) != record["sha256"]:
            raise SafetyError("LIVE_STAGE_ARTIFACT_CHANGED")
        raw_by_field[field] = raw
    receipt = dict(binding, complete=True, approved=False, release_ready=False,
                   intent_sha256=digest(intent), artifacts={field: {key: value for key, value in record.items() if key != "payload_path"}
                                                            for field, record in intent["artifacts"].items()})
    if receipt_path.exists():
        if canonical(load_json(receipt_path, 128 * 1024)) != canonical(receipt):
            raise SafetyError("LIVE_STAGE_RECEIPT_CHANGED")
        # A completed artifact's disappearance/corruption is not a failed-save
        # retry. Do not silently replace material that was previously accepted.
        for field, (path, _, limit) in specifications.items():
            if _stage_read(path, limit) != raw_by_field[field]:
                raise SafetyError("LIVE_STAGE_ARTIFACT_CHANGED")
        return receipt
    for field, (path, _, _) in specifications.items():
        _stage_exact(path, raw_by_field[field])
    atomic_json(receipt_path, receipt)
    return receipt


def _stage_commit(stages, name, binding, specifications, raw_by_field):
    """The complete payload is durable before intent; intent before aliases."""
    saved = _stage_saved(stages, name, binding, specifications)
    if saved is not None:
        return saved
    artifacts = {}
    for field, (path, payload, limit) in specifications.items():
        raw = raw_by_field[field]
        if not isinstance(raw, bytes) or not 1 <= len(raw) <= limit:
            raise SafetyError("LIVE_STAGE_FILE_SIZE_INVALID")
        _stage_exact(payload, raw)
        artifacts[field] = {"path": str(path), "payload_path": str(payload), "sha256": digest_bytes(raw), "bytes": len(raw)}
    atomic_json(stages / (name + "-intent.json"), {"binding": binding, "artifacts": artifacts})
    return _stage_saved(stages, name, binding, specifications)


def _checked_core(issue, dataset, stages):
    core = {key: deepcopy(value) for key, value in dataset.items() if key not in ("narratives", "selected_language")}
    validate_dataset(issue, dict(core, narratives={}, selected_language=None))
    specifications = {"facts": (stages / "checked-product-facts.json", stages / "checked-product-facts.payload.json", MAX_CORE_FILE)}
    binding = _stage_binding(issue, core, "core")
    saved = _stage_saved(stages, "core", binding, specifications)
    if saved is None:
        saved = _stage_commit(stages, "core", binding, specifications, {"facts": canonical(core) + b"\n"})
    return core, saved


def _legacy_workbook_origin(issue, core, data, media):
    """Adopt only exact bytes proved by an original same-revision manifest.

    A preexisting workbook by itself is not provenance. Historical final facts
    remain untouched; a core match does not authorize adding new prose to them.
    """
    candidates = [(language, data / ("manifest-" + language + ".json")) for language in sorted(LANGUAGES)]
    candidates = [(language, path) for language, path in candidates if path.exists()]
    if not candidates:
        return None
    facts = load_json(data / "final_report_data.json", MAX_CORE_FILE)
    validate_dataset(issue, facts)
    if digest({key: value for key, value in facts.items() if key not in ("narratives", "selected_language")}) != digest(core):
        raise SafetyError("LIVE_IMMUTABLE_FACTS_CHANGED")
    raw_workbook, origin = None, None
    for language, manifest_path in candidates:
        prior = load_json(manifest_path, 128 * 1024)
        if not isinstance(prior, dict):
            raise SafetyError("LIVE_PACKAGE_INTEGRITY_CHANGED")
        if prior.get("price_display_version") != PRICE_DISPLAY_VERSION:
            raise SafetyError("LIVE_RENDER_VERSION_CHANGED")
        if (prior.get("issue_id") != issue["issue_id"] or type(prior.get("revision")) is not int
                or prior["revision"] != issue["revision"] or prior.get("language") != language
                or prior.get("facts_hash") != digest(facts) or prior.get("selection_hash") != issue["selection_hash"]
                or prior.get("approved") is not False or prior.get("release_ready") is not False):
            raise SafetyError("LIVE_PACKAGE_INTEGRITY_CHANGED")
        expected = {"schema_version": 2 if facts["schema_version"] == 3 else 1, "evidence_type": facts["evidence_type"], "counts": issue["counts"],
                    "concept_ids": [product["concept_id"] for product in facts["products"]],
                    "projected_ordinal": issue["projected_ordinal"], "language_state": issue["language_state"]}
        if facts.get("fit_policy_id") == FIT_POLICY_ID:
            expected.update(fit_policy_id=FIT_POLICY_ID, criterion_ids=facts["criterion_ids"], kpi_total=10,
                            workbook_contract=workbook_contract(facts))
        if facts["schema_version"] == 3:
            expected.update(recommendation_contract=RECOMMENDATION_CONTRACT,
                            asset_receipt_hash=digest(facts["report_asset_receipt"]))
        if facts.get("fit_policy_id") != FIT_POLICY_ID and any(key in prior for key in ("fit_policy_id", "criterion_ids", "kpi_total", "workbook_contract")):
            raise SafetyError("LIVE_PACKAGE_INTEGRITY_CHANGED")
        if "opportunity_policy_id" in facts:
            expected["opportunity_policy_id"] = facts["opportunity_policy_id"]
        elif "opportunity_policy_id" in prior:
            raise SafetyError("LIVE_PACKAGE_INTEGRITY_CHANGED")
        if (any(canonical(prior.get(key)) != canonical(value) for key, value in expected.items())
                or not isinstance(prior.get("artifacts"), dict)):
            raise SafetyError("LIVE_PACKAGE_INTEGRITY_CHANGED")
        if "internal_stage" in prior:
            validate_stage_provenance(facts, language, prior["internal_stage"])
        for field, path, limit in (("facts", data / "final_report_data.json", MAX_CORE_FILE),
                                   ("workbook", data / "shortlist.xlsx", MAX_STAGE_FILE),
                                   ("html", media / ("report-" + language + ".html"), MAX_STAGE_FILE)):
            record = prior.get("artifacts", {}).get(field, {})
            raw = _stage_read(path, limit)
            if not isinstance(record, dict) or record.get("path") != str(path) or record.get("sha256") != digest_bytes(raw):
                raise SafetyError("LIVE_PACKAGE_INTEGRITY_CHANGED")
            if field == "workbook":
                if raw_workbook is not None and raw_workbook != raw:
                    raise SafetyError("LIVE_PACKAGE_INTEGRITY_CHANGED")
                raw_workbook = raw
        if origin is None:
            origin = {"manifest_path": str(manifest_path), "manifest_hash": digest(prior),
                      "language": language, "facts_hash": prior["facts_hash"]}
    from .delivery_package import _workbook
    _workbook(raw_workbook, dict(core, narratives={}, selected_language=None))
    return raw_workbook, origin


def write_issue_workbook(issue, dataset, destination):
    """Save checked facts and one narrative-independent internal workbook.

    This is an incomplete, unapproved stage in the original issue/revision;
    it neither freezes final prose nor creates a delivery manifest.
    """
    validate_dataset(issue, dataset, root=destination if dataset.get("schema_version") == 3 else None)
    data, media, stages, _ = _stage_paths(issue, destination)
    with Journal(stages / "stage_writer").locked():
        core, core_receipt = _checked_core(issue, dataset, stages)
        binding = dict(_stage_binding(issue, core, "workbook"), core_receipt_hash=digest(core_receipt))
        legacy = None
        existing_intent = stages / "workbook-intent.json"
        if existing_intent.exists():
            saved_intent = load_json(existing_intent, 128 * 1024)
            if not isinstance(saved_intent, dict) or not isinstance(saved_intent.get("binding"), dict):
                raise SafetyError("LIVE_STAGE_BINDING_CHANGED")
            if saved_intent.get("binding", {}).get("generation_version") == LEGACY_WORKBOOK_GENERATION_VERSION:
                legacy = _legacy_workbook_origin(issue, core, data, media)
                if legacy is None:
                    raise SafetyError("LIVE_STAGE_LEGACY_ORIGIN_MISSING")
        else:
            legacy = _legacy_workbook_origin(issue, core, data, media)
        if legacy is not None:
            binding.update(generation_version=LEGACY_WORKBOOK_GENERATION_VERSION, legacy_origin=legacy[1])
        specifications = {"workbook": (stages / "shortlist.xlsx", stages / "shortlist.payload.xlsx", MAX_STAGE_FILE)}
        saved = _stage_saved(stages, "workbook", binding, specifications)
        if saved is not None:
            return saved
        checked = dict(core, narratives={}, selected_language=None)
        raw = legacy[0] if legacy is not None else _workbook_bytes(issue, checked, deterministic=True, root=destination)
        from .delivery_package import _workbook
        _workbook(raw, checked)
        return _stage_commit(stages, "workbook", binding, specifications, {"workbook": raw})


def write_issue_edition(issue, dataset, destination, language="en"):
    """Save one checked HTML edition independently of the workbook/other G.

    Hinglish binds its actual source English; English does not bind a pending
    or later-completed Hinglish edition. Native PDF ownership stays with the
    existing workflow's saved intent and inspection receipt.
    """
    if language not in LANGUAGES:
        raise SafetyError("LIVE_LANGUAGE_MISSING")
    _, _, stages, media_stages = _stage_paths(issue, destination)
    with Journal(stages / "stage_writer").locked():
        core, core_receipt = _checked_core(issue, dataset, stages)
        narratives = dataset.get("narratives", {})
        if not isinstance(narratives, dict) or language not in narratives:
            raise SafetyError("LIVE_LANGUAGE_MISSING")
        english = narratives.get("en") if language == "hi-Latn" else None
        if language == "hi-Latn" and english is None:
            raise SafetyError("REPORT_SOURCE_ENGLISH_REQUIRED")
        # Validate only this edition's dependencies. An unrelated malformed or
        # unfinished writing result must not stop a healthy sibling's stage.
        own = {language: deepcopy(narratives[language])}
        if english is not None:
            own["en"] = deepcopy(english)
        checked = dict(core, narratives=own, selected_language=None)
        validate_dataset(issue, checked)
        from .report_narration import make_narrative_input
        binding = dict(_stage_binding(issue, core, "edition"), core_receipt_hash=digest(core_receipt),
                       language=language, narrative_hash=digest(narratives[language]),
                       writing_slots_hash=digest(make_narrative_input(checked, language, english=english)),
                       source_english_hash=digest(english) if english is not None else None)
        specifications = {"html": (media_stages / ("report-" + language + ".html"),
                                    media_stages / ("report-" + language + ".payload.html"), MAX_STAGE_FILE)}
        intent_path = stages / ("edition-" + language + "-intent.json")
        safe_path(intent_path)
        if intent_path.exists():
            intent = load_json(intent_path, 128 * 1024)
            saved_binding = intent.get("binding") if isinstance(intent, dict) else None
            saved_version = saved_binding.get("render_version") if isinstance(saved_binding, dict) else None
            if saved_version in _supported_render_versions(core):
                # Only select the declared historical layout; _stage_saved
                # still verifies every other binding and the original bytes.
                binding["render_version"] = saved_version
        saved = _stage_saved(stages, "edition-" + language, binding, specifications)
        if saved is not None:
            return saved
        raw = render_html(issue, checked, language=language, root=destination).encode("utf-8")
        return _stage_commit(stages, "edition-" + language, binding, specifications, {"html": raw})


def write_issue_package(issue, dataset, destination, language="en"):
    """destination is an output/project root; media never goes outside media/.

    Facts and one shared workbook live at data/outputs/<issue>/revision-<n>/;
    HTML lives at media/live/reports/<issue>/revision-<n>/. First two language
    editions share identical facts/workbook. A changed dataset needs a new
    revision. Returned manifest is draft-only, not upload/publication approval.
    """
    validate_dataset(issue, dataset, root=destination if dataset.get("schema_version") == 3 else None)
    if language not in LANGUAGES or language not in dataset["narratives"]:
        raise SafetyError("LIVE_LANGUAGE_MISSING")
    root = safe_path(destination)
    issue_id, revision = identifier(issue["issue_id"]), identifier("revision-" + str(issue["revision"]))
    data_dir = safe_path(runtime_path(root, "data", "outputs"), issue_id, revision)
    media_dir = safe_path(runtime_path(root, "media", "live", "reports"), issue_id, revision)
    facts_path, workbook_path = data_dir / "final_report_data.json", data_dir / "shortlist.xlsx"
    html_path, manifest_path = media_dir / ("report-" + language + ".html"), data_dir / ("manifest-" + language + ".json")
    for path in (facts_path, workbook_path, html_path, manifest_path):
        safe_path(path)
    facts_hash = digest(dataset)
    if facts_path.exists() and digest(load_json(facts_path, 4 * 1024 * 1024)) != facts_hash:
        raise SafetyError("LIVE_IMMUTABLE_FACTS_CHANGED")
    # A second language cannot bless a changed workbook or facts file from the
    # first language. Verify the already accepted shared bytes before reuse.
    for existing_language in sorted(LANGUAGES):
        prior_path = data_dir / ("manifest-" + existing_language + ".json")
        if prior_path.exists():
            prior = load_json(prior_path)
            if "price_display_version" in prior and (type(prior["price_display_version"]) is not str
                                                      or prior["price_display_version"] != PRICE_DISPLAY_VERSION):
                raise SafetyError("LIVE_PRICE_DISPLAY_VERSION_INVALID")
            if not manifest_path.exists() and "price_display_version" not in prior:
                # An old shared workbook cannot acquire a newly rendered
                # edition under a different display contract in the same revision.
                raise SafetyError("LIVE_RENDER_VERSION_CHANGED")
            old_html = media_dir / ("report-" + existing_language + ".html")
            try:
                valid = (prior["facts_hash"] == facts_hash and facts_path.exists() and workbook_path.exists() and old_html.exists()
                         and digest_bytes(facts_path.read_bytes()) == prior["artifacts"]["facts"]["sha256"]
                         and digest_bytes(workbook_path.read_bytes()) == prior["artifacts"]["workbook"]["sha256"]
                         and digest_bytes(old_html.read_bytes()) == prior["artifacts"]["html"]["sha256"])
            except KeyError:
                valid = False
            if not valid:
                raise SafetyError("LIVE_PACKAGE_INTEGRITY_CHANGED")
    if manifest_path.exists():
        manifest = load_json(manifest_path)
        if (manifest.get("facts_hash") != facts_hash or not html_path.exists()
                or digest_bytes(html_path.read_bytes()) != manifest["artifacts"]["html"]["sha256"]
                or not workbook_path.exists() or digest_bytes(workbook_path.read_bytes()) != manifest["artifacts"]["workbook"]["sha256"]):
            raise SafetyError("LIVE_PACKAGE_INTEGRITY_CHANGED")
        if "internal_stage" in manifest:
            validate_stage_provenance(dataset, language, manifest["internal_stage"])
        return manifest
    stage_dir = safe_path(data_dir, "internal-stages")
    staged = dataset["schema_version"] == 3 or any((stage_dir / name).exists() for name in ("core-intent.json", "core-receipt.json", "workbook-intent.json"))
    stage_provenance = None
    if staged:
        choice = dataset["selected_language"]
        required = LANGUAGES if issue["projected_ordinal"] <= 2 else ({choice} if choice in LANGUAGES else set())
        if not required or not required.issubset(dataset["narratives"]):
            raise SafetyError("LIVE_FINAL_WRITING_NOT_COMPLETE")
        workbook_receipt = write_issue_workbook(issue, dataset, destination)
        edition_receipt = write_issue_edition(issue, dataset, destination, language=language)
        workbook = _stage_read(Path(workbook_receipt["artifacts"]["workbook"]["path"]))
        encoded_html = _stage_read(Path(edition_receipt["artifacts"]["html"]["path"]))
        stage_provenance = {"stage_version": INTERNAL_STAGE_VERSION,
                            "core_facts_hash": edition_receipt["core_facts_hash"],
                            "workbook_generation_version": workbook_receipt["generation_version"],
                            "edition_render_version": edition_receipt["render_version"],
                            "narrative_hash": edition_receipt["narrative_hash"],
                            "writing_slots_hash": edition_receipt["writing_slots_hash"],
                            "source_english_hash": edition_receipt["source_english_hash"],
                            "workbook_receipt_hash": digest(workbook_receipt),
                            "edition_receipt_hash": digest(edition_receipt)}
        # Promotion never overwrites an old or interrupted mismatching alias.
        for path, raw in ((workbook_path, workbook), (html_path, encoded_html)):
            if path.exists() and _stage_read(path) != raw:
                raise SafetyError("LIVE_PACKAGE_INTEGRITY_CHANGED")
    else:
        # Original, unstaged historical calls retain the whole-dataset and XLSX
        # format contract. New checked stages use their own pinned exact bytes.
        encoded_html = render_html(issue, dataset, language=language, root=destination).encode("utf-8")
        workbook = workbook_path.read_bytes() if workbook_path.exists() else _workbook_bytes(issue, dataset, root=destination)
    # A packet is accepted only when this final manifest exists. Interrupted
    # writes are not a complete package; no partially written file is published.
    if not facts_path.exists():
        atomic_json(facts_path, dataset)
    if not workbook_path.exists():
        atomic_bytes(workbook_path, workbook)
    if staged:
        _stage_exact(html_path, encoded_html)
    else:
        atomic_bytes(html_path, encoded_html)
    artifacts = {"facts": {"path": str(facts_path), "sha256": digest_bytes(facts_path.read_bytes())},
                 "workbook": {"path": str(workbook_path), "sha256": digest_bytes(workbook)},
                 "html": {"path": str(html_path), "sha256": digest_bytes(encoded_html)}}
    manifest = {"schema_version": 2 if dataset["schema_version"] == 3 else 1, "issue_id": issue_id, "revision": issue["revision"], "language": language,
                "facts_hash": facts_hash, "selection_hash": issue["selection_hash"], "evidence_type": dataset["evidence_type"],
                "concept_ids": [product["concept_id"] for product in dataset["products"]], "counts": deepcopy(issue["counts"]),
                "projected_ordinal": issue["projected_ordinal"], "language_state": issue["language_state"],
                "price_display_version": PRICE_DISPLAY_VERSION,
                "approved": False, "release_ready": False, "artifacts": artifacts}
    if dataset.get("fit_policy_id") == FIT_POLICY_ID:
        manifest.update(fit_policy_id=FIT_POLICY_ID, criterion_ids=deepcopy(dataset["criterion_ids"]), kpi_total=10,
                        workbook_contract=workbook_contract(dataset))
    if dataset["schema_version"] == 3:
        manifest.update(recommendation_contract=RECOMMENDATION_CONTRACT,
                        asset_receipt_hash=digest(dataset["report_asset_receipt"]))
    if "opportunity_policy_id" in dataset:
        manifest["opportunity_policy_id"] = dataset["opportunity_policy_id"]
    if stage_provenance is not None:
        manifest["internal_stage"] = stage_provenance
    atomic_json(manifest_path, manifest)
    return manifest


def validate_stage_provenance(dataset, language, provenance):
    """Check additive final-draft stage bindings without requiring stage files.

    Final packets remain independently verifiable after their internal payload
    working files are archived. Their actual promoted bytes are checked by the
    ordinary final artifact hashes and delivery validators.
    """
    from .report_narration import make_narrative_input
    english = dataset["narratives"].get("en") if language == "hi-Latn" else None
    core = {key: value for key, value in dataset.items() if key not in ("narratives", "selected_language")}
    own = {language: deepcopy(dataset["narratives"][language])}
    if english is not None:
        own["en"] = deepcopy(english)
    checked = dict(core, narratives=own, selected_language=None)
    generation = provenance.get("workbook_generation_version") if isinstance(provenance, dict) else None
    current_generation = _generation_version(dataset)
    if generation not in (current_generation, LEGACY_WORKBOOK_GENERATION_VERSION):
        raise SafetyError("LIVE_STAGE_PROVENANCE_CHANGED")
    render_version = provenance.get("edition_render_version")
    if render_version not in _supported_render_versions(dataset):
        raise SafetyError("LIVE_STAGE_PROVENANCE_CHANGED")
    expected = {"stage_version": INTERNAL_STAGE_VERSION, "core_facts_hash": digest(core),
                "workbook_generation_version": generation,
                "edition_render_version": render_version,
                "narrative_hash": digest(dataset["narratives"][language]),
                "writing_slots_hash": digest(make_narrative_input(checked, language, english=english)),
                "source_english_hash": digest(english) if english is not None else None}
    if (not isinstance(provenance, dict) or set(provenance) != set(expected) | {"workbook_receipt_hash", "edition_receipt_hash"}
            or any(provenance.get(key) != value for key, value in expected.items())
            or any(not isinstance(provenance.get(key), str) or not re.fullmatch(r"[a-f0-9]{64}", provenance[key])
                   for key in ("workbook_receipt_hash", "edition_receipt_hash"))):
        raise SafetyError("LIVE_STAGE_PROVENANCE_CHANGED")


def digest_bytes(value):
    import hashlib
    return hashlib.sha256(value).hexdigest()
