"""Current-ten fit joins, using fictional source data and no live services."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))
from test_report_selection import fixture, NOW
from product_radar.fit_policy import ACTIVE_CRITERION_IDS, FIT_POLICY_ID, evaluate_fit
from product_radar.service import product_fit, policy_for_fit, validate_issue
from product_radar.report_selection import build_report_pool, selection_scope
from product_radar.issue_allocation import allocate_issues
from product_radar.live_outputs import build_issue_dataset, validate_dataset, render_html, _workbook_bytes, write_issue_package, digest_bytes
from product_radar.delivery_package import _workbook, write_delivery_manifest, validate_delivery_manifest
from product_radar.report_narration import generate_report_narratives, make_narrative_input
from product_radar.research_memory import ResearchMemory, key
from product_radar.storage import Journal
from product_radar.supplier_effort import record_supplier_effort
from product_radar.models import SafetyError, digest
from product_radar.research_flow import assessment_item, current_assessment_fit


def current_fixture(count=10, matches=6):
    memory = fixture(count)
    for record in memory.state["assessments"].values():
        rows = [row for row in record["result"]["criteria"] if row["id"] in ACTIVE_CRITERION_IDS]
        for n, row in enumerate(rows):
            row["state"] = "Meets" if n < matches else "Unclear"
        record.update(fit_policy_id=FIT_POLICY_ID, contract_version="research-e-fit10-v1")
        record["result"]["criteria"] = rows
    return memory


def selection(memory):
    pool = build_report_pool(memory, NOW, fit_policy_id=FIT_POLICY_ID)
    issue = allocate_issues(pool["eligible"], ["2026-10-07"], [], [], policy=policy_for_fit(FIT_POLICY_ID))["issues"][0]
    facts = [pool["facts_by_concept"][row["concept_id"]] for row in issue["products"]]
    return pool, issue, build_issue_dataset(issue, facts, evidence_type="synthetic")


class WritingFixture:
    identity = "synthetic-fit-writer"
    def __init__(self):
        self.calls = 0
    def __call__(self, role, payload, folder):
        self.calls += 1
        assert role == "G"
        result = []
        for item in payload["items"]:
            row = deepcopy(item["output_template"])
            if item["kind"] == "product":
                row["reason"] = "{{title}} is a fictional example." if item["language"] == "en" else "{{title}} sirf ek test example hai."
                row["checklist_phrases"] = [{"id": criterion["id"], "text": "Fixture judgment" if item["language"] == "en" else "Test ki rai"} for criterion in item["criteria"]]
            result.append(row)
        return {"results": result}, {"input_tokens": 0, "output_tokens": 0}


class CurrentFitPipelineTests(unittest.TestCase):
    def test_market_only_updates_reuse_exact_practical_fit_without_model_work(self):
        from product_radar.programme import _fit
        initial = current_fixture(1)
        initial.state["assessments"]["c000"]["contract_version"] = "research-e-fit10-v2"
        item = assessment_item(initial.state, "c000")
        original_assessment = deepcopy(initial.state["assessments"]["c000"])
        # A new dated observation of the same listing changes market facts,
        # not the original descriptive checklist reference.
        newer = deepcopy(initial.state["observations"]["meta000"])
        newer["facts"].update(active=False, started_at="2026-10-06T08:00:00+00:00")
        newer["observed_at"] = NOW
        newer["acquisition_id"] = "new-synthetic-market-capture"
        initial.state["observations"]["new-market-meta"] = newer
        # Newly matched retailer context has no identity/discovery association.
        retailer = deepcopy(initial.state["observations"]["amazon000"])
        retailer.update(source="flipkart", semantic_id="new-listing-context", observed_at=NOW)
        retailer["facts"].update(list_type="product_search", position=None,
                                 price={"displayed": 249, "currency": "INR"})
        initial.state["observations"]["new-retailer-context"] = retailer
        initial.state["market_reviews"]["new-market-review"] = {
            "concept_id": "c000", "review_version": "marketplace-relevance-v1",
            "description_hash": digest(initial.state["concepts"]["c000"]["description"]), "reviewed_at": NOW,
            "input_hash": digest(retailer), "relevance": [{"id": "new-retailer-context", "relationship": "relevant_family",
                                                          "reason": "Synthetic matching listing context."}]}
        self.assertEqual(assessment_item(initial.state, "c000"), item)
        self.assertIsNotNone(current_assessment_fit(initial.state, "c000", expected_policy_id=FIT_POLICY_ID))
        class NoModel:
            identity = "never-called-fit-reuse-fixture"
            def __call__(self, *_args):
                raise AssertionError("A market-only update must not repeat the full checklist.")
        with tempfile.TemporaryDirectory() as directory:
            memory = ResearchMemory(Path(directory) / "research", Journal(Path(directory) / "finance"), evidence_type="synthetic")
            for table, rows in initial.state.items():
                for record_id, row in rows.items():
                    memory.put(table, record_id, row)
            result = _fit(memory, {"concept_ids": ["c000"], "deferrals": {}},
                          {"at": NOW, "runner": NoModel(), "fit_limit": 5, "retry_seconds": 300})
            self.assertEqual(result["current_accepted_reused"], 1)
            self.assertEqual(result["prepared"], 0)
            self.assertEqual(result["outcomes"], [])
            self.assertEqual(memory.snapshot()["assessments"]["c000"], original_assessment)
            self.assertEqual(memory.snapshot()["jobs"], {})

    def test_changed_mechanism_or_new_associated_descriptive_evidence_needs_real_fit_review(self):
        for mutation in ("mechanism", "associated_description", "original_text"):
            memory = current_fixture(1)
            if mutation == "mechanism":
                memory.state["concepts"]["c000"]["description"]["mechanism"] = "Powered motor, not manual"
            elif mutation == "associated_description":
                row = deepcopy(memory.state["observations"]["meta000"])
                row.update(semantic_id="changed-mechanism", semantic={"title": "Different powered purchase unit"})
                memory.state["observations"]["new-descriptive-evidence"] = row
                description = {"id": "changed-mechanism", "kind": "product",
                    "products": [deepcopy(memory.state["concepts"]["c000"]["description"])],
                    "reason": "Synthetic newly associated descriptive evidence."}
                memory.state["descriptions"]["changed-mechanism"] = {"result": description, "job_id": "fixture-description"}
                child = key("child", ["changed-mechanism", 0, digest(description)])
                memory.state["associations"][child] = {
                    "concept_id": "c000", "semantic_id": "changed-mechanism", "status": "resolved"}
            else:
                memory.state["observations"]["meta000"]["semantic"]["title"] = "Different powered product description"
            with self.subTest(mutation=mutation):
                self.assertIsNone(current_assessment_fit(memory.state, "c000", expected_policy_id=FIT_POLICY_ID))
                pool = build_report_pool(memory, NOW, fit_policy_id=FIT_POLICY_ID)
                self.assertIn("checklist_review_required", pool["decisions"][0]["codes"])

    def test_current_assessment_requires_the_matching_scoring_contract(self):
        for contract in (None, "research-abe-v1"):
            memory = current_fixture(1)
            if contract is None:
                memory.state["assessments"]["c000"].pop("contract_version")
            else:
                memory.state["assessments"]["c000"]["contract_version"] = contract
            pool = build_report_pool(memory, NOW, fit_policy_id=FIT_POLICY_ID)
            self.assertEqual(pool["eligible"], [])
            self.assertIn("checklist_review_required", pool["decisions"][0]["codes"])
        self.assertEqual(len(build_report_pool(current_fixture(1), NOW, fit_policy_id=FIT_POLICY_ID)["eligible"]), 1)

    def test_exact_original_ids_five_unknowns_and_current_dataset(self):
        pool, issue, dataset = selection(current_fixture(matches=5))
        self.assertEqual(len(pool["eligible"]), 10)
        self.assertTrue(issue["ready_for_review"])
        self.assertEqual(dataset["schema_version"], 2)
        for product in dataset["products"]:
            self.assertEqual(product["criterion_ids"], list(ACTIVE_CRITERION_IDS))
            self.assertEqual(product["fit_tier"], "decent")
            self.assertEqual(product_fit(product)["unknown"], 5)
        self.assertEqual(validate_issue(issue["products"])["products"], 10)

    def test_unknown_below_threshold_is_pending_not_false_rejection(self):
        memory = current_fixture(1, matches=4)
        pool = build_report_pool(memory, NOW, fit_policy_id=FIT_POLICY_ID)
        self.assertEqual(pool["decisions"][0]["decision"], "deferred")
        self.assertIn("fit_information_pending", pool["decisions"][0]["codes"])
        self.assertEqual(len(pool["needed_supplier_work"]), 0)

    def test_supported_negative_below_threshold_really_excludes(self):
        memory = current_fixture(1, matches=4)
        for row in memory.state["assessments"]["c000"]["result"]["criteria"][4:]:
            row["state"] = "Does not meet"
        pool = build_report_pool(memory, NOW, fit_policy_id=FIT_POLICY_ID)
        self.assertEqual(pool["decisions"][0]["decision"], "excluded")

    def test_old_policy_requires_projection_without_changing_source_facts(self):
        memory = fixture(1)
        before = deepcopy(memory.state["observations"])
        pool = build_report_pool(memory, NOW, fit_policy_id=FIT_POLICY_ID)
        self.assertEqual(pool["eligible"], [])
        self.assertIn("checklist_review_required", pool["decisions"][0]["codes"])
        self.assertEqual(memory.state["observations"], before)

    def test_good_tier_priority_survives_category_and_input_order(self):
        memory = current_fixture(12, matches=5)
        record = memory.state["assessments"]["c011"]
        record["result"]["criteria"][5]["state"] = "Meets"
        pool, issue, dataset = selection(memory)
        self.assertEqual(issue["products"][0]["concept_id"], "c011")
        self.assertEqual(dataset["products"][0]["fit_tier"], "good")

    def test_tampered_policy_id_count_tier_and_criterion_state_are_rejected(self):
        _, issue, dataset = selection(current_fixture())
        for field, value in (("kpi_total", 13), ("criterion_ids", list(range(1, 11))), ("fit_tier", "decent"), ("kpi_matches", 10)):
            bad = deepcopy(dataset); bad["products"][0][field] = value
            with self.subTest(field=field), self.assertRaises(SafetyError):
                validate_dataset(issue, bad)
        bad = deepcopy(dataset); bad["products"][0]["criteria"][0]["state"] = "Unclear"
        with self.assertRaises(SafetyError):
            validate_dataset(issue, bad)

    def test_failed_real_supplier_effort_allows_product_but_prelaunch_does_not(self):
        initial = current_fixture(10)
        initial.state["supplier_assessments"].clear()
        with tempfile.TemporaryDirectory() as directory:
            memory = ResearchMemory(Path(directory) / "research", Journal(Path(directory) / "finance"), evidence_type="synthetic")
            for table, records in initial.state.items():
                for record_id, record in records.items():
                    memory.put(table, record_id, record)
            for concept_id, concept in initial.state["concepts"].items():
                receipt = {"schema_version": 2, "id": "attempt-one", "requested_url": "https://supplier.example.org/item",
                    "final_url": None, "checked_at": NOW, "http_status": None, "execution_stage": "attempt_started",
                    "attempt_started_at": NOW, "execution_ref": "public-fetch-" + "a" * 32, "reason": "PUBLIC_NETWORK_ERROR"}
                record_supplier_effort(memory, concept_id=concept_id, description=concept["description"],
                    pass_id="pass-" + concept_id, receipts=[receipt], receipt_reference="synthetic-owned-receipt", outcome="failed")
            pool, issue, dataset = selection(memory)
            self.assertEqual(len(pool["eligible"]), 10)
            product = dataset["products"][0]
            self.assertEqual(product["supplier_search_status"], "blocked")
            self.assertTrue(product["supplier_effort_satisfied"])
            payload = make_narrative_input(dataset, "en")
            self.assertEqual(payload["items"][1]["required_limitation"], "Supplier search did not produce any usable results.")
            self.assertTrue(issue["products"][0]["source_revisions"]["source_checks"])

    def test_both_languages_all_ten_phrases_workbook_and_no_call_replay(self):
        _, issue, dataset = selection(current_fixture())
        with tempfile.TemporaryDirectory() as directory:
            memory = ResearchMemory(Path(directory) / "research", Journal(Path(directory) / "finance"), evidence_type="synthetic")
            runner = WritingFixture()
            result = generate_report_narratives(memory, issue, dataset, runner)
            self.assertEqual(result["failures"], {})
            self.assertEqual(result["complete_languages"], ["en", "hi-Latn"])
            calls = runner.calls
            replay = generate_report_narratives(memory, issue, dataset, runner)
            self.assertEqual(runner.calls, calls)
            self.assertEqual(result["dataset"], replay["dataset"])
            for language in ("en", "hi-Latn"):
                html = render_html(issue, result["dataset"], language=language)
                self.assertIn("6/10", html)
                self.assertNotIn("6/13", html)
                self.assertEqual(len(result["dataset"]["narratives"][language]["products"][0]["checklist_phrases"]), 10)
            workbook = _workbook_bytes(issue, result["dataset"])
            _workbook(workbook, result["dataset"])
            from openpyxl import load_workbook
            book = load_workbook(BytesIO(workbook), read_only=True)
            self.assertEqual(book.sheetnames, ["Shortlist", "Evidence", "Supplier leads", "Checklist"])
            self.assertEqual(book["Checklist"].max_row, 106)
            book.close()

    def test_current_and_legacy_fit_do_not_mix_in_one_issue(self):
        pool, issue, _ = selection(current_fixture())
        mixed = deepcopy(issue["products"])
        legacy = fixture(1)
        mixed[0] = build_report_pool(legacy, NOW)["eligible"][0]
        with self.assertRaisesRegex(SafetyError, "ISSUE_FIT_POLICY_MISMATCH"):
            validate_issue(mixed)
        result = allocate_issues(mixed, ["2026-10-07"], [], [])
        self.assertEqual(len(result["issues"][0]["products"]), 9)
        self.assertFalse(result["issues"][0]["ready_for_review"])

    def test_malformed_candidate_does_not_stop_ten_healthy_products(self):
        pool, _, _ = selection(current_fixture())
        for bad in ({"concept_id": "not an id"}, {"concept_id": "invalid-date", "usable_until": []}, None):
            result = allocate_issues([*pool["eligible"], bad], ["2026-10-07"], [], [], policy=policy_for_fit(FIT_POLICY_ID))
            self.assertTrue(result["issues"][0]["ready_for_review"])
            self.assertEqual(len(result["issues"][0]["products"]), 10)
            self.assertTrue(any(row.get("diagnostic_code") for row in result["decisions"]))

    def test_reservation_rejects_changed_header_before_saving(self):
        from product_radar.issue_allocation import reserve_issue, allocation_snapshot_hash
        _, issue, _ = selection(current_fixture())
        with tempfile.TemporaryDirectory() as directory:
            memory = ResearchMemory(Path(directory) / "research", Journal(Path(directory) / "finance"), evidence_type="synthetic")
            for change in ({"fit_policy_id": "fit-13-v1"}, {"criterion_ids": [99]}, {"kpi_total": True}):
                bad = dict(deepcopy(issue), **change)
                with self.assertRaisesRegex(SafetyError, "ISSUE_FIT_POLICY_MISMATCH"):
                    reserve_issue(memory, bad, expected_input_hash=issue["input_hash"], expected_snapshot_hash=allocation_snapshot_hash(memory.snapshot()))
            self.assertEqual(memory.snapshot()["issue_reservations"], {})

    def test_complete_current_bilingual_packet_has_bound_checklist_workbook(self):
        import fitz
        _, issue, dataset = selection(current_fixture())
        with tempfile.TemporaryDirectory() as directory:
            memory = ResearchMemory(Path(directory) / "research", Journal(Path(directory) / "finance"), evidence_type="synthetic")
            dataset = generate_report_narratives(memory, issue, dataset, WritingFixture())["dataset"]
            drafts, receipts = {}, {}
            for language in ("en", "hi-Latn"):
                draft = write_issue_package(issue, dataset, directory, language=language)
                path = Path(draft["artifacts"]["html"]["path"]).with_suffix(".pdf")
                with fitz.open() as document:
                    page = document.new_page(width=595, height=842)
                    page.insert_text((30, 40), "Synthetic fit packet, not a client report.")
                    raw = document.tobytes()
                from product_radar.storage import atomic_bytes
                atomic_bytes(path, raw)
                drafts[language] = draft
                receipts[language] = {"path": str(path), "html_sha256": draft["artifacts"]["html"]["sha256"],
                    "pdf_sha256": digest_bytes(raw), "pages": 1, "portrait": True, "bytes": len(raw)}
            manifest = write_delivery_manifest(issue, drafts, receipts, root=directory)
            self.assertEqual(manifest["fit_policy_id"], FIT_POLICY_ID)
            self.assertEqual(manifest["workbook_contract"], "fit-checklist-v1")
            validate_delivery_manifest(manifest, issue, root=directory)
            bad = deepcopy(dataset)
            bad["narratives"]["en"]["products"][0]["checklist_phrases"][0]["text"] = "Saves 99 percent"
            with self.assertRaisesRegex(SafetyError, "LIVE_CHECKLIST_PHRASES_INVALID"):
                validate_dataset(issue, bad)


if __name__ == "__main__":
    unittest.main()
