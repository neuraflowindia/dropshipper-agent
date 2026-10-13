"""Real file parsing with synthetic facts; no uploads or client publication."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import fitz
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from product_radar.delivery_package import build_delivery_manifest, validate_delivery_manifest, write_delivery_manifest, _workbook
from product_radar.live_outputs import (LEGACY_PRICE_DISPLAY_VERSION, PRICE_DISPLAY_VERSION, digest_bytes,
                                       price_display, write_issue_package, _workbook_bytes)
from product_radar.issue_allocation import _fingerprint
from product_radar.models import SafetyError, digest
from product_radar.storage import atomic_bytes, atomic_json
from test_live_outputs import fixture


def packet(root, *, ordinal=1, languages=("en", "hi-Latn"), count=10, updates=0, raw_quote=False):
    issue, dataset = fixture(count)
    issue["projected_ordinal"] = ordinal
    issue["language_state"] = "bilingual_trial" if ordinal <= 2 else "language_pending"
    if raw_quote:
        dataset["products"][0]["prices"][0].update(amount=None, raw_displayed="₹149–₹189", quantity=None,
                                                   unit="not_stated", conditions="Displayed quote basis not verified.")
    for product in issue["products"][:updates]:
        product.update(status="Updated", material_update_reviewed=True, update_reason="A supported new useful observation.")
    for product in dataset["products"][:updates]:
        product.update(status="Updated", material_update_reviewed=True, update_reason="A supported new useful observation.")
    if updates:
        issue["counts"] = {"products": count, "new": count - updates, "updates": updates}
        issue["evidence_fingerprint"] = _fingerprint(issue["products"])
        issue["selection_hash"] = digest({field: issue[field] for field in ("issue_id", "client_id", "revision", "date", "products", "evidence_fingerprint")})
        dataset["selection_hash"] = issue["selection_hash"]
    drafts, receipts = {}, {}
    for language in languages:
        draft = write_issue_package(issue, dataset, root, language=language)
        path = Path(draft["artifacts"]["html"]["path"]).with_suffix(".pdf")
        with fitz.open() as doc:
            page = doc.new_page(width=595, height=842)
            page.insert_text((30, 40), "Synthetic PDF for delivery-binding tests only.")
            raw = doc.tobytes()
        path.write_bytes(raw)
        drafts[language] = draft
        receipts[language] = {"path": str(path), "html_sha256": draft["artifacts"]["html"]["sha256"],
                              "pdf_sha256": digest_bytes(raw), "pages": 1, "portrait": True, "bytes": len(raw)}
    return issue, dataset, drafts, receipts


class DeliveryPackageTests(unittest.TestCase):
    def build(self, root, issue, drafts, receipts, **kwargs):
        return build_delivery_manifest(issue, drafts, receipts, root=root, **kwargs)

    def legacy_packet(self, root):
        # Simulate the exact pre-marker deterministic formatter, not a range
        # of accepted labels. Only disposable synthetic fixture files change.
        def legacy(price, language="en"):
            return price_display(price, language, render_version=LEGACY_PRICE_DISPLAY_VERSION)
        with patch("product_radar.live_outputs.price_display", side_effect=legacy):
            issue, dataset, drafts, receipts = packet(root, raw_quote=True)
        for language, draft in drafts.items():
            del draft["price_display_version"]
            atomic_json(Path(root) / "data/outputs" / issue["issue_id"] / ("revision-" + str(issue["revision"])) /
                        ("manifest-" + language + ".json"), draft)
        return issue, dataset, drafts, receipts

    def test_unversioned_legacy_raw_quote_packet_validates_without_changing_any_bytes(self):
        with tempfile.TemporaryDirectory() as root:
            issue, dataset, drafts, receipts = self.legacy_packet(root)
            manifest = write_delivery_manifest(issue, drafts, receipts, root=root)
            before = {path: path.read_bytes() for path in Path(root).rglob("*") if path.is_file()}
            validate_delivery_manifest(manifest, issue, root=root)
            self.assertEqual(before, {path: path.read_bytes() for path in Path(root).rglob("*") if path.is_file()})
            self.assertTrue(all("price_display_version" not in draft for draft in manifest["language_manifests"].values()))
            self.assertEqual(self.build(root, issue, drafts, receipts), manifest)
            self.assertEqual(manifest["facts_hash"], digest(dataset))

    def test_new_raw_quote_packet_declares_current_display_contract(self):
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = packet(root, raw_quote=True)
            self.assertTrue(all(draft["price_display_version"] == PRICE_DISPLAY_VERSION for draft in drafts.values()))
            manifest = self.build(root, issue, drafts, receipts)
            validate_delivery_manifest(manifest, issue, root=root)

    def test_current_packet_cannot_accept_legacy_workbook_labels_even_with_updated_byte_hashes(self):
        with tempfile.TemporaryDirectory() as root:
            issue, dataset, drafts, receipts = packet(root, raw_quote=True)
            path = Path(drafts["en"]["artifacts"]["workbook"]["path"])
            book = load_workbook(path)
            book["Shortlist"].cell(7, 7).value = price_display(dataset["products"][0]["prices"][0],
                                                               render_version=LEGACY_PRICE_DISPLAY_VERSION)
            raw = BytesIO()
            book.save(raw)
            book.close()
            atomic_bytes(path, raw.getvalue())
            for language, draft in drafts.items():
                draft["artifacts"]["workbook"]["sha256"] = digest_bytes(raw.getvalue())
                atomic_json(path.parent / ("manifest-" + language + ".json"), draft)
            with self.assertRaisesRegex(SafetyError, "DELIVERY_WORKBOOK_MISMATCH"):
                self.build(root, issue, drafts, receipts)

    def test_legacy_packet_keeps_artifact_hash_protection(self):
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = self.legacy_packet(root)
            path = Path(drafts["en"]["artifacts"]["workbook"]["path"])
            atomic_bytes(path, path.read_bytes() + b"changed")
            with self.assertRaisesRegex(SafetyError, "DELIVERY_ARTIFACT_CHANGED"):
                self.build(root, issue, drafts, receipts)

    def test_unknown_or_mixed_language_price_display_versions_are_refused(self):
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = packet(root, raw_quote=True)
            path = Path(drafts["en"]["artifacts"]["workbook"]["path"])
            original = deepcopy(drafts)
            for value in (None, "unrecognised-version", LEGACY_PRICE_DISPLAY_VERSION, False, True, 1.0, 2.0, "2", 3):
                with self.subTest(value=value):
                    drafts = deepcopy(original)
                    drafts["en"]["price_display_version"] = value
                    atomic_json(path.parent / "manifest-en.json", drafts["en"])
                    with self.assertRaisesRegex(SafetyError, "PRICE_DISPLAY_VERSION_INVALID"):
                        self.build(root, issue, drafts, receipts)
            drafts = deepcopy(original)
            del drafts["hi-Latn"]["price_display_version"]
            for language, draft in drafts.items():
                atomic_json(path.parent / ("manifest-" + language + ".json"), draft)
            with self.assertRaisesRegex(SafetyError, "SHARED_RENDER_VERSION_MISMATCH"):
                self.build(root, issue, drafts, receipts)

    def test_complete_trial_binds_every_language_same_facts_and_portal_files_but_not_approval(self):
        with tempfile.TemporaryDirectory() as root:
            issue, dataset, drafts, receipts = packet(root)
            result = self.build(root, issue, drafts, receipts)
            self.assertTrue(result["complete"])
            self.assertFalse(result["approved"])
            self.assertFalse(result["release_ready"])
            self.assertEqual(result["facts_hash"], digest(dataset))
            self.assertEqual(result["languages"], ["en", "hi-Latn"])
            self.assertEqual(result["concept_ids"], [p["concept_id"] for p in issue["products"]])
            self.assertEqual({f["name"] for f in result["files"]}, {"shortlist.xlsx", "report-en.pdf", "report-hi-Latn.pdf"})
            for file in result["files"]:
                self.assertEqual(file["key"], f'clients/{issue["client_id"]}/issues/{issue["issue_id"]}/r1/{file["name"]}')
            self.assertNotIn("server_activated_at", result)
            self.assertNotIn("scheduled_at", result)
            validate_delivery_manifest(result, issue, root=root)

    def test_both_first_ordinals_require_language_pair(self):
        for ordinal in (1, 2):
            with tempfile.TemporaryDirectory() as root:
                issue, _, drafts, receipts = packet(root, ordinal=ordinal, languages=("en",))
                with self.assertRaisesRegex(SafetyError, "REQUIRED_LANGUAGE_MISSING"):
                    self.build(root, issue, drafts, receipts)

    def test_missing_modified_or_wrongly_bound_pdf_does_not_complete(self):
        for change in ("missing", "modified", "html", "pages", "portrait", "path", "size"):
            with self.subTest(change=change), tempfile.TemporaryDirectory() as root:
                issue, _, drafts, receipts = packet(root)
                path = Path(receipts["en"]["path"])
                if change == "missing":
                    path.unlink()
                elif change == "modified":
                    path.write_bytes(path.read_bytes() + b"modified")
                elif change == "html":
                    receipts["en"]["html_sha256"] = "a" * 64
                elif change == "pages":
                    receipts["en"]["pages"] = 2
                elif change == "portrait":
                    receipts["en"]["portrait"] = False
                elif change == "path":
                    receipts["en"]["path"] = str(path.parent / "wrong.pdf")
                elif change == "size":
                    receipts["en"]["bytes"] += 1
                with self.assertRaises(SafetyError):
                    self.build(root, issue, drafts, receipts)

    def test_later_unset_language_stays_pending_and_selected_language_is_single(self):
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = packet(root, ordinal=3)
            pending = self.build(root, issue, drafts, receipts)
            self.assertEqual(pending["state"], "language_pending")
            self.assertFalse(pending["complete"])
            self.assertEqual(pending["files"], [])
            ready = self.build(root, issue, drafts, receipts, selected_language="hi-Latn")
            self.assertTrue(ready["complete"])
            self.assertEqual(ready["languages"], ["hi-Latn"])
            self.assertEqual([f["name"] for f in ready["files"]], ["shortlist.xlsx", "report-hi-Latn.pdf"])
            self.assertFalse(ready["approved"])

    def test_ordinal_reprojection_requires_trial_pair_and_changes_final_hash(self):
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = packet(root, ordinal=3, languages=("en",))
            ready = self.build(root, issue, drafts, receipts, selected_language="en")
            with self.assertRaisesRegex(SafetyError, "REQUIRED_LANGUAGE_MISSING"):
                self.build(root, issue, drafts, receipts, publication_ordinal=2, selected_language="en")
            with self.assertRaises(SafetyError):
                validate_delivery_manifest(ready, issue, root=root, publication_ordinal=4)

    def test_existing_manifest_replays_and_never_overwrites_language_change(self):
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = packet(root, ordinal=3)
            first = write_delivery_manifest(issue, drafts, receipts, root=root)
            path = Path(root) / "data" / "outputs" / issue["issue_id"] / "revision-1" / "delivery-manifest.json"
            raw = path.read_bytes()
            self.assertEqual(first, write_delivery_manifest(issue, drafts, receipts, root=root))
            with self.assertRaisesRegex(SafetyError, "IMMUTABLE_MANIFEST_CHANGED"):
                write_delivery_manifest(issue, drafts, receipts, root=root, selected_language="en")
            self.assertEqual(path.read_bytes(), raw)

    def test_modified_facts_workbook_html_or_draft_manifest_are_rejected(self):
        for field in ("facts", "workbook", "html", "manifest"):
            with self.subTest(field=field), tempfile.TemporaryDirectory() as root:
                issue, _, drafts, receipts = packet(root)
                if field == "manifest":
                    drafts["en"]["approved"] = True
                else:
                    path = Path(drafts["en"]["artifacts"][field]["path"])
                    path.write_bytes(path.read_bytes() + b"changed")
                with self.assertRaises(SafetyError):
                    self.build(root, issue, drafts, receipts)

    def test_fake_approval_or_changed_package_hash_does_not_validate(self):
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = packet(root)
            result = self.build(root, issue, drafts, receipts)
            for changes in ({"approved": True}, {"release_ready": True}, {"complete": 1}, {"package_hash": "a" * 64}, {"revision": 2}):
                with self.assertRaises(SafetyError):
                    validate_delivery_manifest(dict(result, **changes), issue, root=root)

    def test_counts_selection_revision_and_novelty_cannot_change_after_freeze(self):
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = packet(root)
            for changes in ({"counts": {"products": 9, "new": 9, "updates": 0}}, {"selection_hash": "a" * 64}, {"revision": 2}):
                with self.assertRaises(SafetyError):
                    self.build(root, dict(issue, **changes), drafts, receipts)
            changed = deepcopy(issue)
            for product in changed["products"][:3]:
                product.update(status="Updated", material_update_reviewed=True, update_reason="New useful detail")
            changed["counts"] = {"products": 10, "new": 7, "updates": 3}
            with self.assertRaises(SafetyError):
                self.build(root, changed, drafts, receipts)

    def test_nonpdf_or_landscape_cannot_be_blessed_by_new_receipt(self):
        for landscape in (False, True):
            with tempfile.TemporaryDirectory() as root:
                issue, _, drafts, receipts = packet(root)
                if landscape:
                    with fitz.open() as doc:
                        page = doc.new_page(width=842, height=595)
                        page.insert_text((30, 40), "Landscape fixture")
                        raw = doc.tobytes()
                else:
                    raw = b"Not a PDF"
                Path(receipts["en"]["path"]).write_bytes(raw)
                receipts["en"].update(pdf_sha256=digest_bytes(raw), bytes=len(raw))
                with self.assertRaises(SafetyError):
                    self.build(root, issue, drafts, receipts)

    def test_symlink_escape_rejected(self):
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = packet(root)
            path = Path(receipts["en"]["path"])
            alternate = path.parent / "other.pdf"
            path.rename(alternate)
            path.symlink_to(alternate)
            with self.assertRaises(SafetyError):
                self.build(root, issue, drafts, receipts)

    def test_workbook_concept_change_rejected_even_if_draft_hash_updated(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = packet(root)
            path = Path(drafts["en"]["artifacts"]["workbook"]["path"])
            book = load_workbook(path)
            book["Shortlist"]["I7"] = "differentconcept"
            book.save(path)
            book.close()
            for language in drafts:
                drafts[language]["artifacts"]["workbook"]["sha256"] = digest_bytes(path.read_bytes())
                atomic_json(path.parent / ("manifest-" + language + ".json"), drafts[language])
            with self.assertRaisesRegex(SafetyError, "WORKBOOK_MISMATCH"):
                self.build(root, issue, drafts, receipts)

    def test_current_four_sheet_workbook_rejects_unsafe_actual_hyperlink_targets(self):
        from test_fit_pipeline import current_fixture, selection
        _, issue, dataset = selection(current_fixture())
        original = _workbook_bytes(issue, dataset)
        for target in ("javascript:alert(1)", "file://example.org/share", "https://example.org/?client_secret=fixture",
                       "https://example.org/?SESSION_TOKEN=fixture", "https://example.org/?api_key=",
                       "https://example.org/search%0Atext", "https://example.org/\\redirect"):
            book = load_workbook(BytesIO(original))
            book["Evidence"]["G7"].hyperlink = target
            changed = BytesIO(); book.save(changed); book.close()
            with self.subTest(target=target), self.assertRaisesRegex(SafetyError, "^DELIVERY_WORKBOOK_LINK_INVALID$"):
                _workbook(changed.getvalue(), dataset)

    def test_unknown_checklist_rows_without_references_roundtrip_as_blank_cells(self):
        from test_fit_pipeline import current_fixture, selection
        memory = current_fixture()
        for assessment in memory.state["assessments"].values():
            for criterion in assessment["result"]["criteria"][6:]:
                criterion.update(basis="not_stated", references=[])
        _, issue, dataset = selection(memory)
        raw = _workbook_bytes(issue, dataset)
        _workbook(raw, dataset)
        book = load_workbook(BytesIO(raw), read_only=True)
        try:
            self.assertIsNone(book["Checklist"]["H13"].value)
            self.assertEqual(book["Checklist"].max_row, 106)
        finally:
            book.close()
        self.assertEqual(dataset["products"][0]["criteria"][6]["references"], [])

    def test_blank_reference_normalisation_does_not_hide_missing_or_invented_refs(self):
        from test_fit_pipeline import current_fixture, selection
        memory = current_fixture()
        for assessment in memory.state["assessments"].values():
            assessment["result"]["criteria"][6].update(basis="not_stated", references=[])
        _, issue, dataset = selection(memory)
        original = _workbook_bytes(issue, dataset)
        for address, value in (("H7", None), ("H13", "invented-reference")):
            book = load_workbook(BytesIO(original))
            book["Checklist"][address] = value
            changed = BytesIO(); book.save(changed); book.close()
            with self.subTest(address=address), self.assertRaisesRegex(SafetyError, "^DELIVERY_WORKBOOK_MISMATCH$"):
                _workbook(changed.getvalue(), dataset)

    def test_extra_unsafe_workbook_link_cannot_hide_outside_source_columns(self):
        issue, dataset = fixture()
        book = load_workbook(BytesIO(_workbook_bytes(issue, dataset)))
        book["Shortlist"]["A7"].hyperlink = "file://example.org/share"
        changed = BytesIO(); book.save(changed); book.close()
        with self.assertRaisesRegex(SafetyError, "^DELIVERY_WORKBOOK_LINK_INVALID$"):
            _workbook(changed.getvalue(), dataset)

    def test_safe_actual_links_preserve_current_and_legacy_workbook_bytes(self):
        from test_fit_pipeline import current_fixture, selection
        _, current_issue, current_dataset = selection(current_fixture())
        legacy_issue, legacy_dataset = fixture()
        for issue, dataset, sheets in ((current_issue, current_dataset, 4), (legacy_issue, legacy_dataset, 3)):
            raw = _workbook_bytes(issue, dataset)
            original = bytes(raw)
            _workbook(raw, dataset)
            self.assertEqual(raw, original)
            book = load_workbook(BytesIO(raw))
            self.assertEqual(len(book.sheetnames), sheets)
            self.assertTrue(book["Evidence"]["G7"].hyperlink.target.startswith("https://"))
            book.close()

    def test_unsafe_workbook_link_is_rejected_even_after_draft_byte_hashes_are_updated(self):
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = packet(root)
            path = Path(drafts["en"]["artifacts"]["workbook"]["path"])
            book = load_workbook(path)
            book["Evidence"]["G7"].hyperlink = "javascript:alert(1)"
            changed = BytesIO(); book.save(changed); book.close()
            atomic_bytes(path, changed.getvalue())
            for language, draft in drafts.items():
                draft["artifacts"]["workbook"]["sha256"] = digest_bytes(changed.getvalue())
                atomic_json(path.parent / ("manifest-" + language + ".json"), draft)
            with self.assertRaisesRegex(SafetyError, "^DELIVERY_WORKBOOK_LINK_INVALID$"):
                self.build(root, issue, drafts, receipts)

    def test_10_12_15_keep_allowed_update_rounding_and_supported_facts(self):
        for count, updates in ((10, 2), (12, 2), (15, 3)):
            with tempfile.TemporaryDirectory() as root:
                issue, _, drafts, receipts = packet(root, ordinal=3, languages=("en",), count=count, updates=updates)
                result = self.build(root, issue, drafts, receipts, selected_language="en")
                self.assertEqual(result["new_count"], count - updates)
                self.assertEqual(result["product_count"], count)
                self.assertTrue(result["complete"])

    def test_price_label_change_rejected_even_after_byte_hash_replacement(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as root:
            issue, _, drafts, receipts = packet(root)
            path = Path(drafts["en"]["artifacts"]["workbook"]["path"])
            book = load_workbook(path)
            book["Shortlist"]["G7"] = "Retail listing: ₹0"
            book.save(path)
            book.close()
            for language in drafts:
                drafts[language]["artifacts"]["workbook"]["sha256"] = digest_bytes(path.read_bytes())
                atomic_json(path.parent / ("manifest-" + language + ".json"), drafts[language])
            with self.assertRaisesRegex(SafetyError, "WORKBOOK_MISMATCH"):
                self.build(root, issue, drafts, receipts)


if __name__ == "__main__":
    unittest.main()
