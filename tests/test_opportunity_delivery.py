"""The opportunity score travels through selection, files and delivery without AI choosing it."""

from html import escape
from copy import deepcopy
from io import BytesIO
import unittest

from openpyxl import load_workbook

from test_fit_pipeline import current_fixture, older_format_scores, selection
from product_radar.delivery_package import _workbook
from product_radar.discovery import MARKET_SITES
from product_radar.live_outputs import _workbook_bytes, render_html, validate_dataset
from product_radar.models import SafetyError, digest
from product_radar.opportunity_score import (OPPORTUNITY_FIELDS, POLICY_ID, crowding_text, evaluate_opportunity,
                                             opportunity_fields, score_text)
from product_radar.service import issue_fit_policy, product_fit


def narratives(dataset):
    return {language: {"title": "Fixture research", "products": [
        {"concept_id": product["concept_id"], "reason": "Fixture product research.", "points": [],
         "limitation": "Supplier terms need checking.", "update_sentence": "",
         "checklist_phrases": [{"id": row["id"], "text": "Fixture judgment" if language == "en" else "Test ki rai"}
                               for row in product["criteria"]]}
        for product in dataset["products"]]} for language in ("en", "hi-Latn")}


@older_format_scores
class OpportunityDeliveryTests(unittest.TestCase):
    def test_score_is_frozen_and_workbook_roundtrips_for_both_fit_levels(self):
        for matches, total in ((5, 42), (6, 45)):  # 7 want (2 Meta + 5 platform lists) + 25 room + product
            _, issue, dataset = selection(current_fixture(matches=matches))
            self.assertEqual(issue["opportunity_policy_id"], POLICY_ID)
            self.assertEqual(dataset["opportunity_policy_id"], POLICY_ID)
            for product, frozen in zip(dataset["products"], issue["products"]):
                self.assertTrue(all(product[field] == frozen[field] for field in OPPORTUNITY_FIELDS))
                self.assertEqual((product["opportunity_score"], product["opportunity_label"]), (total, "sweet_spot"))
            raw = _workbook_bytes(issue, dataset)
            _workbook(raw, dataset)
            sheet = load_workbook(BytesIO(raw), read_only=True)["Shortlist"]
            header = next(sheet.iter_rows(min_row=6, max_row=6, values_only=True))
            self.assertEqual(header[-7:], ("Score", "Want /20", "Early /30", "Room /25", "Product /25", "Label", "Crowding"))
            first = next(sheet.iter_rows(min_row=7, max_row=7, values_only=True))
            self.assertEqual(list(first[-7:]), [total, 7, 0, 25, total - 32, "Sweet spot", "Low"])

    def test_pdf_shows_the_code_owned_score_and_crowding_lines_in_both_languages(self):
        _, issue, dataset = selection(current_fixture())
        dataset["narratives"] = narratives(dataset)
        product = dataset["products"][0]
        parts = product["opportunity_basis"]["parts"]
        counts = parts["room"]["items"]["crowding"]["facts"]["counts"]
        for language in ("en", "hi-Latn"):
            html = render_html(issue, dataset, language=language)
            # The code-owned lines name the score ring and the crowding gauge; their
            # figures are the printed parts and crowding facts (never model prose).
            self.assertTrue(f'aria-label="{escape(score_text(product, language), quote=True)}"' in html)
            self.assertTrue(f'aria-label="{escape(crowding_text(product, language), quote=True)}"' in html)
            page = html.split(f'data-concept-id="{product["concept_id"]}"', 1)[1]
            for part in ("want", "early", "room", "product"):
                self.assertTrue(f'<span class="part-val"><b>{parts[part]["points"]}</b>/{parts[part]["max"]}</span>' in page, part)
            self.assertTrue(f'Amazon {counts["relevant_listings"]["amazon"]} · Flipkart {counts["relevant_listings"]["flipkart"]}' in page)
        self.assertIn("Amazon/Flipkart crowding: Low (1 relevant listings", crowding_text(product))
        self.assertIn("Amazon/Flipkart par bheed: Kam", crowding_text(product, "hi-Latn"))

    def test_modified_score_label_or_basis_is_not_a_presentation_change(self):
        _, issue, original = selection(current_fixture())
        for field, value in (("opportunity_score", 99), ("opportunity_label", "early_bet"),
                             ("opportunity_basis", {"settings_version": "x", "parts": {}, "eligible": True,
                                                    "reason": "opportunity_scored"})):
            changed = deepcopy(original); changed["products"][0][field] = value
            with self.subTest(field=field), self.assertRaises(SafetyError):
                validate_dataset(issue, changed)

    def test_header_cannot_hide_a_mixed_policy_or_unmarked_product(self):
        _, issue, _ = selection(current_fixture())
        for field in OPPORTUNITY_FIELDS:
            changed = deepcopy(issue); changed["products"][0].pop(field)
            with self.subTest(field=field), self.assertRaises(SafetyError):
                issue_fit_policy(changed)
        changed = deepcopy(issue); changed.pop("opportunity_policy_id")
        with self.assertRaisesRegex(SafetyError, "ISSUE_OPPORTUNITY_POLICY_MISMATCH"):
            issue_fit_policy(changed)

    def test_boolean_or_float_scores_are_not_integers(self):
        _, _, dataset = selection(current_fixture())
        for value in (True, 45.0):
            changed = deepcopy(dataset["products"][0]); changed["opportunity_score"] = value
            with self.subTest(value=value), self.assertRaisesRegex(SafetyError, "OPPORTUNITY_SCORE_INVALID"):
                product_fit(changed)

    def test_high_crowding_cannot_join_an_issue_that_had_ten_low_or_medium_types(self):
        _, issue, _ = selection(current_fixture(11))
        low = list(issue["products"])
        self.assertGreaterEqual(len(low), 10)
        crowded = deepcopy(low[0]); crowded["concept_id"] = "crowded-type"
        basis = crowded["opportunity_basis"]["parts"]
        market = {"sources": {site: {"available": False, "positive": None, "used_ids": []} for site in MARKET_SITES}}
        market["sources"]["meta"] = {"available": True, "positive": False,
                                     "used_ids": basis["want"]["items"]["facebook_ads"]["evidence_ids"],
                                     "depth": {"advertisers": 1, "recent_start_advertisers": None}}
        lists = basis["want"]["items"]["platform_lists"]["facts"]
        crowded.update(opportunity_fields(evaluate_opportunity(
            {"fit_policy_id": "fit-10-v1", "kpi_matches": crowded["kpi_matches"], "kpi_total": 10}, market,
            presence={"at": lists["presence_at"], "listed": lists["listed"], "on_list": lists["on_list"], "new_arrival": {}},
            crowding={"level": "high", "rule": "bestseller_list", "thresholds_version": "crowding-thresholds-v1",
                      "counts": None, "evidence_ids": []})))
        self.assertEqual(crowded["opportunity_label"], "proven_but_crowded")
        with self.assertRaisesRegex(SafetyError, "ISSUE_HIGH_CROWDING_RULE"):
            issue_fit_policy(dict(issue, products=low[:10] + [crowded]))
        self.assertEqual(issue_fit_policy(dict(issue, products=low[:9] + [crowded])), "fit-10-v1")

    def retired_tier_packet(self, policy_id, basis):
        """Facts frozen under the retired priority tiers (v1 three groups, v2 seven sites)."""
        _, issue, dataset = selection(current_fixture())
        issue["opportunity_policy_id"] = dataset["opportunity_policy_id"] = policy_id
        for records in (issue["products"], dataset["products"]):
            for product in records:
                for field in OPPORTUNITY_FIELDS:
                    product.pop(field, None)
                product.update(opportunity_policy_id=policy_id, opportunity_priority=1, opportunity_basis=deepcopy(basis))
        issue["selection_hash"] = digest({field: issue[field] for field in (
            "issue_id", "client_id", "revision", "date", "products", "evidence_fingerprint")})
        dataset["selection_hash"] = issue["selection_hash"]
        return issue, dataset

    def test_packets_ranked_by_the_retired_priority_tiers_are_refused(self):
        for policy_id, basis in (("opportunity-priority-v1", {"available_families": 2, "positive_families": 1}),
                                 ("opportunity-priority-v2", {"available_sources": 2, "positive_sources": 0,
                                                              "source_ids": ["amazon", "meta"], "positive_source_ids": [],
                                                              "evidence_by_source": {"amazon": ["amazon000"], "meta": ["meta000"]}})):
            issue, dataset = self.retired_tier_packet(policy_id, basis)
            with self.subTest(policy_id=policy_id), self.assertRaises(SafetyError):
                validate_dataset(issue, dataset)
            with self.subTest(policy_id=policy_id), self.assertRaises(SafetyError):
                issue_fit_policy(issue)


class OlderContractTypeTests(unittest.TestCase):
    def test_older_contract_types_wait_for_the_crowding_check(self):
        """Without the test double: an older-format type has no accepted meaning,
        so no crowding answer, and is never report-eligible under the score."""
        from product_radar.fit_policy import FIT_POLICY_ID
        from product_radar.report_selection import build_report_pool
        from test_report_selection import NOW
        pool = build_report_pool(current_fixture(1), NOW, fit_policy_id=FIT_POLICY_ID)
        self.assertEqual(pool["eligible"], [])
        decision = pool["decisions"][0]
        self.assertIn("crowding_pending", decision["codes"])
        self.assertEqual((decision["opportunity_policy_id"], decision["opportunity_label"]), (POLICY_ID, "too_quiet"))
        self.assertEqual(decision["opportunity_basis"]["reason"], "crowding_pending")
        work = pool["needed_market_work"][0]
        self.assertEqual((work["code"], work["blocking"]), ("crowding_pending", True))


if __name__ == "__main__":
    unittest.main()
