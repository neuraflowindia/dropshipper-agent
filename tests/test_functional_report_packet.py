"""Real synthetic producers through new report files, never publication."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

import fitz

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from openpyxl import load_workbook
from product_radar.delivery_package import _workbook
from product_radar.issue_allocation import allocate_issues, selection_fingerprint, reserve_issue, allocation_snapshot_hash
from product_radar.live_outputs import (build_issue_dataset, validate_dataset, render_html,
    _workbook_bytes, write_issue_workbook, write_issue_package)
from product_radar.models import SafetyError, digest
from product_radar.report_assets import prepare_report_asset_plan, freeze_report_assets, attach_report_assets
from product_radar.report_examples import RECOMMENDATION_CONTRACT, WORKBOOK_HEADERS
from product_radar.report_narration import generate_report_narratives, make_narrative_input
from product_radar.research_ai import Jobs
from product_radar.report_selection import build_report_pool
from product_radar.report_workflow import prepare_review_package, freeze_review_input
from product_radar.product_images import capture_product_image, inspect_image
from product_radar.programme import _functional_report_revision, _reports, prepare_programme
from product_radar.service import policy_for_fit
from product_radar.storage import atomic_bytes, load_json
from test_fit_pipeline import WritingFixture
from test_functional_source_market import functional_pool_fixture, NOW
from test_report_workflow import inspect


class RecordingWriter(WritingFixture):
    identity = "functional-packet-recording-writer"

    def __init__(self, bad_id=None):
        super().__init__()
        self.payloads, self.bad_id = [], bad_id

    def __call__(self, role, payload, folder):
        self.payloads.append(deepcopy(payload))
        response, usage = super().__call__(role, payload, folder)
        if payload["language"] == "en":
            for row in response["results"]:
                if row["id"] == self.bad_id:
                    row["references"] = ["invented-slot"]
        return response, usage


def functional_packet_fixture(root, *, count=10, revision=1, memory=None):
    root = Path(root)
    if memory is None:
        memory, _ = functional_pool_fixture(root / "inputs", count=count)
    pool = build_report_pool(memory, as_of=NOW, recommendation_contract=RECOMMENDATION_CONTRACT)
    allocation = allocate_issues(pool["eligible"], [{"issue_id": "functional-test-issue", "client_id": "default",
        "date": "2026-10-07", "revision": revision}], [], [], policy=policy_for_fit("fit-10-v1"))
    issue = allocation["issues"][0]
    if not issue["ready_for_review"]:
        raise AssertionError((issue["counts"], allocation["decisions"][:1]))
    facts = [pool["facts_by_concept"][row["concept_id"]] for row in issue["products"]]
    plan = prepare_report_asset_plan(facts, staging_root=root / "display-stage")
    receipt = freeze_report_assets(plan, root=root, issue_id=issue["issue_id"], revision=revision)
    issue, facts = attach_report_assets(issue, facts, receipt, root=root)
    return memory, pool, issue, build_issue_dataset(issue, facts, evidence_type="synthetic")


class FunctionalPacketTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.shared = tempfile.TemporaryDirectory()
        cls.addClassCleanup(cls.shared.cleanup)
        cls.memory, _ = functional_pool_fixture(Path(cls.shared.name), count=10)

    def setUp(self):
        temporary = tempfile.TemporaryDirectory(); self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.memory, self.pool, self.issue, self.dataset = functional_packet_fixture(self.root, memory=self.memory)

    def test_actual_new_pool_freezes_exact_complete_tuple_and_optional_photo_omissions(self):
        self.assertEqual(self.dataset["schema_version"], 3)
        self.assertEqual(self.dataset["recommendation_contract"], RECOMMENDATION_CONTRACT)
        self.assertEqual(self.issue["counts"], {"products": 10, "new": 10, "updates": 0})
        self.assertEqual(self.issue["selection_hash"], selection_fingerprint(self.issue))
        validate_dataset(self.issue, self.dataset, root=self.root)
        self.assertTrue(all(row["display_asset"] is None for row in self.dataset["products"]))
        self.assertTrue(all(row["example_products"] for row in self.dataset["products"]))

    def test_new_workbook_has_real_fifth_sheet_and_exact_linked_examples(self):
        raw = _workbook_bytes(self.issue, self.dataset, root=self.root)
        _workbook(raw, self.dataset)
        book = load_workbook(BytesIO(raw))
        try:
            self.assertEqual(book.sheetnames, ["Shortlist", "Evidence", "Supplier leads", "Checklist", "Example products"])
            self.assertEqual(tuple(cell.value for cell in book["Example products"][6]), WORKBOOK_HEADERS)
            self.assertTrue(book["Example products"]["I7"].hyperlink)
            self.assertIn("http", book["Example products"]["I7"].hyperlink.target)
        finally:
            book.close()

    def test_workbook_actual_example_tamper_is_refused(self):
        raw = _workbook_bytes(self.issue, self.dataset, root=self.root)
        book = load_workbook(BytesIO(raw)); book["Example products"]["E7"] = "Invented seller"
        output = BytesIO(); book.save(output); book.close()
        with self.assertRaisesRegex(SafetyError, "WORKBOOK_MISMATCH"):
            _workbook(output.getvalue(), self.dataset)

    def test_g_v4_both_languages_and_new_writer_manifest_are_connected(self):
        runner = WritingFixture()
        result = generate_report_narratives(self.memory, self.issue, self.dataset, runner)
        self.assertEqual(result["failures"], {}, result)
        self.assertEqual(result["complete_languages"], ["en", "hi-Latn"])
        written = result["dataset"]
        self.assertIn("Example products", render_html(self.issue, written, root=self.root))
        self.assertNotIn('src="None"', render_html(self.issue, written, root=self.root))
        draft = write_issue_package(self.issue, written, self.root, language="en")
        self.assertEqual(draft["schema_version"], 2)
        self.assertEqual(draft["workbook_contract"], "fit-checklist-examples-v2")
        self.assertEqual(draft["asset_receipt_hash"], digest(self.dataset["report_asset_receipt"]))

    def test_owned_receipt_revision_does_not_repeat_unchanged_writing(self):
        runner = WritingFixture()
        first = generate_report_narratives(self.memory, self.issue, self.dataset, runner)
        calls = runner.calls
        _, _, later_issue, later_data = functional_packet_fixture(self.root, revision=2, memory=self.memory)
        later = generate_report_narratives(self.memory, later_issue, later_data, runner, chunk_size=3)
        self.assertEqual(later["failures"], {})
        self.assertEqual(runner.calls, calls)
        self.assertEqual(first["dataset"]["narratives"], later["dataset"]["narratives"])
        self.assertNotEqual(first["facts_hash"], later["facts_hash"])
        receipts = [row for row in self.memory.snapshot()["source_checks"].values()
                    if row.get("kind") == "report-writing-reuse-v1" and row["revision"] == 2
                    and row["facts_hash"] == later["facts_hash"]
                    and row["job_id"] in {ref["job_id"] for refs in later["writing_item_refs"].values() for ref in refs.values()}]
        self.assertEqual(len(receipts), 22)
        repeated = generate_report_narratives(self.memory, later_issue, later_data, runner, chunk_size=3)
        self.assertEqual(repeated["failures"], {})
        self.assertEqual(runner.calls, calls)
        self.assertEqual(len([row for row in self.memory.snapshot()["source_checks"].values()
            if row.get("kind") == "report-writing-reuse-v1" and row["facts_hash"] == later["facts_hash"]
            and row["job_id"] in {ref["job_id"] for refs in later["writing_item_refs"].values() for ref in refs.values()}]), 22)

    def test_accepted_repair_reuses_with_explicit_new_packet_audit_receipt(self):
        bad_id = self.dataset["products"][0]["concept_id"]
        class RepairOnceWriter(WritingFixture):
            identity = "functional-accepted-repair-audit"
            def __call__(self, role, payload, folder):
                response, usage = super().__call__(role, payload, folder)
                if payload["language"] == "en" and "narration_repair" not in payload:
                    for row in response["results"]:
                        if row["id"] == bad_id:
                            row["references"] = ["invented-slot"]
                return response, usage
        runner = RepairOnceWriter()
        first = generate_report_narratives(self.memory, self.issue, self.dataset, runner)
        self.assertEqual(first["failures"], {})
        source = first["writing_item_refs"]["en"][bad_id]
        self.assertIn("narration_repair", self.memory.snapshot()["jobs"][source["job_id"]]["request"]["payload"])
        calls = runner.calls
        _, _, later_issue, later_data = functional_packet_fixture(self.root, revision=2, memory=self.memory)
        later = generate_report_narratives(self.memory, later_issue, later_data, runner)
        self.assertEqual(later["failures"], {})
        self.assertEqual(runner.calls, calls)
        receipts = [row for row in self.memory.snapshot()["source_checks"].values()
                    if row.get("kind") == "report-writing-reuse-v1" and row["revision"] == 2
                    and row["facts_hash"] == later["facts_hash"]
                    and row["job_id"] in {ref["job_id"] for refs in later["writing_item_refs"].values() for ref in refs.values()}]
        self.assertEqual(len(receipts), 22)
        self.assertTrue(any(row["job_id"] == source["job_id"] and row["item_id"] == bad_id for row in receipts))

    def test_legacy_or_unknown_tuple_cannot_claim_functional_packet(self):
        for field, changed in (("schema_version", 2), ("recommendation_contract", "unknown")):
            with self.subTest(field=field):
                dataset = deepcopy(self.dataset); dataset[field] = changed
                with self.assertRaises(SafetyError):
                    validate_dataset(self.issue, dataset)

    def test_workbook_can_finish_without_waiting_for_model_or_an_edition(self):
        stage = write_issue_workbook(self.issue, self.dataset, self.root)
        self.assertTrue(Path(stage["artifacts"]["workbook"]["path"]).is_file())
        self.assertFalse(stage.get("release_ready", False))

    def test_new_hold_requires_actual_owned_assets_and_frozen_input(self):
        memory, pool, issue, _ = functional_packet_fixture(self.root / "hold-test")
        state = memory.snapshot()
        options = {"expected_input_hash": issue["input_hash"],
            "expected_snapshot_hash": allocation_snapshot_hash(state), "as_of": NOW}
        with self.assertRaisesRegex(SafetyError, "ISSUE_REPORT_ASSETS_REQUIRED"):
            reserve_issue(memory, issue, **options)
        self.assertEqual(memory.snapshot()["issue_reservations"], state["issue_reservations"])
        facts = [pool["facts_by_concept"][row["concept_id"]] for row in issue["products"]]
        freeze_review_input(memory, issue, facts, root=self.root / "hold-test")
        saved = reserve_issue(memory, issue, report_root=self.root / "hold-test", **options)
        self.assertEqual(saved["selection_hash"], issue["selection_hash"])

    def test_actual_public_photo_survives_original_loss_in_render_and_workbook(self):
        from test_report_assets import image_bytes
        memory, _ = functional_pool_fixture(self.root / "photo-inputs", count=10, with_images=True)
        pool = build_report_pool(memory, as_of=NOW, recommendation_contract=RECOMMENDATION_CONTRACT)
        photo_root = self.root / "public-photos"
        raw, requests = image_bytes((400, 200)), []

        def fetch(url):
            requests.append(url)
            return {"state": "usable", "reason": "Synthetic original photo.", "checked_at": NOW,
                    "requested_url": url, "metadata": inspect_image(raw)}, raw

        candidate = pool["product_facts"][0]["image_candidates"][0]
        cid = pool["product_facts"][0]["concept_id"]
        capture_product_image(memory, cid, candidate, media_root=photo_root, fetcher=fetch)
        with patch("product_radar.product_images.MEDIA_ROOT", photo_root):
            pool = build_report_pool(memory, as_of=NOW, recommendation_contract=RECOMMENDATION_CONTRACT)
        issue = allocate_issues(pool["eligible"], [{"issue_id": "owned-photo-issue", "date": "2026-10-07",
            "client_id": "default", "revision": 1}], {}, {})["issues"][0]
        facts = [pool["facts_by_concept"][row["concept_id"]] for row in issue["products"]]
        issue, facts = _functional_report_revision(self.root, issue, facts, "synthetic", media_root=photo_root)
        dataset = build_issue_dataset(issue, facts, evidence_type="synthetic")
        photo = next(row for row in dataset["products"] if row["concept_id"] == cid)
        self.assertIsNotNone(photo["display_asset"])
        writer = RecordingWriter(); writer.identity = "functional-packet-owned-photo"
        written = generate_report_narratives(memory, issue, dataset, writer)["dataset"]
        Path(photo["image_provenance"]["path"]).unlink()
        for staged in (self.root / "media" / "live" / "report-assets-stage").rglob("*.jpg"):
            staged.unlink()
        with patch("product_radar.product_images.MEDIA_ROOT", photo_root):
            # Rebuild the actual pool after the shared original disappears.
            # Its default reader must omit the missing source image; recovery
            # may adopt only the exact already-owned revision's photograph.
            current_pool = build_report_pool(memory, as_of=NOW, recommendation_contract=RECOMMENDATION_CONTRACT)
        self.assertIsNone(current_pool["facts_by_concept"][cid].get("image_provenance"))
        original_issue = allocate_issues(current_pool["eligible"], [{"issue_id": "owned-photo-issue", "date": "2026-10-07",
            "client_id": "default", "revision": 1}], {}, {})["issues"][0]
        original_facts = [current_pool["facts_by_concept"][row["concept_id"]] for row in original_issue["products"]]
        recovered_issue, recovered_facts = _functional_report_revision(self.root, original_issue, original_facts,
            "synthetic", media_root=photo_root, source_state=memory.snapshot())
        self.assertEqual(recovered_issue["revision"], issue["revision"])
        self.assertEqual(recovered_facts, facts)
        validate_dataset(recovered_issue, build_issue_dataset(recovered_issue, recovered_facts, evidence_type="synthetic"), root=self.root)
        validate_dataset(issue, written, root=self.root)
        self.assertIn("data:image/jpeg;base64,", render_html(issue, written, root=self.root))
        book = load_workbook(BytesIO(_workbook_bytes(issue, written, root=self.root)))
        try:
            self.assertGreaterEqual(len(book["Shortlist"]._images), 1)
        finally:
            book.close()
        self.assertEqual(len(requests), 1)

    def test_owned_assets_without_input_resume_same_revision_then_freeze_before_hold(self):
        memory, _ = functional_pool_fixture(self.root / "orphan-inputs", count=10)
        pool = build_report_pool(memory, as_of=NOW, recommendation_contract=RECOMMENDATION_CONTRACT)
        issue = allocate_issues(pool["eligible"], [{"issue_id": "orphan-asset-issue", "date": "2026-10-07",
            "client_id": "default", "revision": 1}], {}, {})["issues"][0]
        facts = [pool["facts_by_concept"][row["concept_id"]] for row in issue["products"]]
        first, frozen = _functional_report_revision(self.root, issue, facts, "synthetic")
        again, repeated = _functional_report_revision(self.root, issue, facts, "synthetic")
        self.assertEqual(first, again)
        self.assertEqual(frozen, repeated)
        self.assertEqual(memory.snapshot()["issue_reservations"], {})
        freeze_review_input(memory, first, frozen, root=self.root)
        current, current_facts = _functional_report_revision(self.root, issue, facts, "synthetic")
        self.assertEqual((current, current_facts), (first, frozen))
        reserve_issue(memory, first, report_root=self.root, as_of=NOW,
            expected_input_hash=first["input_hash"], expected_snapshot_hash=allocation_snapshot_hash(memory.snapshot()))

    def test_programme_joins_real_new_pool_assets_hold_and_bilingual_draft(self):
        memory, ids = functional_pool_fixture(self.root / "programme-inputs", count=10)
        record = prepare_programme(memory, "functional-programme", source_run_id="functional-source-run",
                                   concept_ids=ids, as_of=NOW)
        runner, renders = RecordingWriter(), []
        runner.identity = "functional-programme-writer"

        def render(html, pdf):
            renders.append(str(pdf))
            with fitz.open() as document:
                page = document.new_page(width=595, height=842)
                page.insert_text((40, 40), "Synthetic functional programme, not a client report.")
                raw = document.tobytes()
            atomic_bytes(pdf, raw)
            from product_radar.live_outputs import digest_bytes
            return dict(inspect(raw), path=str(pdf), html_sha256=digest_bytes(html.read_bytes()), pdf_sha256=digest_bytes(raw))

        context = {"at": NOW, "root": self.root, "runner": runner, "render_pdf": render,
                   "inspect_pdf": inspect, "research_advancing": False, "retry_seconds": 3600}
        from product_radar.work_turns import bind_work_turn, open_work_turn
        turn = open_work_turn(memory, record["programme_id"], at=NOW, report_ready=True)
        with bind_work_turn(memory, turn["id"]):
            first = _reports(memory, record, context)
            calls = runner.calls
            again = _reports(memory, record, context)
        completed = [row for row in first["outcomes"] if row.get("complete")]
        self.assertEqual(len(completed), 1, first)
        self.assertEqual(completed[0]["state"], "draft_ready_for_review")
        self.assertEqual(len(renders), 2)
        self.assertEqual(runner.calls, calls)
        self.assertEqual(len(memory.snapshot()["issue_reservations"]), 1)
        self.assertEqual(memory.snapshot()["published_history"], {})
        manifest = load_json(completed[0]["manifest_path"], 512 * 1024)
        self.assertEqual(manifest["schema_version"], 2)
        self.assertEqual(manifest["recommendation_contract"], RECOMMENDATION_CONTRACT)
        self.assertTrue(any(row.get("complete") for row in again["outcomes"]))

    def test_one_bad_english_row_does_not_block_healthy_translations_or_repeat_them(self):
        # A different v4 caller now correctly reuses exact healthy originals.
        # Start this bad-answer scenario without unrelated saved writing.
        self.memory, self.pool, self.issue, self.dataset = functional_packet_fixture(self.root / "invalid-row")
        bad_id = self.dataset["products"][0]["concept_id"]
        runner = RecordingWriter(bad_id)
        runner.identity = "functional-packet-recording-writer-one-invalid-row"
        first = generate_report_narratives(self.memory, self.issue, self.dataset, runner)
        self.assertEqual(set(first["failures"]["en"]), {bad_id})
        self.assertEqual(set(first["failures"]["hi-Latn"]), {bad_id})
        self.assertEqual(first["complete_languages"], [])
        self.assertEqual(len(first["writing_item_refs"]["en"]), 10)
        self.assertEqual(len(first["writing_item_refs"]["hi-Latn"]), 10)
        translated = [item for payload in runner.payloads if payload["language"] == "hi-Latn" for item in payload["items"]]
        self.assertEqual(len(translated), 10)
        self.assertNotIn(bad_id, [row["id"] for row in translated])
        self.assertTrue(all(row["source_english"] is not None for row in translated))
        self.assertTrue(all(row["source_english_ref"]["job_id"] in self.memory.snapshot()["jobs"] for row in translated))
        calls = runner.calls
        again = generate_report_narratives(self.memory, self.issue, self.dataset, runner, chunk_size=2)
        self.assertEqual(runner.calls, calls)
        self.assertEqual(again["failures"], first["failures"])
        self.assertEqual(first["writing_item_refs"], again["writing_item_refs"])

    def test_photo_hashes_and_unrelated_products_are_not_writing_cache_inputs(self):
        payload = make_narrative_input(self.dataset, "en")
        self.assertNotIn("facts_hash", payload)
        self.assertNotIn("selection_hash", payload)
        self.assertEqual(payload["writing_contract"], "writing-input-v1")
        altered = deepcopy(self.dataset)
        altered["revision"] += 1
        altered["selection_hash"] = "0" * 64
        altered["report_asset_receipt"] = {"synthetic": "not supplied to the model"}
        self.assertEqual(payload, make_narrative_input(altered, "en"))

    def test_one_reuse_receipt_save_failure_preserves_all_completed_rows(self):
        runner = RecordingWriter()
        runner.identity = "functional-packet-reuse-local-save-failure"
        first = generate_report_narratives(self.memory, self.issue, self.dataset, runner)
        self.assertFalse(first["failures"])
        calls, original_put = runner.calls, self.memory.put
        bad_id = self.dataset["products"][0]["concept_id"]

        def put(kind, record_id, record):
            if (kind == "source_checks" and record.get("kind") == "report-writing-reuse-v1"
                    and record.get("item_id") == bad_id):
                raise OSError("synthetic local file failure")
            return original_put(kind, record_id, record)

        with patch.object(self.memory, "put", side_effect=put):
            again = generate_report_narratives(self.memory, self.issue, self.dataset, runner)
        self.assertEqual(runner.calls, calls)
        self.assertEqual(again["failures"], {})
        self.assertEqual(again["complete_languages"], ["en", "hi-Latn"])
        self.assertEqual(len(again["writing_item_refs"]["en"]), 11)
        self.assertEqual(set(again["problem_save_failures"]["en"]), {bad_id})

    def test_pending_title_keeps_actual_healthy_original_job_owners(self):
        # An earlier healthy title must not replace this intentionally pending
        # original merely because the class shares its default research fixture.
        self.memory, self.pool, self.issue, self.dataset = functional_packet_fixture(self.root / "pending-title")
        class PendingWriter(RecordingWriter):
            identity = "functional-packet-pending-title"

            def __call__(self, role, payload, folder):
                raise TimeoutError("synthetic unresolved original")

        payload = make_narrative_input(self.dataset, "en")
        with self.assertRaises(SafetyError):
            Jobs(self.memory, PendingWriter()).run("G", dict(payload, items=payload["items"][:1]))
        runner = RecordingWriter(); runner.identity = PendingWriter.identity
        result = generate_report_narratives(self.memory, self.issue, self.dataset, runner)
        self.assertEqual(set(result["failures"]["en"]), {payload["items"][0]["id"]})
        self.assertEqual(len(result["writing_item_refs"]["en"]), 10)
        self.assertEqual(len(result["writing_item_refs"]["hi-Latn"]), 10)
        self.assertTrue(all(ref["job_id"] in self.memory.snapshot()["jobs"]
                            for refs in result["writing_item_refs"].values() for ref in refs.values()))
        self.assertTrue(all(item["kind"] != "report_title" for request in runner.payloads for item in request["items"]))
        calls = runner.calls
        again = generate_report_narratives(self.memory, self.issue, self.dataset, runner, chunk_size=2)
        self.assertEqual(runner.calls, calls)
        self.assertEqual(again["writing_item_refs"], result["writing_item_refs"])

    def test_hinglish_only_keeps_internal_english_provenance_without_an_english_delivery(self):
        runner = RecordingWriter(); runner.identity = "functional-packet-hi-only"
        result = generate_report_narratives(self.memory, self.issue, self.dataset, runner, languages=("hi-Latn",))
        self.assertEqual(result["complete_languages"], ["hi-Latn"])
        self.assertEqual(set(result["dataset"]["narratives"]), {"en", "hi-Latn"})
        payload = make_narrative_input(result["dataset"], "hi-Latn", english=result["dataset"]["narratives"]["en"])
        self.assertTrue(all("source_english_ref" in item for item in payload["items"]))
        altered = deepcopy(result["dataset"])
        altered["narratives"]["en"]["products"][0]["reason"] = "Invented {{title}} description."
        with self.assertRaisesRegex(SafetyError, "REPORT_WRITING_ACCEPTANCE_INVALID"):
            validate_dataset(self.issue, altered)

    def test_actual_workflow_builds_new_delivery_tuple_and_reuses_exact_files(self):
        runner, renders = RecordingWriter(), []

        def render(html, pdf):
            renders.append(str(pdf))
            with fitz.open() as document:
                page = document.new_page(width=595, height=842)
                page.insert_text((40, 40), "Synthetic functional report workflow; not a client report.")
                raw = document.tobytes()
            atomic_bytes(pdf, raw)
            from product_radar.live_outputs import digest_bytes
            return dict(inspect(raw), path=str(pdf), html_sha256=digest_bytes(html.read_bytes()), pdf_sha256=digest_bytes(raw))

        facts = [self.pool["facts_by_concept"][row["concept_id"]] for row in self.issue["products"]]
        first = prepare_review_package(self.memory, self.issue, facts, root=self.root,
            runner=runner, render_pdf=render, inspect_pdf=inspect)
        self.assertTrue(first["complete"], first)
        manifest = load_json(first["manifest_path"], 512 * 1024)
        self.assertEqual(manifest["schema_version"], 2)
        self.assertEqual(manifest["recommendation_contract"], RECOMMENDATION_CONTRACT)
        self.assertEqual(manifest["workbook_contract"], "fit-checklist-examples-v2")
        self.assertEqual(len(manifest["files"]), 3)
        self.assertFalse(manifest["approved"])
        calls = runner.calls
        again = prepare_review_package(self.memory, self.issue, facts, root=self.root,
            runner=runner, render_pdf=render, inspect_pdf=inspect)
        self.assertTrue(again["complete"])
        self.assertEqual(again["package_hash"], first["package_hash"])
        self.assertEqual(runner.calls, calls)
        self.assertEqual(len(renders), 2)


if __name__ == "__main__":
    unittest.main()
