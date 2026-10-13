"""Saved-route report contracts, using only disposable synthetic research."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))

from test_fit_pipeline import NOW, current_fixture, WritingFixture
from test_report_selection import OBSERVED
from test_live_outputs import fixture as legacy_fixture
from test_trends_browser import metadata, weekly
from product_radar.delivery_package import _workbook
from product_radar.fit_policy import FIT_POLICY_ID
from product_radar.issue_allocation import allocate_issues, allocation_snapshot_hash, reserve_issue
from product_radar.live_outputs import (LOGO, build_issue_dataset, validate_dataset, render_html,
                                       _workbook_bytes, display_market_checks, write_issue_workbook, write_issue_edition,
                                       SOURCE_WORKBOOK_GENERATION_VERSION, SOURCE_HTML_RENDER_VERSION,
                                       WORKBOOK_GENERATION_VERSION, HTML_RENDER_VERSION)
from product_radar.models import SafetyError, canonical, digest
from product_radar.market_validation import REVIEW_VERSION
from product_radar.providers.trends_browser import parse_trends_csv
from product_radar.report_narration import generate_report_narratives, make_narrative_input
from product_radar.report_selection import build_report_pool, selection_scope
from product_radar.report_source_facts import CONTRACT, FIELDS, ROUTES, validate_route_facts
from product_radar.research_memory import ResearchMemory, change
from product_radar.source_coverage import prepare_source_tasks, execute_source_run, source_coverage
from product_radar.storage import Journal
from product_radar.trends import compute_trends_signal


def source_fixture(count=10):
    memory = current_fixture(count)
    for source, route in (("meta", "meta-ads"), ("amazon", "amazon-bestsellers")):
        capture_id = "fixture-capture-" + source
        memory.state["captures"][capture_id] = {"source": source, "acquisition_id": "fixture-" + source,
            "observed_at": OBSERVED, "reference": {"route": route}}
        for row in memory.state["observations"].values():
            if row["source"] == source:
                row["capture_id"] = capture_id
    return memory


def select(memory):
    pool = build_report_pool(memory, NOW, fit_policy_id=FIT_POLICY_ID)
    issue = allocate_issues(pool["eligible"], ["2026-10-07"], [], [])["issues"][0]
    facts = [pool["facts_by_concept"][row["concept_id"]] for row in issue["products"]]
    return pool, issue, build_issue_dataset(issue, facts, evidence_type="synthetic")


class SavedSourceReportTests(unittest.TestCase):
    def real_memory(self, directory, count=10):
        memory = ResearchMemory(Path(directory) / "research", Journal(Path(directory) / "finance"), evidence_type="synthetic")
        state = source_fixture(count).state
        memory.commit("synthetic-report-source-seed", [change(table, record_id, row)
            for table, values in state.items() for record_id, row in values.items()])
        return memory

    def question(self, route, concept_id="c000", query="fixture product"):
        return {"route": route, "request": {"query": query, "country": "IN"}, "concept_ids": [concept_id]}

    def collect(self, memory, questions, state="blocked", *, failure_scope=None, actual=True):
        prepare_source_tasks(memory, questions, run_id="synthetic-report-routes", as_of=NOW)
        def handler(context):
            if actual:
                context["record_dispatch"]()
            return {"schema_version": 1, "route": context["route"], "acquisition_id": context["acquisition_id"],
                "observed_at": NOW, "state": state, "dispatch_started": actual,
                "reference": context["acquisition_id"] + "/synthetic-receipt.json", "capture_ids": [],
                "observation_ids": [], "target_results": [], "reason": "Synthetic route outcome only.",
                **({"failure_scope": failure_scope} if failure_scope else {})}
        with patch("product_radar.source_coverage.utc_now", return_value=NOW):
            return execute_source_run(memory, "synthetic-report-routes", as_of=NOW,
                handlers={q["route"]: handler for q in questions}, bounded_call=lambda callback, context, timeout: callback(context))

    def product(self, data, concept_id="c000"):
        return next(row for row in data["products"] if row["concept_id"] == concept_id)

    def route(self, product, route):
        return next(row for row in product["source_route_checks"] if row["route"] == route)

    def test_all_fourteen_routes_are_saved_factual_blocks_without_new_collection(self):
        memory = source_fixture()
        pool, issue, dataset = select(memory)
        product = self.product(dataset)
        self.assertEqual(product["source_route_contract"], CONTRACT)
        self.assertEqual([row["route"] for row in product["source_route_checks"]], list(ROUTES))
        self.assertEqual(self.route(product, "meta-ads")["outcome"], "completed_relevant")
        self.assertIs(self.route(product, "meta-ads")["actual_attempt"], True)
        self.assertEqual(self.route(product, "blinkit")["outcome"], "not_attempted")
        self.assertIs(self.route(product, "blinkit")["actual_attempt"], False)
        self.assertEqual(memory.reads, 1)
        self.assertEqual(memory.state["jobs"], {})
        frozen = next(row for row in issue["products"] if row["concept_id"] == "c000")
        self.assertEqual(frozen["source_route_checks"], product["source_route_checks"])
        self.assertEqual(frozen["source_scope_hash"], digest(selection_scope(memory.state, "c000", source_route_contract=CONTRACT)))
        self.assertEqual(pool["facts_by_concept"]["c000"]["source_route_checks"], product["source_route_checks"])

    def test_actual_blocked_request_and_untried_shared_pause_have_different_answers(self):
        with tempfile.TemporaryDirectory() as directory:
            memory = self.real_memory(directory)
            result = self.collect(memory, [self.question("blinkit"), self.question("blinkit", "c001", "other product")],
                                  failure_scope="route")
            self.assertEqual(result["handler_calls"], 1)
            _, _, dataset = select(memory)
            first, second = (self.route(self.product(dataset, cid), "blinkit") for cid in ("c000", "c001"))
            attempted, untried = (first, second) if first["actual_attempt"] else (second, first)
            self.assertEqual(attempted["state"], "Check blocked")
            self.assertIs(attempted["actual_attempt"], True)
            self.assertEqual(untried["state"], "Not tried")
            self.assertIs(untried["actual_attempt"], False)
            self.assertEqual(untried["outcome"], "blocked")
            self.assertTrue(untried["receipt_ids"])

    def test_completed_empty_is_not_described_as_untried_or_a_failed_product(self):
        with tempfile.TemporaryDirectory() as directory:
            memory = self.real_memory(directory)
            self.collect(memory, [self.question("flipkart-search")], state="empty")
            _, _, dataset = select(memory)
            row = self.route(self.product(dataset), "flipkart-search")
            self.assertEqual(row["outcome"], "completed_no_results")
            self.assertEqual(row["state"], "No usable result found")
            self.assertIs(row["actual_attempt"], True)
            self.assertEqual(row["measurements"], [])

    def test_saved_acquisition_without_match_stays_review_pending(self):
        with tempfile.TemporaryDirectory() as directory:
            memory = self.real_memory(directory)
            self.collect(memory, [self.question("flipkart-search")], state="usable")
            _, _, dataset = select(memory)
            row = self.route(self.product(dataset), "flipkart-search")
            self.assertEqual(row["outcome"], "interpretation_pending")
            self.assertEqual(row["state"], "Results need review")
            self.assertIs(row["actual_attempt"], True)
            self.assertEqual(row["measurements"], [])

    def test_report_view_reads_no_external_staged_progress_and_default_live_view_is_unchanged(self):
        with tempfile.TemporaryDirectory() as directory:
            memory = self.real_memory(directory)
            prepare_source_tasks(memory, [self.question("blinkit")], run_id="synthetic-report-routes", as_of=NOW)
            with patch("product_radar.source_coverage._effective_route_pauses", side_effect=AssertionError("No staged-file read")):
                _, _, dataset = select(memory)
                self.assertEqual(self.route(self.product(dataset), "blinkit")["state"], "Not tried")
                with self.assertRaisesRegex(AssertionError, "No staged-file read"):
                    source_coverage(memory, concept_ids=["c000"], as_of=NOW)
            with self.assertRaisesRegex(SafetyError, "SOURCE_COVERAGE_OPTIONS_INVALID"):
                source_coverage(memory, include_staged_progress=0)

    def test_future_receipt_does_not_turn_an_earlier_pending_attempt_into_a_completed_check(self):
        with tempfile.TemporaryDirectory() as directory:
            memory = self.real_memory(directory)
            self.collect(memory, [self.question("flipkart-search")], state="empty")
            state = memory.snapshot()
            task_id = next(task_id for task_id, task in state["query_tasks"].items() if task["route"] == "flipkart-search")
            receipt_id = state["query_tasks"][task_id]["receipt_refs"][-1]
            old = state["source_checks"][receipt_id]
            later = deepcopy(old)
            later["at"] = "2026-10-06T11:00:00+00:00"
            later["envelope"]["observed_at"] = later["at"]
            memory.put("source_checks", receipt_id, later, previous=old)
            _, _, dataset = select(memory)
            row = self.route(self.product(dataset), "flipkart-search")
            self.assertEqual(row["outcome"], "in_progress")
            self.assertEqual(row["state"], "Check in progress")
            self.assertIs(row["actual_attempt"], True)
            self.assertEqual(row["receipt_ids"], [])
            self.assertIsNone(row["last_attempt"])

    def test_future_task_and_applicability_do_not_supply_an_earlier_answer(self):
        with tempfile.TemporaryDirectory() as directory:
            memory = self.real_memory(directory)
            future = "2026-10-06T11:00:00+00:00"
            prepare_source_tasks(memory, [self.question("zepto"), {"route": "company-pages", "concept_ids": ["c000"],
                "not_applicable_reason": "Synthetic later decision."}], run_id="later-questions", as_of=future)
            _, _, dataset = select(memory)
            for route in ("zepto", "company-pages"):
                row = self.route(self.product(dataset), route)
                self.assertEqual(row["outcome"], "not_attempted")
                self.assertEqual(row["task_ids"], [])
                self.assertIs(row["actual_attempt"], False)

    def test_ad_counts_are_sample_activity_not_sales_and_unknown_is_not_zero(self):
        memory = source_fixture()
        memory.state["observations"]["meta000"]["facts"].update(active=None, advertiser_id=None)
        _, _, dataset = select(memory)
        row = self.route(self.product(dataset), "meta-ads")
        metrics = {metric["name"]: metric["value"] for metric in row["measurements"]}
        self.assertEqual(metrics["relevant_ads"], 1)
        self.assertIsNone(metrics["advertisers"])
        self.assertIsNone(metrics["active_ads"])
        self.assertEqual(metrics["unknown_active_status"], 1)
        self.assertEqual(metrics["missing_advertiser_ids"], 1)
        self.assertIn("Active".casefold(), row["detail"].casefold())
        self.assertIn("not stated", row["detail"])
        self.assertNotIn("units sold", row["detail"])

    def test_same_semantic_refresh_reuses_the_accepted_match_but_changed_semantics_stay_pending(self):
        for changed in (False, True):
            with self.subTest(changed=changed):
                memory = source_fixture()
                observation = deepcopy(memory.state["observations"]["meta000"])
                observation.update(observed_at=NOW, capture_id="new-meta-capture", acquisition_id="new-meta-acquisition")
                if changed:
                    observation["semantic_id"] = "genuinely-changed-semantic"
                    observation["semantic"] = {"title": "Synthetic unrelated replacement description"}
                memory.state["observations"]["fresh-meta"] = observation
                memory.state["captures"]["new-meta-capture"] = {"source": "meta", "acquisition_id": "new-meta-acquisition",
                    "observed_at": NOW, "reference": {"route": "meta-ads"}}
                _, _, dataset = select(memory)
                row = self.route(self.product(dataset), "meta-ads")
                self.assertEqual(row["outcome"], "interpretation_pending" if changed else "completed_relevant")
                self.assertIs(row["actual_attempt"], True)
                if changed:
                    self.assertEqual(row["measurements"], [])
                else:
                    self.assertEqual(row["measurements"][0]["value"], 1)
                    self.assertIn("fresh-meta", row["measurements"][0]["evidence_ids"])

    def test_older_actual_check_stays_tried_without_becoming_a_current_market_vote(self):
        from product_radar.discovery import latest_observations
        memory = source_fixture()
        old = "2026-09-15T10:00:00+00:00"
        memory.state["observations"]["meta000"]["observed_at"] = old
        memory.state["captures"]["fixture-capture-meta"]["observed_at"] = old
        _, _, dataset = select(memory)
        product = self.product(dataset)
        row = self.route(product, "meta-ads")
        self.assertEqual(row["outcome"], "completed_relevant")
        self.assertIs(row["actual_attempt"], True)
        self.assertEqual(row["observed_at"], old)
        self.assertEqual(row["measurements"], [])
        self.assertNotIn("meta", product["opportunity_basis"]["source_ids"])
        self.assertNotIn("meta000", latest_observations(memory.state["observations"], NOW))
        self.assertIn("meta000", latest_observations(memory.state["observations"], NOW, max_age_days=None))
        for option in (False, 0, -1):
            with self.subTest(option=option), self.assertRaisesRegex(SafetyError, "OBSERVATION_AGE_OPTION_INVALID"):
                latest_observations(memory.state["observations"], NOW, max_age_days=option)

    def add_context(self, memory, source, ref, *, facts, source_item=None):
        capture_id = "capture-" + ref
        semantic_id = "semantic-" + ref
        observation = {"source": source, "source_item": source_item or ref, "capture_id": capture_id,
            "acquisition_id": "acquisition-" + ref, "observed_at": OBSERVED, "semantic_id": semantic_id,
            "semantic": {"title": "Synthetic related product type"}, "facts": deepcopy(facts)}
        memory.state["captures"][capture_id] = {"source": source, "acquisition_id": observation["acquisition_id"],
            "observed_at": OBSERVED, "reference": {"route": facts.get("source_route", source)}}
        memory.state["observations"][ref] = observation
        memory.state["market_reviews"]["review-" + ref] = {"concept_id": "c000", "review_version": REVIEW_VERSION,
            "description_hash": digest(memory.state["concepts"]["c000"]["description"]), "reviewed_at": NOW,
            "input_hash": digest(observation), "relevance": [{"id": ref, "relationship": "relevant_family", "reason": "Synthetic matching type."}]}

    def test_amazon_list_rank_and_search_order_do_not_become_sales(self):
        memory = source_fixture()
        self.add_context(memory, "amazon", "amazon-search-context", facts={"url": "https://example.org/search-item",
            "source_route": "amazon-search", "list_type": "product_search", "position": 99,
            "price": {"displayed": "₹299", "currency": "INR"}, "limitations": ["product_search_not_sales_proof"]})
        _, _, dataset = select(memory)
        product = self.product(dataset)
        listing = self.route(product, "amazon-bestsellers")
        generic = self.route(product, "amazon-search")
        self.assertEqual(next(m["value"] for m in listing["measurements"] if m["name"] == "list_rank_bounds"), [1, 1])
        self.assertIn("not units sold", listing["detail"])
        self.assertNotIn("list_rank_bounds", {m["name"] for m in generic["measurements"]})
        self.assertEqual(next(m["value"] for m in generic["measurements"] if m["name"] == "relevant_listings"), 1)
        self.assertEqual(product["opportunity_basis"]["source_ids"], ["amazon", "meta"])

    def test_current_unrelated_and_unclear_are_reviewed_outcomes_while_retired_reviews_need_review(self):
        for relationship, version, outcome in (("unrelated", REVIEW_VERSION, "completed_no_results"),
                                               ("unclear", REVIEW_VERSION, "insufficient"),
                                               ("relevant_family", "marketplace-relevance-v1", "interpretation_pending")):
            with self.subTest(relationship=relationship, version=version):
                memory = source_fixture()
                self.add_context(memory, "flipkart", "flipkart-related", facts={"url": "https://example.org/flipkart",
                    "source_route": "flipkart-search", "list_type": "product_search", "price": {"displayed": "₹299", "currency": "INR"}})
                review = memory.state["market_reviews"]["review-flipkart-related"]
                review["review_version"] = version
                review["relevance"][0]["relationship"] = relationship
                _, _, dataset = select(memory)
                row = self.route(self.product(dataset), "flipkart-search")
                self.assertEqual(row["outcome"], outcome)
                self.assertIs(row["actual_attempt"], True)
                self.assertEqual(row["measurements"], [])

    def test_latest_current_matching_review_wins_not_dictionary_insertion_order(self):
        memory = source_fixture()
        self.add_context(memory, "flipkart", "flipkart-related", facts={"url": "https://example.org/flipkart",
            "source_route": "flipkart-search", "list_type": "product_search", "price": {"displayed": None, "currency": "INR"}})
        current = memory.state["market_reviews"]["review-flipkart-related"]
        current["relevance"][0]["relationship"] = "unrelated"
        older = deepcopy(current)
        older["reviewed_at"] = OBSERVED
        older["relevance"][0]["relationship"] = "relevant_family"
        memory.state["market_reviews"]["older-inserted-later"] = older
        _, _, dataset = select(memory)
        self.assertEqual(self.route(self.product(dataset), "flipkart-search")["outcome"], "completed_no_results")

    def test_quickcommerce_preserves_available_unavailable_unknown_and_checked_location(self):
        memory = source_fixture()
        for n, value in enumerate((True, False, None)):
            self.add_context(memory, "blinkit", "blinkit-context-" + str(n), facts={"source_route": "blinkit",
                "url": "https://example.org/blinkit/item-" + str(n), "list_type": "quickcommerce_context",
                "available": value, "price": {"displayed": None, "currency": "INR"}, "image_urls": [],
                "location": {"name": "Synthetic Delhi location"}, "limitations": []})
        _, _, dataset = select(memory)
        row = self.route(self.product(dataset), "blinkit")
        metrics = {metric["name"]: metric["value"] for metric in row["measurements"]}
        self.assertEqual((metrics["relevant_listings"], metrics["available_listings"], metrics["unavailable_listings"],
                          metrics["unknown_availability"]), (3, 1, 1, 1))
        self.assertEqual(metrics["displayed_prices"], 0)
        self.assertIn("Synthetic Delhi location", row["detail"])
        self.assertNotIn("sales", row["detail"])

    def test_combined_quickcommerce_counts_keep_all_actual_locations_and_same_item_checks(self):
        memory = source_fixture()
        for name, available in (("Synthetic Delhi", True), ("Synthetic Mumbai", False)):
            self.add_context(memory, "blinkit", "blinkit-" + name.split()[-1], source_item="same-blinkit-item", facts={
                "source_route": "blinkit", "url": "https://example.org/blinkit", "list_type": "quickcommerce_context",
                "available": available, "price": {"displayed": None, "currency": "INR"},
                "location": {"name": name}, "limitations": []})
        _, _, dataset = select(memory)
        row = self.route(self.product(dataset), "blinkit")
        metrics = {metric["name"]: metric["value"] for metric in row["measurements"]}
        self.assertEqual((metrics["relevant_listings"], metrics["available_listings"], metrics["unavailable_listings"]), (2, 1, 1))
        self.assertIn("Combined saved sample", row["detail"])
        self.assertIn("Synthetic Delhi", row["detail"])
        self.assertIn("Synthetic Mumbai", row["detail"])
        self.assertNotIn("Checked location: Synthetic Delhi", row["detail"])
        self.assertEqual(self.product(dataset)["opportunity_basis"]["source_ids"].count("blinkit"), 1)

    def test_quickcommerce_same_location_replacement_is_only_counted_once(self):
        memory = source_fixture()
        for ref, available in (("blinkit-older", True), ("blinkit-newer", False)):
            self.add_context(memory, "blinkit", ref, source_item="same-blinkit-item", facts={"source_route": "blinkit",
                "url": "https://example.org/blinkit", "list_type": "quickcommerce_context", "available": available,
                "price": {"displayed": None, "currency": "INR"}, "location": {"name": "Synthetic Delhi"}, "limitations": []})
        memory.state["observations"]["blinkit-newer"]["observed_at"] = NOW
        _, _, dataset = select(memory)
        row = self.route(self.product(dataset), "blinkit")
        metrics = {metric["name"]: metric["value"] for metric in row["measurements"]}
        self.assertEqual((metrics["relevant_listings"], metrics["available_listings"], metrics["unavailable_listings"]), (1, 0, 1))
        self.assertEqual(row["measurements"][0]["evidence_ids"], ["blinkit-newer"])

    def test_trends_shows_actual_relative_interest_numbers_not_a_sales_estimate(self):
        memory = source_fixture()
        parsed = parse_trends_csv(weekly(), metadata(collected_at=OBSERVED))
        memory.state["trends_evidence"]["synthetic-chart"] = {"concept_id": "c000",
            "description_hash": digest(memory.state["concepts"]["c000"]["description"]), "parsed": parsed}
        _, _, dataset = select(memory)
        row = self.route(self.product(dataset), "google-trends")
        metrics = {metric["name"]: metric["value"] for metric in row["measurements"]}
        self.assertEqual(metrics["chart_intervals"], len(parsed["series"]))
        self.assertIn("previous_mean_bounds", metrics)
        self.assertIn("recent_mean_bounds", metrics)
        self.assertIn("relative interest index, not sales", row["detail"])

    def test_fractional_trends_numbers_are_unrounded_decimal_text_through_workbook_and_hashes(self):
        memory = source_fixture()
        parsed = parse_trends_csv(weekly([20] * 45 + [20, 20, 20, 21] + [30, 30, 30, 31] + [100]), metadata(collected_at=OBSERVED))
        memory.state["trends_evidence"]["fractional-chart"] = {"concept_id": "c000",
            "description_hash": digest(memory.state["concepts"]["c000"]["description"]), "parsed": parsed}
        original = compute_trends_signal(parsed, now=NOW)
        _, _, dataset = select(memory)
        product = self.product(dataset)
        metrics = {metric["name"]: metric["value"] for metric in self.route(product, "google-trends")["measurements"]}
        self.assertEqual(metrics["previous_mean_bounds"], ["20.25", "20.25"])
        self.assertEqual(metrics["recent_mean_bounds"], ["30.25", "30.25"])
        self.assertEqual(original["previous_mean_bounds"], [20.25, 20.25])
        self.assertEqual(compute_trends_signal(parsed, now=NOW), original)
        self.assertIn("20.25", canonical(product["source_route_checks"]).decode())
        validate_route_facts(product)
        _workbook(_workbook_bytes(select(memory)[1], dataset), dataset)
        self.assertEqual(digest(product["source_route_checks"]), digest(deepcopy(product["source_route_checks"])))

    def test_changed_used_receipt_invalidates_only_affected_product_scope_and_draft(self):
        with tempfile.TemporaryDirectory() as directory:
            memory = self.real_memory(directory)
            self.collect(memory, [self.question("blinkit")])
            pool, issue, dataset = select(memory)
            state = memory.snapshot()
            before_first = digest(selection_scope(state, "c000", source_route_contract=CONTRACT))
            before_other = digest(selection_scope(state, "c001", source_route_contract=CONTRACT))
            row = self.route(self.product(dataset), "blinkit")
            receipt_id = row["receipt_ids"][0]
            self.assertEqual(self.product(dataset)["source_revisions"]["source_checks"][receipt_id], digest(state["source_checks"][receipt_id]))
            snapshot_hash = allocation_snapshot_hash(state)
            old = state["source_checks"][receipt_id]
            memory.put("source_checks", receipt_id, dict(old, reason="Changed synthetic saved receipt."), previous=old)
            changed = memory.snapshot()
            self.assertNotEqual(digest(selection_scope(changed, "c000", source_route_contract=CONTRACT)), before_first)
            self.assertEqual(digest(selection_scope(changed, "c001", source_route_contract=CONTRACT)), before_other)
            with self.assertRaisesRegex(SafetyError, "ISSUE_SOURCE_FACTS_CHANGED"):
                reserve_issue(memory, issue, expected_input_hash=issue["input_hash"], expected_snapshot_hash=snapshot_hash)

    def test_new_task_or_applicability_invalidates_scoped_facts_but_other_product_does_not(self):
        with tempfile.TemporaryDirectory() as directory:
            memory = self.real_memory(directory)
            before = digest(selection_scope(memory.snapshot(), "c000", source_route_contract=CONTRACT))
            prepare_source_tasks(memory, [self.question("zepto", "c001")], run_id="other-product", as_of=NOW)
            self.assertEqual(digest(selection_scope(memory.snapshot(), "c000", source_route_contract=CONTRACT)), before)
            prepare_source_tasks(memory, [{"route": "company-pages", "concept_ids": ["c000"],
                "not_applicable_reason": "No supplier company link exists in this synthetic scope."}], run_id="own-product", as_of=NOW)
            self.assertNotEqual(digest(selection_scope(memory.snapshot(), "c000", source_route_contract=CONTRACT)), before)

    def test_frozen_route_facts_override_untrusted_supplied_rows_and_refuse_after_build_changes(self):
        _, issue, dataset = select(source_fixture())
        facts = deepcopy(dataset["products"])
        self.product({"products": facts})["source_route_checks"][0]["state"] = "Forged measured sales"
        built = build_issue_dataset(issue, facts, evidence_type="synthetic")
        self.assertEqual(built["products"][0]["source_route_checks"], issue["products"][0]["source_route_checks"])
        built["products"][0]["source_route_checks"][0]["actual_attempt"] = False
        with self.assertRaisesRegex(SafetyError, "LIVE_SOURCE_ROUTE_FACTS_CHANGED"):
            validate_dataset(issue, built)

    def test_unbound_photo_is_refused_both_at_build_and_direct_dataset_validation(self):
        _, issue, dataset = select(source_fixture())
        for changes in ({"image_path": str(LOGO)}, {"image_provenance": {"path": str(LOGO)}},
                        {"image_path": str(LOGO), "image_provenance": {"path": str(LOGO)}}):
            with self.subTest(changes=changes):
                bad = deepcopy(dataset)
                bad["products"][0].update(changes)
                with self.assertRaisesRegex(SafetyError, "LIVE_PHOTO_BINDING_INVALID"):
                    validate_dataset(issue, bad)
                with self.assertRaisesRegex(SafetyError, "LIVE_PHOTO_BINDING_INVALID"):
                    build_issue_dataset(issue, bad["products"], evidence_type="synthetic")
        self.assertIsNone(dataset["products"][0]["image_path"])

    def test_future_photo_capture_is_not_in_an_earlier_report_and_default_reader_stays_compatible(self):
        from test_product_images import sample_bytes
        from product_radar.product_images import VERSION, _candidate, inspect_image, current_public_image
        from product_radar.research_memory import key
        from product_radar.storage import atomic_bytes
        memory = source_fixture()
        observation = memory.state["observations"]["meta000"]
        observation["facts"]["image_urls"] = ["https://example.org/original-photo.png"]
        candidate = {"observation_id": "meta000", "url": observation["facts"]["image_urls"][0],
                     "observed_at": observation["observed_at"], "relationship": "relevant_family", "source": "meta"}
        binding = _candidate(memory.state, "c000", candidate)
        capture_id = key("product-image", [VERSION, binding])
        with tempfile.TemporaryDirectory() as directory:
            raw = sample_bytes()
            image_path = Path(directory) / "photo.png"
            atomic_bytes(image_path, raw)
            record = {"kind": "public_product_image", "version": VERSION, "state": "usable",
                "at": "2026-10-06T12:00:00+00:00", **binding,
                "image": {"path": str(image_path), **inspect_image(raw)}}
            memory.state["source_checks"][capture_id] = record
            self.assertIsNotNone(current_public_image(memory.state, "c000", [candidate], media_root=directory))
            self.assertIsNone(current_public_image(memory.state, "c000", [candidate], media_root=directory, as_of=NOW))
            with patch("product_radar.product_images.MEDIA_ROOT", Path(directory)):
                _, _, dataset = select(memory)
                self.assertIsNone(self.product(dataset)["image_path"])

    def test_unmarked_legacy_local_image_contract_is_unchanged(self):
        issue, dataset = legacy_fixture()
        dataset["products"][0]["image_path"] = str(LOGO)
        validate_dataset(issue, dataset)
        self.assertIn('alt="Synthetic product 1"', render_html(issue, dataset))
        self.assertTrue(all(field not in dataset["products"][0] for field in FIELDS))

    def supplier_leads(self, memory, count=4):
        original = deepcopy(memory.state["supplier_assessments"]["supplier-record0"])
        for n in range(1, count):
            record = deepcopy(original)
            lead_id = "additional-supplier-" + str(n)
            record["input"]["id"] = lead_id
            record["result"]["id"] = lead_id
            record["input"]["pages"][0]["url"] = "https://supplier-" + str(n) + ".example.org/item"
            record["input_hash"] = digest(record["input"])
            memory.state["supplier_assessments"]["additional-record-" + str(n)] = record

    def test_every_saved_supplier_and_every_route_reaches_both_languages_and_excel(self):
        memory = source_fixture()
        self.supplier_leads(memory)
        _, issue, dataset = select(memory)
        with tempfile.TemporaryDirectory() as directory:
            jobs = ResearchMemory(Path(directory) / "research", Journal(Path(directory) / "finance"), evidence_type="synthetic")
            written = generate_report_narratives(jobs, issue, dataset, WritingFixture())["dataset"]
        product = self.product(written)
        self.assertEqual(len(product["suppliers"]), 4)
        for language in ("en", "hi-Latn"):
            html = render_html(issue, written, language=language)
            section = html.split('data-concept-id="c000"', 1)[1].split('</section>', 1)[0]
            for supplier in product["suppliers"]:
                self.assertIn(supplier["name"], section)
            for route in product["source_route_checks"]:
                self.assertIn(route["source"].replace("&", "&amp;"), section)
                self.assertIn(route["state"], section)
                self.assertIn(route["detail"].replace("&", "&amp;"), section)
        raw = _workbook_bytes(issue, written)
        _workbook(raw, written)
        book = load_workbook(BytesIO(raw), read_only=True)
        try:
            rows = [row for row in book["Evidence"].iter_rows(min_row=7, values_only=True) if row[-1] == "c000"]
            self.assertTrue({row["source"] for row in product["source_route_checks"]} <= {row[2] for row in rows})
            self.assertEqual(sum(row[-1] == "c000" for row in book["Supplier leads"].iter_rows(min_row=7, values_only=True)), 4)
        finally:
            book.close()
        payload = make_narrative_input(written, "en")
        self.assertLess(len(canonical(payload)), 100 * 1024)
        self.assertEqual([fact["text"] for fact in next(item for item in payload["items"] if item["id"] == "c000")["facts"]
                         if fact["id"].startswith("market-") and fact["id"] != "market-label"],
                         [fact["source"] + ": " + fact["state"] + ". " + fact["detail"] for fact in product["market_checks"]])

    def test_new_workbook_cannot_silently_drop_route_or_supplier_rows_even_with_valid_product_ids(self):
        memory = source_fixture()
        self.supplier_leads(memory)
        _, issue, dataset = select(memory)
        raw = _workbook_bytes(issue, dataset)
        for sheet in ("Evidence", "Supplier leads"):
            with self.subTest(sheet=sheet):
                book = load_workbook(BytesIO(raw))
                book[sheet].delete_rows(7)
                changed = BytesIO()
                book.save(changed); book.close()
                with self.assertRaisesRegex(SafetyError, "DELIVERY_WORKBOOK_(?:LINK_)?MISMATCH"):
                    _workbook(changed.getvalue(), dataset)

    def test_marked_workbook_checks_each_real_hyperlink_not_just_open_source_text(self):
        _, issue, dataset = select(source_fixture())
        raw = _workbook_bytes(issue, dataset)
        for name, column in (("Evidence", 7), ("Supplier leads", 9)):
            for target in (None, "https://unrelated.example.org/not-this-product"):
                with self.subTest(name=name, target=target):
                    book = load_workbook(BytesIO(raw))
                    cell = next(row[column - 1] for row in book[name].iter_rows(min_row=7) if row[column - 1].hyperlink)
                    if target is None:
                        cell.hyperlink = None
                    else:
                        cell.hyperlink.target = target
                    changed = BytesIO()
                    book.save(changed); book.close()
                    with self.assertRaisesRegex(SafetyError, "DELIVERY_WORKBOOK_LINK_MISMATCH"):
                        _workbook(changed.getvalue(), dataset)

    def test_metric_unknown_and_bad_boolean_nan_ranges_are_validated(self):
        _, _, dataset = select(source_fixture())
        for value in (True, float("nan"), 1.5, [-1, 3], [3, 2], [None, None], "1e2", "1.0", "01", ".5",
                      "0.00", "1000000000.1", "0." + "1" * 31):
            with self.subTest(value=value):
                product = deepcopy(self.product(dataset))
                self.route(product, "meta-ads")["measurements"][0]["value"] = value
                with self.assertRaises(SafetyError):
                    validate_route_facts(product)

    def test_malformed_outcome_or_metric_name_has_a_fixed_validation_failure(self):
        _, _, dataset = select(source_fixture())
        for field, value in (("outcome", {}), ("metric_name", []), ("metric_name", {})):
            with self.subTest(field=field, value=value):
                product = deepcopy(self.product(dataset))
                row = self.route(product, "meta-ads")
                if field == "metric_name":
                    row["measurements"][0]["name"] = value
                else:
                    row[field] = value
                with self.assertRaises(SafetyError):
                    validate_route_facts(product)

    def test_current_stages_pin_new_render_versions_without_changing_legacy_contract(self):
        _, issue, dataset = select(source_fixture())
        with tempfile.TemporaryDirectory() as directory:
            jobs = ResearchMemory(Path(directory) / "research", Journal(Path(directory) / "finance"), evidence_type="synthetic")
            written = generate_report_narratives(jobs, issue, dataset, WritingFixture())["dataset"]
            workbook = write_issue_workbook(issue, written, directory)
            english = write_issue_edition(issue, written, directory, language="en")
            hinglish = write_issue_edition(issue, written, directory, language="hi-Latn")
            self.assertEqual(workbook["generation_version"], SOURCE_WORKBOOK_GENERATION_VERSION)
            self.assertEqual(english["render_version"], SOURCE_HTML_RENDER_VERSION)
            self.assertEqual(hinglish["render_version"], SOURCE_HTML_RENDER_VERSION)
            replay = write_issue_edition(issue, written, directory, language="en")
            self.assertEqual(replay, english)
            self.assertLess(len(canonical(written)), 4 * 1024 * 1024)
        old_issue, old_dataset = legacy_fixture()
        with tempfile.TemporaryDirectory() as directory:
            self.assertEqual(write_issue_workbook(old_issue, old_dataset, directory)["generation_version"], WORKBOOK_GENERATION_VERSION)
            self.assertEqual(write_issue_edition(old_issue, old_dataset, directory)["render_version"], HTML_RENDER_VERSION)


if __name__ == "__main__":
    unittest.main()
