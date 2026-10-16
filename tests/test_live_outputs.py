"""Synthetic packet checks; these are not real client research or publication."""

from copy import deepcopy
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
from product_radar.issue_allocation import allocate_issues, _fingerprint
from product_radar.live_outputs import (LEGACY_PRICE_DISPLAY_VERSION, PRICE_DISPLAY_VERSION, build_issue_dataset, price_display, render_html, supplier_display,
                                        supplier_terms_display, validate_dataset, write_issue_package, _supplier_minimum_order)
from product_radar.models import SafetyError, digest
from product_radar.service import CATEGORIES
from product_radar.storage import atomic_json
from test_fit_pipeline import older_format_scores


def fixture(count=10):
    pool = [{"concept_id": f"test{n:03d}", "category": CATEGORIES[n % 4], "identity_reviewed": True,
             "kpi_states": ["Meets", "Meets", "Not assessed"] + ["Meets"] * 5 + ["Unclear"] * 5,
             "evidence_refs": [f"testobs{n:03d}"]} for n in range(count)]
    issue = allocate_issues(pool, ["2026-10-13"], [], [])["issues"][0]
    facts = [{"concept_id": product["concept_id"], "title": "Synthetic product " + str(rank), "market_label": "Worth Watching",
              "market_checks": [{"source": "Meta", "state": "Positive", "detail": "Synthetic ad evidence for a fixture only.",
                                 "url": "https://example.org/ad", "observed_at": "2026-10-06T10:00:00+00:00"}],
              "prices": [{"kind": "retail", "amount": "149", "currency": "INR", "unit": "per_pack", "quantity": 2, "moq": None,
                          "conditions": "Delivery charge not stated.", "source_url": "https://example.org/retail", "observed_at": "2026-10-06T10:00:00+00:00"},
                         {"kind": "supplier", "amount": "45", "currency": "INR", "unit": "bulk_quote", "quantity": None, "moq": 100,
                          "conditions": "Bulk quote; not a verified one-piece dispatch offer.", "source_url": "https://example.org/supplier", "observed_at": "2026-10-06T10:00:00+00:00"}],
              "suppliers": [], "supplier_search_status": "complete", "source_refs": [{"id": f"testref{rank}", "url": "https://example.org/ad", "source": "Fixture source", "observed_at": "2026-10-06T10:00:00+00:00"}]}
             for rank, product in enumerate(issue["products"], 1)]
    narratives = {language: {"title": "Product picks" if language == "en" else "Product sujhaav", "products": [
        {"concept_id": product["concept_id"], "reason": "{{title}} is a fictional fixture, not a real recommendation." if language == "en" else "{{title}} sirf test ka example hai, asli recommendation nahi hai.",
         "points": ["The saved checklist has {{kpi_matches}} matches." if language == "en" else "Saved checklist mein {{kpi_matches}} points match hue hain."],
         "limitation": "Single-piece dispatch is not confirmed." if language == "en" else "Single-piece dispatch confirm nahi hai.", "update_sentence": ""}
        for product in issue["products"]]} for language in ("en", "hi-Latn")}
    dataset = build_issue_dataset(issue, facts, narratives=narratives, evidence_type="synthetic")
    return issue, dataset


@older_format_scores
class LiveOutputTests(unittest.TestCase):
    def photo_fixture(self):
        issue, dataset = fixture()
        path = ROOT / "media/NeuraFlow/long_maroon_path_transparent_1024.png"
        from product_radar.product_images import inspect_image
        photo = {"id": "synthetic-photo-record", "owner_id": issue["products"][0]["concept_id"],
                 "source_kind": "public_product_image", "relationship": "relevant_family",
                 "path": str(path), "sha256": inspect_image(path.read_bytes())["sha256"],
                 "source_url": "https://example.org/fixture-photo.png", "observed_at": "2026-10-06T10:00:00+00:00"}
        # The approved logo is only a real-byte synthetic image fixture. It is
        # not being attached as a product photograph to an actual report.
        issue["products"][0].update(image_provenance=photo,
            source_revisions={"source_checks": {photo["id"]: "a" * 64}})
        issue["selection_hash"] = digest({field: issue[field] for field in
            ("issue_id", "client_id", "revision", "date", "products", "evidence_fingerprint")})
        dataset.update(selection_hash=issue["selection_hash"])
        dataset["products"][0].update(image_provenance=deepcopy(photo), image_path=photo["path"])
        return issue, dataset, path

    def test_frozen_original_photo_is_verified_and_family_caption_is_explicit(self):
        issue, dataset, path = self.photo_fixture()
        with patch("product_radar.product_images.MEDIA_ROOT", path.parent):
            validate_dataset(issue, dataset)
            html = render_html(issue, dataset)
        self.assertTrue('<figcaption class="plate-note">Photo of a similar product type</figcaption>' in html)

    def test_changed_photo_binding_is_not_silently_accepted(self):
        for field, value in (("sha256", "f" * 64), ("owner_id", "different-product"),
                             ("relationship", "unrelated"), ("source_url", "https://example.org/different.png")):
            issue, dataset, path = self.photo_fixture()
            dataset["products"][0]["image_provenance"][field] = value
            with patch("product_radar.product_images.MEDIA_ROOT", path.parent), \
                    self.assertRaisesRegex(SafetyError, "LIVE_PHOTO_BINDING_INVALID"):
                validate_dataset(issue, dataset)

    def test_frozen_photo_with_wrong_saved_bytes_is_not_rendered_as_verified(self):
        issue, dataset, path = self.photo_fixture()
        issue["products"][0]["image_provenance"]["sha256"] = "f" * 64
        dataset["products"][0]["image_provenance"]["sha256"] = "f" * 64
        issue["selection_hash"] = digest({field: issue[field] for field in
            ("issue_id", "client_id", "revision", "date", "products", "evidence_fingerprint")})
        dataset["selection_hash"] = issue["selection_hash"]
        with patch("product_radar.product_images.MEDIA_ROOT", path.parent), \
                self.assertRaisesRegex(SafetyError, "LIVE_PHOTO_BYTES_CHANGED"):
            validate_dataset(issue, dataset)

    def supplier_fixture(self, *, purchase="No", dispatch="Yes"):
        return {"lead_id": "fixture-lead", "name": "Synthetic supplier", "url": "https://example.org/supplier",
                "single_piece_dispatch": "not_supported", "subscription": "not_stated", "min_order": None,
                "ecommerce_support": "Not stated", "reason": "A fixture requires bulk stock before sending individual parcels.",
                "terms": [{"field": "single_piece_purchase", "basis": "observed", "value": purchase},
                          {"field": "direct_customer_dispatch", "basis": "observed", "value": dispatch}]}

    def test_combined_supplier_label_does_not_deny_supported_individual_delivery(self):
        supplier = self.supplier_fixture()
        before = deepcopy(supplier)
        english = supplier_display(supplier)
        hinglish = supplier_display(supplier, "hi-Latn")
        self.assertIn("Single-piece sourcing and dispatch not supported", english)
        self.assertIn("Single-piece purchase: No", english)
        self.assertIn("Direct customer dispatch: Yes", english)
        self.assertIn("Single-piece kharid: Nahi", hinglish)
        self.assertIn("Seedha customer ko dispatch: Haan", hinglish)
        self.assertEqual(supplier, before)

    def test_individual_supplier_capabilities_require_observed_yes_or_no(self):
        for terms in (None, [], [{"field": "single_piece_purchase", "basis": "inferred", "value": "Yes"}],
                      [{"field": "single_piece_purchase", "basis": "conflicting", "value": "Conflicting"}],
                      [{"field": "single_piece_purchase", "basis": "observed", "value": "Yes"}] * 2):
            supplier = dict(self.supplier_fixture(), terms=terms)
            self.assertIn("Single-piece purchase: Unknown", supplier_terms_display(supplier))
            self.assertIn("Direct customer dispatch: Unknown", supplier_terms_display(supplier))
        supplier = self.supplier_fixture(purchase="Yes", dispatch="No")
        self.assertIn("Single-piece purchase: Yes", supplier_terms_display(supplier))
        self.assertIn("Direct customer dispatch: No", supplier_terms_display(supplier))

    def test_observed_compound_minimum_order_is_preserved_verbatim(self):
        for text in ("120 Pieces for the flip-lock offer; 96 Piece for the smaller container",
                     "100 Pieces for one offer; 160 Piece for another; 96 Piece for the third",
                     " 500 Roll for the printed shelf liner ", "96 sets of 3"):
            supplier = self.supplier_fixture()
            supplier["terms"].append({"field": "purchase_moq", "basis": "observed", "value": text})
            before = deepcopy(supplier)
            self.assertEqual(_supplier_minimum_order(supplier), text)
            self.assertIsNone(supplier["min_order"])
            self.assertEqual(supplier, before)

    def test_numeric_minimum_order_takes_precedence_over_raw_statement(self):
        supplier = self.supplier_fixture()
        supplier.update(min_order=100)
        supplier["terms"].append({"field": "purchase_moq", "basis": "observed", "value": "100 Pieces for the matched offer"})
        self.assertEqual(_supplier_minimum_order(supplier), 100)
        self.assertIs(type(_supplier_minimum_order(supplier)), int)

    def test_unknown_inferred_blank_duplicate_or_conflicting_minimum_is_not_stated(self):
        observed = {"field": "purchase_moq", "basis": "observed", "value": "120 Pieces"}
        cases = (None, [], [dict(observed, basis="inferred")], [observed, deepcopy(observed)],
                 [dict(observed, basis="conflicting")], [dict(observed, value="Conflicting")],
                 [dict(observed, value="Not stated")], [dict(observed, value="")],
                 [dict(observed, value=" \t")], [dict(observed, value=None)])
        for terms in cases:
            with self.subTest(terms=terms):
                self.assertEqual(_supplier_minimum_order(dict(self.supplier_fixture(), terms=terms)), "Not stated")

    def test_workbook_minimum_order_cell_preserves_observed_compounds_without_column_shift(self):
        issue, dataset = fixture()
        statements = ("120 Pieces for the flip-lock offer; 96 Piece for the smaller container",
                      "100 Pieces; 160 Piece; 96 Piece", "96 sets of 3")
        suppliers = []
        for index, text in enumerate(statements):
            supplier = self.supplier_fixture()
            supplier.update(lead_id="fixture-lead-" + str(index))
            supplier["terms"].append({"field": "purchase_moq", "basis": "observed", "value": text})
            suppliers.append(supplier)
        numeric = self.supplier_fixture()
        numeric.update(lead_id="numeric-lead", min_order=1)
        suppliers.append(numeric)
        dataset["products"][0]["suppliers"] = suppliers
        before = deepcopy(dataset)
        with tempfile.TemporaryDirectory() as folder:
            packet = write_issue_package(issue, dataset, folder)
            book = load_workbook(packet["artifacts"]["workbook"]["path"])
            sheet = book["Supplier leads"]
            self.assertEqual([sheet.cell(row, 6).value for row in range(7, 11)], [*statements, 1])
            self.assertEqual(sheet["F6"].value, "Minimum order")
            self.assertEqual(sheet["I6"].value, "Link")
            self.assertEqual(sheet["J6"].value, "Concept ID")
            self.assertTrue(all(sheet.cell(row, 9).hyperlink.target == "https://example.org/supplier" for row in range(7, 11)))
        self.assertEqual(dataset, before)

    def test_explicit_bulk_quantities_have_one_concise_product_warning_in_both_languages(self):
        issue, dataset = fixture()
        conditions = ["Minimum order quantity is 96 sets of 3; ₹80 per set.", "MOQ: 120 pieces.",
                      "Minimum order quantity: 500 Roll", "Min. order qty = 2 units."]
        sample = dataset["products"][0]["prices"][1]
        dataset["products"][0]["prices"] = [dict(sample, moq=None, conditions=text) for text in conditions]
        before = deepcopy(dataset)
        for language, phrase in (("en", "Some quoted offers have bulk minimum orders."),
                                 ("hi-Latn", "Kuch quotes mein bulk minimum order hai.")):
            html = render_html(issue, dataset, language=language)
            self.assertEqual(html.count('<p class="bulk">'), 1)
            self.assertEqual(html.count(phrase), 1)
            for condition in conditions:
                self.assertIn(condition, html)
        self.assertEqual(dataset, before)

    def test_each_supported_named_minimum_quantity_can_trigger_the_warning(self):
        issue, dataset = fixture()
        price = dataset["products"][0]["prices"][1]
        for conditions in ("MOQ 120 pieces", "Minimum order quantity is 96 sets of 3",
                           "Minimum order quantity: 500 rolls", "MOQ 2", "MOQ 1,000 units"):
            with self.subTest(conditions=conditions):
                price.update(conditions=conditions, moq=None)
                self.assertEqual(render_html(issue, dataset).count('<p class="bulk">'), 1)

    def test_pack_price_capacity_and_moq_one_do_not_create_bulk_minimums(self):
        issue, dataset = fixture()
        price = dataset["products"][0]["prices"][1]
        examples = ("₹130 per set of 6. No MOQ stated for this listing.",
                    "Set of 3; capacities 600ML + 1200ML + 1800ML; ₹80 per set.",
                    "MOQ 1 piece; pack contains 6 items; ₹130 each.",
                    "500 ml bottle, 1200 ml capacity; minimum order quantity not stated.",
                    "Minimum order price: ₹130; six-piece set.", "MOQ: ₹130 per set of 6",
                    "MOQ 500 ml", "Minimum order quantity: 2 kg", "Minimum order: 130 INR",
                    "MOQ 130 rupees", "Minimum order: 130 ₹", None)
        for conditions in examples:
            with self.subTest(conditions=conditions):
                price.update(conditions=conditions, moq=None)
                self.assertFalse('<p class="bulk">' in render_html(issue, dataset))

    def test_bulk_warning_does_not_apply_supplier_minimum_to_retail_offer(self):
        issue, dataset = fixture()
        dataset["products"][0]["prices"] = [dict(dataset["products"][0]["prices"][0], conditions="MOQ 500 pieces")]
        self.assertFalse('<p class="bulk">' in render_html(issue, dataset))

    def test_running_footer_is_a_page_box_and_every_product_keeps_its_sources(self):
        issue, dataset = fixture()
        html = render_html(issue, dataset)
        self.assertTrue('@bottom-left{content:"NeuraFlow · Product Opportunity Brief · 13 October 2026"' in html)
        self.assertTrue('content:"Page " counter(page) " of " counter(pages)' in html)
        self.assertEqual(html.count('<section class="src-sec">'), 10)
        self.assertEqual(html.count('<figure class="plate plate--empty plate--big">'), 10)
        self.assertTrue('.plate--big{height:50mm}' in html)
        self.assertTrue('@page{size:A4 portrait;margin:15mm 16mm 17mm 16mm;background:#F9F3EC;' in html)

    def test_missing_image_heading_is_full_width_with_small_explicit_note(self):
        issue, dataset = fixture()
        title = "Synthetic long product heading that should use the full page width without a blank picture column"
        dataset["products"][0]["title"] = title
        before = deepcopy(dataset)
        for language, note in (("en", "No product photo in this issue"),
                               ("hi-Latn", "Is issue mein product ki tasveer nahi hai")):
            html = render_html(issue, dataset, language=language)
            self.assertEqual(html.count('<figure class="plate plate--empty plate--big">'), 10)
            self.assertEqual(html.count(f'<b>{note}</b>'), 16)  # ten product pages and six cover tiles
            self.assertTrue('<h1 class="h-title">' + title + '</h1>' in html)
            self.assertTrue('.h-title{font-family:var(--head);font-weight:600;font-size:var(--fs-title);' in html)
            self.assertTrue('--fs-title:18pt' in html)
        self.assertEqual(dataset, before)

    def test_available_fixture_image_keeps_original_photographed_hero_dimensions(self):
        issue, dataset = fixture()
        # Existing branding image is used only as a valid local raster fixture;
        # it is not presented as a real product photograph or a client packet.
        dataset["products"][0]["image_path"] = str(ROOT / "media" / "NeuraFlow" / "long_maroon_path_transparent_1024.png")
        html = render_html(issue, dataset)
        self.assertEqual(html.count('<figure class="plate plate--big"><img src="data:image/png;base64,'), 1)
        self.assertEqual(html.count('<figure class="plate plate--empty plate--big">'), 9)
        self.assertTrue('.pa-head{display:grid;grid-template-columns:50mm 1fr 30mm;' in html)
        self.assertTrue('.plate img{max-width:100%;max-height:100%;object-fit:contain;' in html)
        self.assertEqual(html.count('alt="Synthetic product 1"'), 2)  # the product page and its cover tile

    def test_full_width_heading_does_not_omit_prices_conditions_sources_or_dates(self):
        issue, dataset = fixture()
        product = dataset["products"][0]
        product["prices"][1]["conditions"] = "Minimum order quantity is 96 sets of 3. All delivery and packaging conditions require confirmation."
        product["source_refs"].extend([
            {"id": "extra-a", "url": "https://example.org/extra", "source": "Extra source", "observed_at": "2026-10-07T10:00:00+00:00"},
            {"id": "extra-b", "url": "https://example.org/extra", "source": "Extra source", "observed_at": "2026-10-07T10:00:00+00:00"}])
        before = deepcopy(dataset)
        for language in ("en", "hi-Latn"):
            html = render_html(issue, dataset, language=language)
            self.assertTrue(product["prices"][1]["conditions"] in html)
            self.assertTrue("₹149" in html)
            self.assertTrue("₹45" in html)
            self.assertTrue("13 October 2026" in html)
            # Two references to the same page seen on the same day are one dated link.
            self.assertEqual(html.count('href="https://example.org/extra"'), 1)
            self.assertEqual(html.count('Extra source</a> <span>7 Oct</span>'), 1)
            self.assertEqual(html.count('<p class="bulk">'), 1)
        self.assertEqual(dataset, before)

    def test_print_uses_physical_page_margins_and_a_full_bleed_cover(self):
        issue, dataset = fixture()
        html = render_html(issue, dataset)
        self.assertTrue('@page{size:A4 portrait;margin:15mm 16mm 17mm 16mm;background:#F9F3EC;' in html)
        self.assertTrue('@page cover{margin:0;' in html)
        self.assertTrue('.cover{page:cover;height:297mm;overflow:hidden}' in html)
        # Screen A4 previews keep their own padding; print relies on the page margins.
        self.assertTrue('.sheet{position:relative;width:210mm;min-height:297mm;margin:0 auto 8mm;padding:15mm 16mm 17mm;' in html)
        self.assertTrue('--fs-label:9pt' in html)

    def test_print_keeps_headings_with_their_content_and_never_splits_rows(self):
        issue, dataset = fixture()
        html = render_html(issue, dataset)
        self.assertTrue('.sheet{break-after:page;page-break-after:always}' in html)
        self.assertTrue('margin-bottom:2mm;break-after:avoid;page-break-after:avoid}' in html)
        for rule in ('.price-list li{display:grid;', '.sup tr,.sup tbody{break-inside:avoid;page-break-inside:avoid}',
                     '.league tr{break-inside:avoid;page-break-inside:avoid}', '.fit-list li{break-inside:avoid;',
                     '.tile{background:var(--card);', '.src-sec{margin-top:3.2mm;break-inside:avoid;'):
            self.assertTrue(rule in html, rule)
        self.assertFalse("overflow:hidden}" in html.split(".cover{page:cover")[1].split("}")[1])

    def test_long_print_flow_preserves_all_offers_suppliers_and_sources_in_order(self):
        issue, dataset = fixture()
        product = dataset["products"][0]
        price = product["prices"][1]
        product["prices"] = [dict(price, conditions="MOQ 120 pieces; " + "Long exact supplier condition. " * 120 + "END-QUOTE-ONE"),
                             dict(price, conditions="MOQ 96 sets of 3; " + "Another exact offer condition. " * 120 + "END-QUOTE-TWO")]
        product["suppliers"] = [self.supplier_fixture()]
        product["suppliers"][0]["reason"] = "Long exact supplier explanation. " * 120 + "END-SUPPLIER-EXPLANATION"
        product["source_refs"].append({"id": "long-print-source", "source": "Last exact source",
            "url": "https://example.org/final-source", "observed_at": "2026-10-07T10:00:00+00:00"})
        before = deepcopy(dataset)
        for language in ("en", "hi-Latn"):
            html = render_html(issue, dataset, language=language)
            for offer in product["prices"]:
                self.assertTrue(offer["conditions"] in html)
            self.assertTrue(product["suppliers"][0]["reason"] in html)
            self.assertTrue('href="https://example.org/final-source"' in html)
            self.assertTrue('Last exact source</a> <span>7 Oct</span>' in html)
            self.assertLess(html.index("END-QUOTE-ONE"), html.index("END-QUOTE-TWO"))
            self.assertLess(html.index("END-QUOTE-TWO"), html.index("END-SUPPLIER-EXPLANATION"))
            self.assertLess(html.index("END-SUPPLIER-EXPLANATION"), html.index("Last exact source"))
        self.assertEqual(dataset, before)

    def test_long_conditions_are_preserved_in_html_and_expand_workbook_rows(self):
        issue, dataset = fixture()
        condition = "Supplier quotation terms remain unchanged and require confirmation. " * 20 + "END-OF-QUOTE"
        dataset["products"][0]["prices"][1]["conditions"] = condition
        supplier = self.supplier_fixture()
        supplier["reason"] = "Original supplier evidence and conditions. " * 30 + "END-OF-SUPPLIER-NOTE"
        dataset["products"][0]["suppliers"] = [supplier]
        html = render_html(issue, dataset)
        self.assertIn(condition, html)
        self.assertIn(supplier["reason"], html)
        with tempfile.TemporaryDirectory() as folder:
            packet = write_issue_package(issue, dataset, folder)
            book = load_workbook(packet["artifacts"]["workbook"]["path"])
            self.assertIn(condition, book["Shortlist"]["H7"].value)
            self.assertIn(supplier["reason"], book["Supplier leads"]["H7"].value)
            self.assertGreater(book["Shortlist"].row_dimensions[7].height, 52)
            self.assertGreater(book["Shortlist"].row_dimensions[7].height, book["Shortlist"].row_dimensions[8].height)
            self.assertGreater(book["Supplier leads"].row_dimensions[7].height, 52)
            self.assertGreater(book["Shortlist"].column_dimensions["H"].width, 32)
            self.assertGreater(book["Supplier leads"].column_dimensions["H"].width, 32)
            for sheet in book:
                self.assertTrue(all(row.height <= 409 for row in sheet.row_dimensions.values() if row.height))

    def test_workbook_height_ceiling_keeps_all_long_text_and_column_positions(self):
        issue, dataset = fixture()
        condition = "Full supplier offer condition. " * 600 + "FINAL-FACT-NOT-OMITTED"
        dataset["products"][0]["prices"][1]["conditions"] = condition
        dataset["products"][0]["suppliers"] = [self.supplier_fixture()]
        with tempfile.TemporaryDirectory() as folder:
            packet = write_issue_package(issue, dataset, folder)
            book = load_workbook(packet["artifacts"]["workbook"]["path"])
            self.assertEqual(book["Shortlist"].row_dimensions[7].height, 409)
            self.assertIn(condition, book["Shortlist"]["H7"].value)
            sheet = book["Supplier leads"]
            self.assertEqual(sheet.max_column, 10)
            self.assertEqual(sheet["D6"].value, "Single-piece sourcing and dispatch")
            self.assertEqual(sheet["I6"].value, "Link")
            self.assertEqual(sheet["J6"].value, "Concept ID")
            self.assertEqual(sheet["I7"].hyperlink.target, "https://example.org/supplier")
            self.assertEqual(sheet["J7"].value, issue["products"][0]["concept_id"])
            self.assertIn("Single-piece purchase: No", sheet["H7"].value)
            self.assertIn("Direct customer dispatch: Yes", sheet["H7"].value)
            self.assertEqual(sheet.freeze_panes, "C7")
            self.assertEqual(sheet.auto_filter.ref, "A6:J7")

    def reviewed_update_fixture(self):
        issue, source = fixture()
        # Manufactured frozen selection representing a later useful-update
        # issue; this does not create real publication history or evidence.
        issue["products"][0].update(status="Updated", material_update_reviewed=True,
                                    update_reason="A reviewed fixture adds a supported supplier option.")
        issue["counts"] = {"products": 10, "new": 9, "updates": 1}
        issue["evidence_fingerprint"] = _fingerprint(issue["products"])
        issue["selection_hash"] = digest({field: issue[field] for field in
                                          ("issue_id", "client_id", "revision", "date", "products", "evidence_fingerprint")})
        source["selection_hash"] = issue["selection_hash"]
        return issue, source

    def test_reviewed_update_fields_come_only_from_frozen_selection_and_survive_saved_facts(self):
        issue, source = self.reviewed_update_fixture()
        source["products"][0].update(material_update_reviewed=False, update_reason="Untrusted supplied facts.")
        for narrative in source["narratives"].values():
            narrative["products"][0].update(material_update_reviewed=False, update_reason="Untrusted prose metadata.")
        dataset = build_issue_dataset(issue, source["products"], narratives=source["narratives"], evidence_type="synthetic")
        product = dataset["products"][0]
        self.assertIs(product["material_update_reviewed"], True)
        self.assertEqual(product["update_reason"], issue["products"][0]["update_reason"])
        self.assertEqual(product["status"], "Updated")
        with tempfile.TemporaryDirectory() as root:
            packet = write_issue_package(issue, dataset, root)
            from product_radar.storage import load_json
            saved = load_json(packet["artifacts"]["facts"]["path"], 4 * 1024 * 1024)
            self.assertIs(saved["products"][0]["material_update_reviewed"], True)
            self.assertEqual(saved["products"][0]["update_reason"], issue["products"][0]["update_reason"])

    def test_updated_admission_cannot_be_changed_or_removed_after_dataset_build(self):
        issue, source = self.reviewed_update_fixture()
        dataset = build_issue_dataset(issue, source["products"], narratives=source["narratives"], evidence_type="synthetic")
        edits = [lambda p: p.update(material_update_reviewed=False), lambda p: p.update(material_update_reviewed=1),
                 lambda p: p.pop("material_update_reviewed"), lambda p: p.update(update_reason="Different unsupported reason."),
                 lambda p: p.pop("update_reason")]
        for edit in edits:
            changed = deepcopy(dataset); edit(changed["products"][0])
            with self.assertRaisesRegex(SafetyError, "LIVE_PROTECTED_FACT_CHANGED"):
                validate_dataset(issue, changed)

    def test_prose_and_supplied_facts_cannot_admit_an_unreviewed_updated_product(self):
        issue, source = self.reviewed_update_fixture()
        issue["products"][0]["material_update_reviewed"] = False
        issue["selection_hash"] = digest({field: issue[field] for field in
                                          ("issue_id", "client_id", "revision", "date", "products", "evidence_fingerprint")})
        source["products"][0].update(material_update_reviewed=True, update_reason="The model says it is reviewed.")
        for narrative in source["narratives"].values():
            narrative["products"][0].update(material_update_reviewed=True, update_reason="The model says it is reviewed.")
        with self.assertRaisesRegex(SafetyError, "ISSUE_UPDATE_EVIDENCE_REQUIRED"):
            build_issue_dataset(issue, source["products"], narratives=source["narratives"], evidence_type="synthetic")

    def test_new_products_do_not_inherit_accidental_update_metadata_from_facts_or_prose(self):
        issue, source = fixture()
        for product in source["products"]:
            product.update(material_update_reviewed=True, update_reason="Injected update claim.")
        for narrative in source["narratives"].values():
            for product in narrative["products"]:
                product.update(material_update_reviewed=True, update_reason="Injected prose update claim.")
        dataset = build_issue_dataset(issue, source["products"], narratives=source["narratives"], evidence_type="synthetic")
        for product in dataset["products"]:
            self.assertEqual(product["status"], "New")
            self.assertNotIn("material_update_reviewed", product)
            self.assertNotIn("update_reason", product)
        changed = deepcopy(dataset)
        changed["products"][0].update(material_update_reviewed=True, update_reason="Post-build claim.")
        with self.assertRaisesRegex(SafetyError, "LIVE_PROTECTED_FACT_CHANGED"):
            validate_dataset(issue, changed)

    def test_new_product_cannot_render_a_model_supplied_update_sentence_in_either_language(self):
        issue, dataset = fixture()
        for language in ("en", "hi-Latn"):
            changed = deepcopy(dataset)
            changed["narratives"][language]["products"][0]["update_sentence"] = "An unreviewed fixture claims a useful update."
            with self.assertRaisesRegex(SafetyError, "LIVE_NARRATIVE_UPDATE_NOT_ADMITTED"):
                render_html(issue, changed, language=language)

    def test_reviewed_update_can_render_its_supplied_update_sentence_without_changing_admission(self):
        issue, source = self.reviewed_update_fixture()
        for narrative in source["narratives"].values():
            narrative["products"][0]["update_sentence"] = "The reviewed fixture adds a supplier option."
        dataset = build_issue_dataset(issue, source["products"], narratives=source["narratives"], evidence_type="synthetic")
        for language in ("en", "hi-Latn"):
            self.assertIn("The reviewed fixture adds a supplier option.", render_html(issue, dataset, language=language))
        self.assertEqual(dataset["products"][0]["update_reason"], issue["products"][0]["update_reason"])

    def test_10_12_15_have_matching_fixed_order_and_correct_counts(self):
        for count in (10, 12, 15):
            issue, dataset = fixture(count)
            self.assertEqual(len(dataset["products"]), count)
            self.assertEqual([p["rank"] for p in dataset["products"]], list(range(1, count + 1)))
            self.assertEqual([p["concept_id"] for p in dataset["products"]], [p["concept_id"] for p in issue["products"]])

    def test_adapter_uses_issue_order_not_supplier_input_order(self):
        issue, dataset = fixture()
        reordered = build_issue_dataset(issue, list(reversed(dataset["products"])), narratives=dataset["narratives"], evidence_type="synthetic")
        self.assertEqual(reordered, dataset)

    def test_extra_missing_or_duplicate_fact_record_is_rejected(self):
        issue, dataset = fixture()
        for facts in (dataset["products"][:-1], dataset["products"] + [dataset["products"][0]]):
            with self.assertRaises(SafetyError):
                build_issue_dataset(issue, facts)

    def test_different_counts_ids_checklist_or_order_cannot_leak_into_report(self):
        issue, dataset = fixture()
        changes = [lambda d: d["products"].reverse(), lambda d: d["products"][0].update(kpi_matches=13),
                   lambda d: d["products"][0].update(category="Pet"), lambda d: d.update(selection_hash="a" * 64)]
        for mutation in changes:
            altered = deepcopy(dataset)
            mutation(altered)
            with self.assertRaises(SafetyError):
                validate_dataset(issue, altered)

    def test_frozen_rank_and_checklist_count_are_integers_not_equal_bool_or_float_values(self):
        issue, dataset = fixture()
        for field, value in (("rank", True), ("rank", 1.0), ("kpi_matches", 7.0)):
            changed = deepcopy(dataset)
            changed["products"][0][field] = value
            with self.assertRaisesRegex(SafetyError, "LIVE_PROTECTED_FACT_CHANGED"):
                validate_dataset(issue, changed)

    def test_current_issue_and_dataset_header_criterion_ids_are_exact_integers(self):
        from test_fit_pipeline import current_fixture, selection
        _, issue, dataset = selection(current_fixture())
        for header in ("issue", "dataset"):
            for number in (True, 1.0):
                changed_issue, changed_dataset = deepcopy(issue), deepcopy(dataset)
                target = changed_issue if header == "issue" else changed_dataset
                target["criterion_ids"][0] = number
                with self.subTest(header=header, number=repr(number)), self.assertRaisesRegex(SafetyError, "^LIVE_FIT_POLICY_MISMATCH$"):
                    validate_dataset(changed_issue, changed_dataset)
        bad_issue = deepcopy(issue); bad_issue["kpi_total"] = 10.0
        with self.assertRaisesRegex(SafetyError, "^LIVE_FIT_POLICY_MISMATCH$"):
            validate_dataset(bad_issue, dataset)
        with self.assertRaisesRegex(SafetyError, "^ISSUE_FIT_POLICY_MISMATCH$"):
            build_issue_dataset(bad_issue, dataset["products"])

    def test_issue_and_dataset_revision_headers_reject_equal_bool_or_float(self):
        issue, dataset = fixture()
        for header in ("issue", "dataset"):
            for number in (True, 1.0):
                changed_issue, changed_dataset = deepcopy(issue), deepcopy(dataset)
                target = changed_issue if header == "issue" else changed_dataset
                target["revision"] = number
                with self.subTest(header=header, number=repr(number)), self.assertRaisesRegex(SafetyError, "^LIVE_DATASET_ISSUE_MISMATCH$"):
                    validate_dataset(changed_issue, changed_dataset)

    def test_current_checklist_phrase_ids_reject_equal_bool_or_float(self):
        from test_fit_pipeline import current_fixture, selection
        _, issue, dataset = selection(current_fixture())
        dataset["narratives"] = {"en": {"title": "Fixture research", "products": [
            {"concept_id": product["concept_id"], "reason": "Fixture product research.", "points": [],
             "limitation": "Supplier terms need checking.", "update_sentence": "",
             "checklist_phrases": [{"id": row["id"], "text": "Fixture judgment"} for row in product["criteria"]]}
            for product in dataset["products"]]}}
        validate_dataset(issue, dataset)
        for number in (True, 1.0):
            changed = deepcopy(dataset)
            changed["narratives"]["en"]["products"][0]["checklist_phrases"][0]["id"] = number
            with self.subTest(number=repr(number)), self.assertRaisesRegex(SafetyError, "^LIVE_CHECKLIST_PHRASES_INVALID$"):
                validate_dataset(issue, changed)

    def test_short_issue_is_not_rendered_as_a_complete_report(self):
        issue = allocate_issues([{"concept_id": f"short{n}", "category": CATEGORIES[0], "identity_reviewed": True,
                                 "kpi_states": ["Meets", "Meets", "Not assessed"] + ["Meets"] * 5 + ["Unclear"] * 5} for n in range(5)], ["2026-10-13"], [], [])["issues"][0]
        with self.assertRaises(SafetyError):
            build_issue_dataset(issue, [{"concept_id": p["concept_id"], "title": "Short fixture", "market_label": "Worth Watching"} for p in issue["products"]])

    def test_retail_and_supplier_prices_keep_unit_pack_and_bulk_conditions(self):
        _, dataset = fixture()
        retail, supplier = dataset["products"][0]["prices"]
        self.assertIn("Retail listing: ₹149", price_display(retail))
        self.assertIn("pack quantity 2", price_display(retail))
        self.assertIn("Supplier quote: ₹45", price_display(supplier))
        self.assertIn("minimum order 100", price_display(supplier))
        self.assertIn("not a verified one-piece", price_display(supplier))

    def test_unknown_price_is_not_zero_and_unbounded_decimal_is_rejected(self):
        _, dataset = fixture()
        price = dataset["products"][0]["prices"][0]
        price["amount"] = None
        self.assertIn("Price not stated", price_display(price))
        self.assertNotIn("₹0", price_display(price))
        for amount in ("NaN", "Infinity", "-1", "1E9999999", 149.0, True):
            with self.assertRaises(SafetyError):
                price_display(dict(price, amount=amount))

    def test_raw_retail_quote_is_observed_without_inventing_normalized_amount(self):
        _, dataset = fixture()
        price = dict(dataset["products"][0]["prices"][0], amount=None, raw_displayed="₹675.00",
                     unit="not_stated", quantity=None,
                     conditions="Pack quantity, tax and shipping not verified. Relevant family example, not an exact product price.")
        before = deepcopy(price)
        for language, prefix, caveat, missing in (
                ("en", "Retail listing: displayed quote ₹675.00", "price basis not verified", "Price not stated"),
                ("hi-Latn", "Retail listing ki keemat: dikhaya gaya quote ₹675.00",
                 "keemat ka basis abhi verify nahi hua hai", "keemat nahi di gayi")):
            with self.subTest(language=language):
                displayed = price_display(price, language)
                self.assertIn(prefix, displayed)
                self.assertIn(caveat, displayed)
                self.assertIn(price["conditions"], displayed)
                self.assertNotIn(missing, displayed)
        self.assertEqual(price, before)
        self.assertIsNone(price["amount"])

    def test_missing_or_blank_raw_quote_still_has_no_stated_price(self):
        _, dataset = fixture()
        price = dict(dataset["products"][0]["prices"][0], amount=None)
        for raw in (None, "", " \t\n"):
            for language, missing in (("en", "Price not stated"), ("hi-Latn", "keemat nahi di gayi")):
                with self.subTest(raw=raw, language=language):
                    displayed = price_display(dict(price, raw_displayed=raw), language)
                    self.assertIn(missing, displayed)
                    self.assertNotIn("displayed quote", displayed)
                    self.assertNotIn("dikhaya gaya quote", displayed)

    def test_normalized_price_and_supplier_display_are_unchanged_by_raw_field(self):
        _, dataset = fixture()
        retail, supplier = dataset["products"][0]["prices"]
        for language in ("en", "hi-Latn"):
            self.assertEqual(price_display(dict(retail, raw_displayed="₹149–199"), language), price_display(retail, language))
            self.assertIn("₹149.00", price_display(dict(retail, amount="149.00", raw_displayed="₹149–199"), language))
            self.assertEqual(price_display(dict(supplier, raw_displayed="₹45"), language), price_display(supplier, language))
            unknown_supplier = dict(supplier, amount=None)
            self.assertEqual(price_display(dict(unknown_supplier, raw_displayed="₹45"), language), price_display(unknown_supplier, language))

    def test_display_contract_versions_are_exact_named_strings_and_legacy_is_deterministic(self):
        _, dataset = fixture()
        price = dict(dataset["products"][0]["prices"][0], amount=None, raw_displayed="₹149–₹189")
        self.assertIn("displayed quote ₹149–₹189", price_display(price, render_version=PRICE_DISPLAY_VERSION))
        self.assertIn("Retail listing: Price not stated", price_display(price, render_version=LEGACY_PRICE_DISPLAY_VERSION))
        self.assertNotIn("displayed quote", price_display(price, render_version=LEGACY_PRICE_DISPLAY_VERSION))
        for value in (False, True, 1.0, 2.0, "1", "2", None, 3):
            with self.subTest(value=value), self.assertRaisesRegex(SafetyError, "PRICE_DISPLAY_VERSION_INVALID"):
                price_display(price, render_version=value)

    def test_saved_manifest_unknown_version_cannot_replay_or_add_a_language(self):
        issue, dataset = fixture()
        with tempfile.TemporaryDirectory() as root:
            draft = write_issue_package(issue, dataset, root)
            path = Path(root) / "data/outputs" / issue["issue_id"] / "revision-1/manifest-en.json"
            for value in (False, True, 1.0, 2.0, "2", None, 3):
                with self.subTest(value=value):
                    atomic_json(path, dict(draft, price_display_version=value))
                    with self.assertRaisesRegex(SafetyError, "PRICE_DISPLAY_VERSION_INVALID"):
                        write_issue_package(issue, dataset, root)
            atomic_json(path, {k: v for k, v in draft.items() if k != "price_display_version"})
            with self.assertRaisesRegex(SafetyError, "RENDER_VERSION_CHANGED"):
                write_issue_package(issue, dataset, root, language="hi-Latn")

    def test_raw_quote_reaches_both_html_languages_and_actual_shared_workbook(self):
        issue, dataset = fixture()
        price = dataset["products"][0]["prices"][0]
        price.update(amount=None, raw_displayed="₹149–₹189", unit="not_stated", quantity=None,
                     conditions="Pack quantity, tax and shipping not verified. Relevant family example, not an exact product price.")
        dataset["narratives"]["en"]["products"][0]["reason"] = "The observed listing is {{retail_price}}."
        dataset["narratives"]["hi-Latn"]["products"][0]["reason"] = "Listing mein {{retail_price}} dikhaya gaya hai."
        before = deepcopy(dataset)
        with tempfile.TemporaryDirectory() as folder:
            packets = [write_issue_package(issue, dataset, folder, language=language) for language in ("en", "hi-Latn")]
            for language, packet in zip(("en", "hi-Latn"), packets):
                html = Path(packet["artifacts"]["html"]["path"]).read_text()
                # The figure leads each price row; the rest of the same exact wording follows it.
                figure, rest = price_display(price, language).partition("; ")[::2]
                self.assertTrue(figure.split(": ", 1)[1] in html)
                self.assertTrue(rest in html)
                self.assertTrue(price["conditions"] in html)
                self.assertNotIn("Retail listing: Price not stated", html)
                self.assertNotIn("Retail listing ki keemat: keemat nahi di gayi", html)
            book = load_workbook(packets[0]["artifacts"]["workbook"]["path"])
            self.assertEqual(book["Shortlist"]["G7"].value, price_display(price))
            self.assertEqual(book["Shortlist"]["H7"].value, price_display(dataset["products"][0]["prices"][1]))
            self.assertEqual(packets[0]["artifacts"]["workbook"], packets[1]["artifacts"]["workbook"])
        self.assertEqual(dataset, before)

    def test_no_suppliers_or_images_have_factual_fallbacks_not_invented_leads(self):
        issue, dataset = fixture()
        html = render_html(issue, dataset)
        self.assertIn("No supplier lead found in the saved research.", html)
        self.assertIn("No product photo in this issue", html)
        text = html.split("</style>", 1)[1]  # the embedded font data is not report text
        self.assertNotIn("DEMO-", text)
        self.assertNotIn("Pet", text)
        dataset["products"][0]["supplier_search_status"] = "not_started"
        self.assertIn("Supplier search has not been completed.", render_html(issue, dataset))

    def test_cover_disclaimer_is_once_not_every_page(self):
        issue, dataset = fixture()
        html = render_html(issue, dataset)
        self.assertEqual(html.count("This issue uses manufactured test data."), 1)

    def test_two_independent_languages_keep_identical_facts_and_negation(self):
        issue, dataset = fixture()
        english = render_html(issue, dataset, language="en")
        hinglish = render_html(issue, dataset, language="hi-Latn")
        self.assertTrue("Single-piece dispatch is not confirmed." in english)
        self.assertTrue("Single-piece dispatch confirm nahi hai." in hinglish)
        for html in (english, hinglish):
            self.assertTrue('<b class="price">₹149</b>' in html)
            self.assertTrue('<b class="price">₹45</b>' in html)
            self.assertEqual(html.count('<td class="c5">7<span class="of">/13</span></td>'), 10)
        self.assertTrue('<p class="sub">Retail listing</p>' in hinglish)
        self.assertTrue('<p class="sub">Supplier aur dropship quote</p>' in hinglish)
        self.assertTrue("minimum order: 100" in hinglish)
        self.assertTrue("Achha sanket mila hai" in hinglish)
        self.assertTrue("Indian dropshipping ke liye 10 product types: is issue mein 10 naye, 0 update." in hinglish)
        self.assertTrue("13 October 2026" in hinglish)
        for product in dataset["products"]:
            self.assertIn(product["concept_id"], english)
            self.assertIn(product["concept_id"], hinglish)
            self.assertIn("#product-" + product["concept_id"], english)
        del dataset["narratives"]["hi-Latn"]
        with self.assertRaisesRegex(SafetyError, "LANGUAGE_MISSING"):
            render_html(issue, dataset, language="hi-Latn")

    def test_narrative_cannot_change_numbers_or_insert_unknown_placeholders(self):
        issue, dataset = fixture()
        for sentence in ("This will make 100% profit.", "The price is {{invented_price}}."):
            changed = deepcopy(dataset)
            changed["narratives"]["en"]["products"][0]["reason"] = sentence
            with self.assertRaisesRegex(SafetyError, "UNPROTECTED_NUMBER"):
                render_html(issue, changed)

    def test_html_text_and_urls_are_not_executable(self):
        issue, dataset = fixture()
        dataset["products"][0]["title"] = "<script>alert('private')</script>"
        html = render_html(issue, dataset)
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)
        for url in ("javascript:alert(1)", "https://person:secret@example.org", "https://example.org/?access_token=private"):
            changed = deepcopy(dataset)
            changed["products"][0]["source_refs"][0]["url"] = url
            with self.assertRaisesRegex(SafetyError, "URL_INVALID"):
                render_html(issue, changed)

    def test_report_links_use_public_reader_secret_and_control_rules(self):
        issue, dataset = fixture()
        invalid = ("https://example.org/?client_secret=fixture", "https://example.org/?SESSION_TOKEN=fixture",
                   "https://example.org/?api_token=fixture", "https://example.org/?api_key=",
                   "https://example.org/?API%5FKEY=fixture", "https://example.org/search%0Atext",
                   "https://example.org/with space", "https://example.org/\\redirect",
                   "https://@example.org/item", "https://example.org:99999/item",
                   "https://example.org/" + "x" * 4096)
        for field, url_field in (("source_refs", "url"), ("market_checks", "url"),
                                 ("prices", "source_url"), ("suppliers", "url")):
            for url in invalid:
                changed = deepcopy(dataset)
                if field == "suppliers":
                    changed["products"][0][field] = [self.supplier_fixture()]
                changed["products"][0][field][0][url_field] = url
                with self.subTest(field=field, url=url), self.assertRaisesRegex(SafetyError, "^LIVE_SOURCE_URL_INVALID$"):
                    render_html(issue, changed)

    def test_public_report_link_is_validated_without_changing_its_fragment(self):
        issue, dataset = fixture()
        link = "https://example.org:443/search?q=organizer&sort=recent#details"
        dataset["products"][0]["source_refs"][0]["url"] = link
        before = deepcopy(dataset)
        html = render_html(issue, dataset)
        self.assertIn('href="https://example.org:443/search?q=organizer&amp;sort=recent#details"', html)
        self.assertEqual(dataset, before)

    def test_immutable_packet_replay_reuses_bytes_and_both_languages_share_workbook(self):
        issue, dataset = fixture()
        with tempfile.TemporaryDirectory() as folder:
            first = write_issue_package(issue, dataset, folder)
            workbook = Path(first["artifacts"]["workbook"]["path"]).read_bytes()
            self.assertEqual(first, write_issue_package(issue, dataset, folder))
            second = write_issue_package(issue, dataset, folder, language="hi-Latn")
            self.assertEqual(first["facts_hash"], second["facts_hash"])
            self.assertEqual(first["artifacts"]["workbook"], second["artifacts"]["workbook"])
            self.assertEqual(Path(second["artifacts"]["workbook"]["path"]).read_bytes(), workbook)
            self.assertIn("/media/live/reports/", first["artifacts"]["html"]["path"])
            self.assertIn("/data/outputs/", first["artifacts"]["workbook"]["path"])
            self.assertFalse(first["approved"])
            self.assertFalse(first["release_ready"])

    def test_changed_facts_or_corrupted_package_requires_a_new_revision_or_recovery(self):
        issue, dataset = fixture()
        with tempfile.TemporaryDirectory() as folder:
            packet = write_issue_package(issue, dataset, folder)
            changed = deepcopy(dataset)
            changed["products"][0]["prices"][0]["amount"] = "199"
            with self.assertRaisesRegex(SafetyError, "IMMUTABLE_FACTS_CHANGED"):
                write_issue_package(issue, changed, folder)
            html_path = Path(packet["artifacts"]["html"]["path"])
            with patch.object(Path, "read_bytes", autospec=True, side_effect=lambda path: b"changed" if path == html_path else original(path)):
                with self.assertRaisesRegex(SafetyError, "INTEGRITY_CHANGED"):
                    write_issue_package(issue, dataset, folder)

    def test_second_language_cannot_bless_corrupted_shared_workbook(self):
        issue, dataset = fixture()
        with tempfile.TemporaryDirectory() as folder:
            packet = write_issue_package(issue, dataset, folder)
            workbook = Path(packet["artifacts"]["workbook"]["path"])
            with patch.object(Path, "read_bytes", autospec=True, side_effect=lambda path: b"changed" if path == workbook else original(path)):
                with self.assertRaisesRegex(SafetyError, "INTEGRITY_CHANGED"):
                    write_issue_package(issue, dataset, folder, language="hi-Latn")

    def test_interrupted_before_workbook_write_recovers_without_new_facts_or_model_work(self):
        issue, dataset = fixture()
        with tempfile.TemporaryDirectory() as folder:
            facts = Path(folder) / "data" / "outputs" / issue["issue_id"] / "revision-1" / "final_report_data.json"
            atomic_json(facts, dataset)
            packet = write_issue_package(issue, dataset, folder)
            self.assertEqual(packet["facts_hash"], digest(dataset))
            self.assertTrue(Path(packet["artifacts"]["workbook"]["path"]).exists())

    def test_excel_formula_prefixes_and_unknown_values_remain_text(self):
        issue, dataset = fixture()
        dataset["products"][0]["title"] = " \t=HYPERLINK(\"https://example.org\")"
        dataset["products"][0]["prices"] = []
        with tempfile.TemporaryDirectory() as folder:
            packet = write_issue_package(issue, dataset, folder)
            book = load_workbook(packet["artifacts"]["workbook"]["path"])
            self.assertFalse(any(cell.data_type == "f" for sheet in book for row in sheet for cell in row))
            self.assertEqual(book["Shortlist"].cell(7, 7).value, "Not stated")
            self.assertEqual(book["Shortlist"].cell(7, 9).value, issue["products"][0]["concept_id"])
            self.assertEqual(book["Shortlist"].freeze_panes, "C7")


original = Path.read_bytes

if __name__ == "__main__":
    unittest.main()
