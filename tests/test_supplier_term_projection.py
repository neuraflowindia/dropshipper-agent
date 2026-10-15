"""Versioned supplier wording; only disposable synthetic report fixtures."""

from copy import deepcopy
from io import BytesIO
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tests"))

from test_report_source_facts import source_fixture, select
from test_report_selection import stated
from product_radar.delivery_package import _workbook
from product_radar.live_outputs import (
    LEGACY_SUPPLIER_TERMS_DISPLAY_VERSION, SUPPLIER_TERMS_DISPLAY_VERSION,
    build_issue_dataset, render_html, supplier_display, supplier_terms_display,
    supplier_workbook_rows, validate_dataset, write_issue_edition,
    write_issue_workbook, _workbook_bytes,
)
from product_radar.models import SafetyError, digest
from product_radar.suppliers import SUPPLIER_FIELDS
from test_fit_pipeline import older_format_scores


def term(field, value, basis="observed"):
    return {"field": field, "value": value, "basis": basis}


def lead(terms=(), relationship="exact_product"):
    return {"lead_id": "synthetic-supplier", "name": "Synthetic supplier",
            "url": "https://example.org/synthetic-supplier",
            "single_piece_dispatch": "not_stated", "subscription": "not_stated",
            "min_order": None, "ecommerce_support": "Not stated",
            "reason": "Synthetic supplier statement.", "relationship": relationship,
            "terms": list(terms)}


def packet(memory=None):
    _, issue, dataset = select(source_fixture() if memory is None else memory)
    dataset["narratives"] = {language: {
        "title": "Synthetic picks" if language == "en" else "Test ke sujhaav",
        "products": [{"concept_id": product["concept_id"],
                      "reason": "{{title}} is a synthetic example." if language == "en" else "{{title}} sirf test ka example hai.",
                      "points": [], "limitation": "", "update_sentence": "",
                      "checklist_phrases": [{"id": row["id"], "text": "Fixture judgment" if language == "en" else "Test ki rai"}
                                            for row in product["criteria"]]}
                     for product in dataset["products"]]}
        for language in ("en", "hi-Latn")}
    validate_dataset(issue, dataset)
    return issue, dataset


def without_marker(issue, dataset):
    issue, dataset = deepcopy(issue), deepcopy(dataset)
    for product in issue["products"] + dataset["products"]:
        product.pop("supplier_terms_display_version", None)
    issue["selection_hash"] = digest({field: issue[field] for field in
        ("issue_id", "client_id", "revision", "date", "products", "evidence_fingerprint")})
    dataset["selection_hash"] = issue["selection_hash"]
    validate_dataset(issue, dataset)
    return issue, dataset


@older_format_scores
class SupplierTermProjectionTests(unittest.TestCase):
    def display(self, supplier, language="en"):
        return supplier_terms_display(supplier, language, terms_version=SUPPLIER_TERMS_DISPLAY_VERSION)

    def test_absent_marker_preserves_exact_original_two_capability_strings(self):
        supplier = lead([term("single_piece_purchase", "No"), term("direct_customer_dispatch", "Yes"),
                         term("cod", "Yes"), term("white_label_shipping", "Yes")])
        self.assertEqual(supplier_terms_display(supplier), "Single-piece purchase: No; Direct customer dispatch: Yes")
        self.assertEqual(supplier_terms_display(supplier, "hi-Latn"),
                         "Single-piece kharid: Nahi; Seedha customer ko dispatch: Haan")
        self.assertEqual(supplier_terms_display(supplier, terms_version=LEGACY_SUPPLIER_TERMS_DISPLAY_VERSION),
                         supplier_terms_display(supplier))
        self.assertNotIn("Cash on delivery", supplier_display(supplier))

    def test_all_seventeen_observed_terms_use_fixed_labels_in_both_languages(self):
        values = {field: "Yes" for field in SUPPLIER_FIELDS}
        values.update(subscription_requirement="Optional", recurring_fee_frequency="Every 30 days",
                      monthly_fee="INR 999 plus tax", purchase_moq="96 sets of 3",
                      india_coverage="Delhi and NCR", setup_fee="INR 499 once", deposit="INR 1500 refundable",
                      per_order_fee="INR 12 per parcel", bulk_stock_prerequisite="No")
        supplier = lead([term(field, values[field]) for field in reversed(SUPPLIER_FIELDS)])
        before = deepcopy(supplier)
        english, hinglish = self.display(supplier), self.display(supplier, "hi-Latn")
        self.assertEqual(english.count(": "), 17)
        self.assertEqual(hinglish.count(": "), 17)
        for text in ("Single-piece purchase: Yes", "Direct customer dispatch: Yes", "Minimum purchase: 96 sets of 3",
                     "Subscription requirement: Optional", "Recurring fee frequency: Every 30 days",
                     "Monthly fee: INR 999 plus tax", "Bulk stock must be bought first: No", "Cash on delivery: Yes",
                     "Delivery coverage: Delhi and NCR", "Product offered under your brand: Yes",
                     "Shipping under your brand: Yes", "Plain packaging without supplier branding: Yes",
                     "Support for online selling: Yes", "Order-system connection: Yes", "Setup fee: INR 499 once",
                     "Deposit: INR 1500 refundable", "Fee for each order: INR 12 per parcel"):
            self.assertIn(text, english)
        for text in ("Single-piece kharid: Haan", "Seedha customer ko dispatch: Haan", "Minimum kharid: 96 sets of 3",
                     "Subscription ki zaroorat: Aapki marzi par", "Baar-baar lagne wali fee kab deni hai: Every 30 days",
                     "Har mahine ki fee: INR 999 plus tax", "Pehle bulk stock kharidna zaroori hai: Nahi",
                     "Delivery par payment: Haan", "Delivery kahan tak hai: Delhi and NCR",
                     "Aapke brand ke naam se product: Haan", "Aapke brand ke naam se shipping: Haan",
                     "Supplier ke brand ke bina saada packing: Haan", "Online bechne mein madad: Haan",
                     "Order system se connection: Haan", "Shuru karne ki fee: INR 499 once",
                     "Jama karne ki rakam: INR 1500 refundable", "Har order ki fee: INR 12 per parcel"):
            self.assertIn(text, hinglish)
        self.assertLess(english.index("Minimum purchase"), english.index("Cash on delivery"))
        self.assertLess(english.index("Cash on delivery"), english.index("Fee for each order"))
        self.assertEqual(supplier, before)

    def test_observed_no_is_a_real_negative_capability_not_an_unknown(self):
        supplier = lead([term("cod", "No"), term("white_label_shipping", "No")])
        self.assertIn("Cash on delivery: No", self.display(supplier))
        self.assertIn("Shipping under your brand: No", self.display(supplier))
        self.assertIn("Delivery par payment: Nahi", self.display(supplier, "hi-Latn"))

    def test_each_subscription_state_is_distinct_without_inferring_from_missing_fees(self):
        for value, expected in (("Required", "Zaroori"), ("Optional", "Aapki marzi par"), ("Not required", "Zaroori nahi")):
            with self.subTest(value=value):
                supplier = lead([term("subscription_requirement", value)])
                self.assertIn("Subscription requirement: " + value, self.display(supplier))
                self.assertIn("Subscription ki zaroorat: " + expected, self.display(supplier, "hi-Latn"))
                self.assertNotIn("Monthly fee", self.display(supplier))

    def test_known_optional_or_conflicting_subscription_does_not_say_terms_are_unstated(self):
        for value, basis, english, hinglish in (
                ("Optional", "observed", "Optional subscription stated by the supplier",
                 "Supplier ne optional subscription ka zikr kiya hai"),
                ("Conflicting", "conflicting", "Supplier subscription statements conflict",
                 "Supplier ki subscription ki shartein aapas mein nahi milti hain")):
            supplier = lead([term("subscription_requirement", value, basis)])
            for language, expected in (("en", english), ("hi-Latn", hinglish)):
                with self.subTest(value=value, language=language):
                    text = supplier_display(supplier, language, terms_version=SUPPLIER_TERMS_DISPLAY_VERSION)
                    self.assertIn(expected, text)
                    self.assertNotIn("Subscription terms not stated", text)
                    self.assertNotIn("Subscription ki shartein nahi di gayi hain", text)
            self.assertIn("Subscription terms not stated", supplier_display(supplier))
            product = {"rank": 1, "title": "Synthetic product", "concept_id": "c000", "suppliers": [supplier],
                       "supplier_terms_display_version": SUPPLIER_TERMS_DISPLAY_VERSION}
            self.assertEqual(supplier_workbook_rows(product)[0][4], english)
            product.pop("supplier_terms_display_version")
            self.assertEqual(supplier_workbook_rows(product)[0][4], "Subscription terms not stated")

    def test_absent_and_not_stated_extras_do_not_pad_or_imply_free_service(self):
        for terms in ([], [term(field, "Not stated", "not_stated") for field in SUPPLIER_FIELDS]):
            supplier = lead(terms)
            self.assertEqual(self.display(supplier), "Single-piece purchase: Unknown; Direct customer dispatch: Unknown")
            self.assertEqual(self.display(supplier, "hi-Latn"),
                             "Single-piece kharid: Abhi clear nahi hai; Seedha customer ko dispatch: Abhi clear nahi hai")
            self.assertNotIn("No subscription", self.display(supplier))

    def test_inferred_invalid_duplicate_and_blank_extras_never_become_confirmed_capabilities(self):
        for rows in ([term("cod", "Yes", "inferred")], [term("cod", "Sometimes")],
                     [term("cod", "Yes"), term("cod", "Yes")], [term("cod", "")],
                     [term("cod", None)], [term("cod", "Unknown")], [term("cod", "Not stated")]):
            with self.subTest(rows=rows):
                self.assertIn("Cash on delivery: Unknown", self.display(lead(rows)))
                self.assertIn("Delivery par payment: Abhi clear nahi hai", self.display(lead(rows), "hi-Latn"))

    def test_conflicting_terms_are_visible_without_selecting_the_favourable_statement(self):
        supplier = lead([term("cod", "Conflicting", "conflicting"),
                         term("subscription_requirement", "Conflicting", "conflicting")])
        self.assertIn("Cash on delivery: Conflicting statements", self.display(supplier))
        self.assertIn("Subscription requirement: Conflicting statements", self.display(supplier))
        self.assertIn("Delivery par payment: Shartein aapas mein nahi milti hain", self.display(supplier, "hi-Latn"))
        self.assertNotIn("Cash on delivery: Yes", self.display(supplier))

    def test_company_only_terms_do_not_confirm_this_product_or_a_product_price(self):
        for relationship in ("relevant_company", "unclear"):
            supplier = lead([term("cod", "Yes")], relationship)
            self.assertIn("Company terms only; supply of this product is not confirmed.", self.display(supplier))
            self.assertIn("is product ki supply abhi confirm nahi hai", self.display(supplier, "hi-Latn"))
            self.assertNotIn("Supplier quote", self.display(supplier))
        self.assertNotIn("Company terms only", self.display(lead([term("cod", "Yes")])))

    def test_fixed_failures_for_unknown_versions_and_languages(self):
        for version in (None, True, {}, "supplier-terms-v9"):
            with self.subTest(version=version), self.assertRaisesRegex(SafetyError, "LIVE_SUPPLIER_TERMS_DISPLAY_VERSION_INVALID"):
                supplier_terms_display(lead(), terms_version=version)
        for language in (None, [], "hi"):
            with self.subTest(language=language), self.assertRaisesRegex(SafetyError, "LIVE_LANGUAGE_INVALID"):
                self.display(lead(), language)

    def test_new_marker_is_frozen_for_every_selected_product_without_memory_changes(self):
        memory = source_fixture()
        before = deepcopy(memory.state)
        pool, issue, dataset = select(memory)
        for rows in (pool["product_facts"], pool["eligible"], issue["products"], dataset["products"]):
            self.assertTrue(rows)
            self.assertTrue(all(row["supplier_terms_display_version"] == SUPPLIER_TERMS_DISPLAY_VERSION for row in rows))
        self.assertEqual(memory.state, before)
        self.assertEqual(memory.state["jobs"], {})

    def test_marker_removal_or_change_cannot_silently_change_a_frozen_packet(self):
        issue, dataset = packet()
        for value in (None, "supplier-terms-v1", "supplier-terms-v9", True):
            changed = deepcopy(dataset)
            if value is None:
                changed["products"][0].pop("supplier_terms_display_version")
            else:
                changed["products"][0]["supplier_terms_display_version"] = value
            with self.subTest(value=value), self.assertRaisesRegex(SafetyError, "LIVE_SUPPLIER_TERMS_BINDING_CHANGED"):
                validate_dataset(issue, changed)

    def test_supplied_marker_cannot_upgrade_an_old_selection(self):
        issue, dataset = without_marker(*packet())
        supplied = deepcopy(dataset["products"])
        for product in supplied:
            product["supplier_terms_display_version"] = SUPPLIER_TERMS_DISPLAY_VERSION
        replay = build_issue_dataset(issue, supplied, narratives=dataset["narratives"], evidence_type="synthetic")
        self.assertTrue(all("supplier_terms_display_version" not in row for row in replay["products"]))
        validate_dataset(issue, replay)
        changed = deepcopy(dataset)
        changed["products"][0]["supplier_terms_display_version"] = SUPPLIER_TERMS_DISPLAY_VERSION
        with self.assertRaisesRegex(SafetyError, "LIVE_SUPPLIER_TERMS_BINDING_CHANGED"):
            validate_dataset(issue, changed)

    def test_frozen_new_marker_requires_exact_version_and_current_source_scope(self):
        for field, value in (("supplier_terms_display_version", "supplier-terms-v1"),
                             ("supplier_terms_display_version", True),
                             ("opportunity_policy_id", "opportunity-priority-v1"),
                             ("source_route_contract", "unknown-source-contract")):
            with self.subTest(field=field, value=value):
                issue, dataset = packet()
                issue["products"][0][field] = value
                dataset["products"][0][field] = value
                issue["selection_hash"] = digest({name: issue[name] for name in
                    ("issue_id", "client_id", "revision", "date", "products", "evidence_fingerprint")})
                dataset["selection_hash"] = issue["selection_hash"]
                expected = "LIVE_FIT_POLICY_MISMATCH" if field == "opportunity_policy_id" else "LIVE_SUPPLIER_TERMS_DISPLAY_VERSION_INVALID"
                with self.assertRaisesRegex(SafetyError, expected):
                    validate_dataset(issue, dataset)

    def test_known_terms_reach_both_html_editions_and_excel_without_changing_product_prices(self):
        memory = source_fixture()
        for field, value in (("cod", "Yes"), ("white_label_shipping", "Yes"), ("bulk_stock_prerequisite", "No"),
                             ("monthly_fee", "INR 999 plus tax"), ("purchase_moq", "96 sets of 3"),
                             ("india_coverage", "Delhi and NCR")):
            stated(memory, field, value)
        issue, dataset = packet(memory)
        product = next(row for row in dataset["products"] if row["concept_id"] == "c000")
        before = deepcopy(dataset)
        english, hinglish = render_html(issue, dataset), render_html(issue, dataset, language="hi-Latn")
        self.assertIn("Cash on delivery: Yes", english)
        self.assertIn("Shipping under your brand: Yes", english)
        self.assertIn("Bulk stock must be bought first: No", english)
        self.assertIn("Monthly fee: INR 999 plus tax", english)
        self.assertIn("Delivery coverage: Delhi and NCR", english)
        self.assertIn("Delivery par payment: Haan", hinglish)
        self.assertIn("Aapke brand ke naam se shipping: Haan", hinglish)
        self.assertIn("Pehle bulk stock kharidna zaroori hai: Nahi", hinglish)
        self.assertIn("Har mahine ki fee: INR 999 plus tax", hinglish)
        self.assertNotIn("Supplier quote: INR 999", english)
        self.assertEqual(dataset, before)
        raw = _workbook_bytes(issue, dataset, deterministic=True)
        _workbook(raw, dataset)
        book = load_workbook(BytesIO(raw), read_only=True, data_only=False)
        try:
            rows = list(book["Supplier leads"].iter_rows(min_row=7, values_only=True))
            row = next(row for row in rows if row[-1] == "c000")
            self.assertEqual(row[5], "96 sets of 3")
            self.assertIn("Cash on delivery: Yes", row[7])
            self.assertIn("Monthly fee: INR 999 plus tax", row[7])
            shortlist = next(row for row in book["Shortlist"].iter_rows(min_row=7, values_only=True) if row[8] == "c000")
            self.assertNotIn("999", shortlist[6])
            self.assertNotIn("999", shortlist[7])
        finally:
            book.close()
        self.assertEqual(supplier_workbook_rows(product)[0][5], "96 sets of 3")

    def test_free_text_terms_are_escaped_as_text_not_instructions_or_html(self):
        memory = source_fixture()
        stated(memory, "india_coverage", "Delhi <script>alert('fixture')</script> & NCR")
        issue, dataset = packet(memory)
        html = render_html(issue, dataset)
        self.assertIn("Delivery coverage: Delhi &lt;script&gt;alert(&#x27;fixture&#x27;)&lt;/script&gt; &amp; NCR", html)
        self.assertNotIn("<script>alert('fixture')</script>", html)
        self.assertIn("<script>", supplier_workbook_rows(dataset["products"][0])[0][7])
        _workbook(_workbook_bytes(issue, dataset, deterministic=True), dataset)

    def test_old_and_new_stages_replay_exactly_but_in_place_contract_upgrade_is_rejected(self):
        memory = source_fixture()
        stated(memory, "cod", "Yes")
        new_issue, new_dataset = packet(memory)
        old_issue, old_dataset = without_marker(new_issue, new_dataset)
        old_html = render_html(old_issue, old_dataset)
        self.assertNotIn("Cash on delivery", old_html)
        self.assertNotIn("Cash on delivery", supplier_workbook_rows(old_dataset["products"][0])[0][7])
        with tempfile.TemporaryDirectory() as directory:
            workbook = write_issue_workbook(old_issue, old_dataset, directory)
            edition = write_issue_edition(old_issue, old_dataset, directory)
            self.assertEqual(write_issue_workbook(old_issue, old_dataset, directory), workbook)
            self.assertEqual(write_issue_edition(old_issue, old_dataset, directory), edition)
            _workbook(Path(workbook["artifacts"]["workbook"]["path"]).read_bytes(), old_dataset)
            with self.assertRaisesRegex(SafetyError, "LIVE_STAGE_BINDING_CHANGED"):
                write_issue_workbook(new_issue, new_dataset, directory)
        with tempfile.TemporaryDirectory() as directory:
            workbook = write_issue_workbook(new_issue, new_dataset, directory)
            english = write_issue_edition(new_issue, new_dataset, directory)
            hinglish = write_issue_edition(new_issue, new_dataset, directory, language="hi-Latn")
            self.assertEqual(write_issue_workbook(new_issue, new_dataset, directory), workbook)
            self.assertEqual(write_issue_edition(new_issue, new_dataset, directory), english)
            self.assertEqual(write_issue_edition(new_issue, new_dataset, directory, language="hi-Latn"), hinglish)
            self.assertNotEqual(workbook["core_facts_hash"], digest({key: value for key, value in old_dataset.items()
                if key not in ("narratives", "selected_language")}))


if __name__ == "__main__":
    unittest.main()
